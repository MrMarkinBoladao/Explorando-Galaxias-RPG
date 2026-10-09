# -*- coding: utf-8 -*-
"""
oraculo_mestre_hist.py — o oráculo das abas da Fase 2: NPCs (G = 200, 6.6), Aventuras (G = 300, 6.11) e Recompensas
(G = 400/410/420, 6.12). Mesma convenção de oraculo_mestre_abas.py: entrada {nome lógico da entrada: valor}, saída
{nome lógico: valor esperado}; nenhuma fórmula é lida e build\\mestre_dados.py / build\\mestre\\* não são importados.
As tabelas do livro estão TRANSCRITAS abaixo com a seção; as listas de sabor (aba Tabelas) são CONTEÚDO e vêm do
.xlsx gerado (valores, não fórmulas — design §3.1), na ordem das vagas do mapa (build\\mestre_mapa.json).
"""
import json
from functools import lru_cache
from pathlib import Path

import oraculo_mestre as S
from oraculo_mestre_abas import num, tx, clamp
from oraculo_mestre_livro import FAIXAS, ORC, DT, VERBA, COMPOS, RACAS, bestiario_md, ancora

BUILD = Path(__file__).resolve().parent
M = S.M

# ---------------------------------------------------------------------------
# Tabelas do livro transcritas (com a seção)
# ---------------------------------------------------------------------------
CAMINHOS = ["A Destruição", "A Inexistência", "A Harmonia", "A Abundância", "A Recordação", "A Erudição",
            "A Euforia", "A Caça", "A Preservação"]                                                    # 06.3
PAPEIS = ["Aliado", "Contratante", "Rival", "Neutro", "Informante", "Vítima", "Antagonista"]            # 6.6 (H11)
# 05 — montagem do nome (ordem 2 = família primeiro); a planilha guarda isso na aba Dados
MONTAGEM = {"Humano": (1, " "), "Xianzhouíta": (2, " "), "Vidyadhara": (1, " "), "Vulpes": (1, " "),
            "Haloviano": (1, " "), "Avginiano": (1, " do clã "), "Intellitron": (1, ", que se chama ")}
SLUG = {"Humano": "humano", "Xianzhouíta": "xianzhouita", "Vidyadhara": "vidyadhara", "Vulpes": "vulpes",
        "Haloviano": "haloviano", "Avginiano": "avginiano", "Intellitron": "intellitron"}
# 27.16 — fichas de cada facção
FACCAO_FICHAS = {
    "Legião da Antimatéria": ["Peão da Antimatéria", "Sargento de Trincheira", "Fuzileiro da Legião",
                              "Centurião Catafracto", "Pretor Vazio-Nove", "Guarda Pretoriano da Antimatéria",
                              "Arcanjo de Ferro-Vazio"],
    "Corporação da Paz Interastral": ["Dron de Vigilância Pacificadora", "Auditora de Risco"],
    "Aliança Xianzhou": ["Lanceiro da Frota de Jade", "Oficial-Lâmina da Frota de Jade", "Centenário Enlouquecido"],
    "Caçadores de Stellaron": ["Operativo de Campo dos Caçadores", "Executora de Contrato"],
    "Tolos Mascarados": ["Palhaço de Pólvora", "Mestre de Cerimônias Mascarado", "O Dramaturgo de Mil Faces"],
    "Cavaleiros da Beleza": ["Noviço da Beleza", "Escultor de Carne", "Vênia, a Primeira Obra"],
    "Culto da Inexistência": ["Devoto do Silêncio", "Arauto da Tempestade Vazia", "Vazia Coroada"],
    "Autômatos": ["Sentinela Enferrujada", "Autômato Taciturno", "Dron de Vigilância Pacificadora"],
}
# 27.17 — faixa da cena (Ferro-Vazio "Faixa alta" = 13-16 e 17-20: H26)
LOCAIS_FAIXA = {"Poço Sete": (1, 1), "Ferro-Vazio": (4, 5), "Hesperin": (5, 5)}
SEM_MATAR = ["fugir com a carga", "aguentar 3 Ciclos", "chegar ao console"]                              # 27.7
# 24.3 — consumíveis (poção: faixa da H9) na ordem da tabela
POCOES = [("Pequena", 15, 0.5, 50), ("Média", 30, 1, 150), ("Grande", 50, 2, 400)]
FAIXA_POCAO = {"Pequena": (1, 2), "Média": (2, 4), "Grande": (4, 5)}
ITENS = [("Kit de primeiros socorros (3 usos)", 1, 100, "Cura 10 fora de combate, por uso"),
         ("Kit de ferramentas", 1, 120, "Permite Testes de Tecnologia que exigem equipamento"),
         ("Kit de pesquisa de campo", 1, 150, "Idem, para Ciência e Pesquisa"),
         ("Lanterna de plasma", 0.5, 30, "Luz, 8 horas de carga"), ("Corda de fibra, 20 m", 1, 25, "Corda"),
         ("Gancho de escalada", 0.5, 60, "Sobe o que não tem escada"),
         ("Ração de viagem (3 dias)", 0.5, 15, "Descanso Longo fora de abrigo"),
         ("Comunicador de curto alcance", 0.5, 80, "Fala com o grupo entre cenas"),
         ("Scanner portátil", 1, 200, "Vantagem em um Teste de Pesquisa por cena, fora de combate"),
         ("Traje de vedação", 2, 250, "Sobrevive a vácuo e atmosfera hostil por uma cena"),
         ("Munição ou célula de reserva", 0.5, 20, "Reposição de projétil ou de célula de energia")]
# 24.1 / 24.2
ARMADURAS = [("Leve", 3, "+1 de Velocidade", 1, 150), ("Média", 5, "—", 2, 300),
             ("Pesada", 6, "2 RD, -2 em Reflexos e nas Perícias de Agilidade, -2 de Velocidade", 3, 500)]
ARMAS = [("Leve", "1d8", "Pessoal", "Agilidade", 0.5, 100), ("Média", "1d10", "Pessoal",
                                                             "Poder ou Agilidade (fixo na criação)", 1, 150),
         ("Pesada", "1d12", "Pessoal", "Poder", 2, 200), ("Disparo curto", "1d8", "Média", "Agilidade", 1, 250),
         ("Disparo longo", "1d10", "Longa", "Agilidade", 2, 350), ("Energia", "2d8", "Longa", "Sincronia", 2, 500)]
# 24.5 — o que a verba compra
COMPRA = ["reposição de poções, uma arma, um kit", "trocar uma armadura, equipar o grupo com poções Médias",
          "equipamento de cena, suborno, passagem, transporte", "contratar gente, comprar silêncio, abrir porta fechada",
          "o custo não é mais dinheiro, e a tabela está aqui para o Mestre ter um teto"]
# 25.2 — Cone de Luz por Nível: (Bônus Maior, Efeito Condicional, parte numérica, PV)
CONE = [("+1 ou +10 PV", "1 efeito simples, 1 vez por combate", 1, 10), ("+1 e +10 PV", "1 efeito, 1 vez por Ciclo", 1, 10),
        ("+2 ou +25 PV", "1 efeito + pequeno ganho de Energia", 2, 25), ("+2 e +25 PV", "2 efeitos", 2, 25),
        ("+3 ou +50 PV", "2 efeitos, um deles podendo conceder Avanço", 3, 50)]
BONUS_MAIOR = ["PV máximos", "Defesa", "Velocidade", "Dano de Ataque Básico", "Dano de Habilidade", "Dano de Ultimate",
               "RD, dentro do teto de 2 + (2 × Eficiência)", "Bônus em Teste de Ataque",
               "Bônus em um Teste de Resistência", "Bônus em uma Perícia"]                               # 25.2
CONJUNTO_2 = ["+1 em um tipo de rolagem", "+1 de Velocidade", "+2 de dano", "+1 RD"]                    # 25.3
# 25.3 — Relíquias: slot, o que dá, Tier I..IV
RELIQUIAS = [("Cabeça", "PV máximos", 10, 20, 35, 50), ("Mãos", "Dano de Ataque Básico", 2, 4, 6, 8),
             ("Tronco", "Defesa", 1, 1, 2, 2), ("Botas", "Velocidade", 2, 3, 4, 5),
             ("Esfera Planar", "Dano de um Elemento escolhido", 2, 4, 6, 8),
             ("Corda de Ligação", "Energia, 1 vez por combate", 10, 15, 20, 25)]
# 26.7 — Ressonâncias
RESSONANCIAS = {5: ("I", "1 uso por combate de uma Habilidade de Nível 3 ou menor sem custo de PH; ou +1 de Velocidade "
                         "permanente"),
                10: ("II", "Sua Ultimate ganha um efeito extra, dentro dos limites de buff e debuff lidos no Nível "
                           "equivalente da sua faixa (capítulo 17)"),
                15: ("III", "Uma Habilidade sua sobe 1 Nível de efeito e passa a custar o PH do Nível novo, sem passar "
                            "do seu Nível máximo"),
                20: ("IV", "Sua Ultimate ativa com 80 de Energia e consome 80 (o teto do medidor continua 100); ou, se "
                           "você tiver a Bênção Avatar do seu Caminho, ela afeta um alvo adicional")}
SUG = "Sugestão da planilha — não é regra do livro"


def tier(n):                                                                                  # 25.3
    return "Tier I" if n <= 6 else "Tier II" if n <= 12 else "Tier III" if n <= 17 else "Tier IV"


# ---------------------------------------------------------------------------
# Listas da aba Tabelas (conteúdo, lido do .xlsx gerado)
# ---------------------------------------------------------------------------

@lru_cache(None)
def _tabelas():
    import openpyxl
    mapa = json.loads((BUILD / "mestre_mapa.json").read_text(encoding="utf-8"))
    arq = BUILD.parent / "Mestre" / mapa["arquivo"]
    ws = openpyxl.load_workbook(arq, read_only=False)["Tabelas"]
    saida = {}
    for id_, t in mapa["tabelas"].items():
        v1 = [ws[c].value for c in t["vagas"]]
        v2 = [ws[c].value for c in t.get("segunda", [])]
        saida[id_[4:]] = (v1, v2)
    return saida


def lista(id_):
    """Valores não vazios da lista, na ordem das vagas."""
    return [tx(v) for v in _tabelas()[id_][0] if v not in (None, "") and tx(v) != ""]


def pares(id_):
    v1, v2 = _tabelas()[id_]
    return [(tx(a), tx(b)) for a, b in zip(v1, v2)]


def sorteia(id_, x):
    L_ = lista(id_)
    return L_[S.escolha(x, len(L_)) - 1] if L_ else ""


def sem_repetir(n, x2, i1):
    if n <= 1:
        return i1
    return (i1 - 1 + 1 + (x2 * (n - 1)) // M) % n + 1


def nome_cultura(raca, x_nome, x_sob):
    nome = sorteia(f"nome.{SLUG[raca]}", x_nome)
    sob = sorteia(f"sobrenome.{SLUG[raca]}", x_sob)
    ordem, sep = MONTAGEM[raca]
    if not sob:
        return nome
    if not nome:
        return sob
    return sob + sep + nome if ordem == 2 else nome + sep + sob


def efetiva_faixa(v, padrao):
    return v if v in FAIXAS else padrao


# ---------------------------------------------------------------------------
# NPCs (6.6)
# ---------------------------------------------------------------------------

def npcs(E, out):
    sem, rol = out["inicio.semente_ef"], S.rolagem_efetiva(E.get("npcs.rolagem"))
    x = lambda c: S.valor(sem, 200, rol, c)  # noqa: E731
    RA, CA, PA = tx(E.get("npcs.raca")), tx(E.get("npcs.caminho")), tx(E.get("npcs.papel"))
    FB, TB = tx(E.get("npcs.bloco_faixa")), tx(E.get("npcs.bloco_tipo"))
    raca = RA if RA in RACAS else RACAS[(x(1) * 7) // M]
    nome = nome_cultura(raca, x(2), x(3))
    caminho = CA if (CA == "Nenhum" or CA in CAMINHOS) else sorteia("caminho_npc", x(4))
    ocup = sorteia("ocupacao", x(5))
    ap = lista("aparencia_traco")
    if ap:
        i1 = S.escolha(x(6), len(ap))
        a1, a2 = ap[i1 - 1], ap[sem_repetir(len(ap), x(7), i1) - 1]
        aparencia = a1 if a2 == a1 else f"{a1}; {a2}"
    else:
        aparencia = ""
    pers = sorteia("personalidade", x(8))
    mc = [v for k, v in pares("motivacao_caminho") if caminho and k == caminho and v]
    if caminho in ("Nenhum", "") or not mc:
        motiv = sorteia("motivacao", x(9))
    else:
        motiv = mc[S.escolha(x(9), len(mc)) - 1]
    segredo, manei, atit = sorteia("segredo", x(10)), sorteia("maneirismo", x(11)), sorteia("atitude", x(12))
    fonte = S.inteiro(x(14), 1, 2)
    fx_g = out["campanha.faixa"]
    gl = sorteia(f"ganchos_{fx_g.replace('-', '_')}", x(13))
    gt = sorteia("gancho_npc", x(13))
    gancho = (gl or gt) if fonte == 1 else (gt or gl)
    papel = PA if PA in PAPEIS else PAPEIS[(x(15) * 7) // M]
    res = {"nome": nome, "raca": raca, "caminho": caminho, "ocupacao": ocup, "aparencia": aparencia,
           "personalidade": pers, "motivacao": motiv, "segredo": segredo, "maneirismo": manei, "atitude": atit,
           "gancho": gancho, "papel": papel}
    for k, v in res.items():
        out[f"npcs.res.{k}"] = v
        out[f"npcs.saida.{k}"] = v
    out["npcs.res.nome.origem"] = f"Tabelas: nome e sobrenome da cultura {raca} (estilo de 05)"
    out["npcs.res.raca.origem"] = "Escolhida" if RA in RACAS else "Sorteada entre as 7 Raças de 05"
    out["npcs.res.caminho.origem"] = ("Escolhido" if (CA == "Nenhum" or CA in CAMINHOS) else
                                      f"Tabelas: Caminho do NPC (06.3; Nenhum ×3, 27.12) · {SUG} (H11)")
    out["npcs.res.motivacao.origem"] = ("Tabelas: Motivação (sem Caminho)" if caminho in ("Nenhum", "") or not mc
                                        else "Tabelas: Motivação por Caminho (27.12)")
    out["npcs.res.papel.origem"] = "Escolhido" if PA in PAPEIS else "Sorteado entre os 7 papéis"
    out["npcs.res.gancho.origem"] = (f"27.18, faixa {fx_g}" if fonte == 1 else "Tabelas: Gancho do NPC") + \
        f" · {SUG} (H20)"
    out["npcs.g.fonte_gancho"] = fonte
    # avisos (os dois sentidos)
    rv = E.get("npcs.rolagem")
    av = ""
    if rv is not None and tx(rv) != "" and not (num(rv) and 1 <= rv <= 1000000 and int(rv) == rv):
        av = "Rolagem nº fora de 1 a 1.000.000: usando 1"
    elif RA and RA != "Sortear" and RA not in RACAS:
        av = "Raça fora da lista do capítulo 05: sorteando"
    elif CA and CA not in ("Sortear", "Nenhum") and CA not in CAMINHOS:
        av = "Caminho fora da lista: sorteando"
    elif PA and PA != "Sortear" and PA not in PAPEIS:
        av = "Papel fora da lista: sorteando"
    elif FB and FB not in ("Nenhum", "Do grupo") and FB not in FAIXAS:
        av = "Bloco de combate fora da lista: sem bloco"
    elif TB and TB not in ("Comum", "Elite"):
        av = "Tipo do bloco fora da lista: usando Comum"
    out["npcs.aviso.entradas"] = av
    out["npcs.aviso.gancho"] = "" if gancho else "Listas de gancho vazias (aba Tabelas)"
    # bloco de combate (28.3)
    bfx = fx_g if FB == "Do grupo" else (FB if FB in FAIXAS else "")
    bt = "Elite" if TB == "Elite" else "Comum"
    p = "npcs.bloco"
    if not bfx:
        out.update({f"{p}.rotulo": "Sem bloco", **{f"{p}.{k}": "" for k in ("pv", "defesa", "rd", "ten", "vel", "ataque",
                                                                              "dt", "tr", "nfraq", "dano", "nafila",
                                                                              "combate")}})
        out[f"{p}.regra"] = "Escolha a faixa do bloco (ou Do grupo) para ver os números."
    else:
        a = ancora(bfx, bt)
        out.update({f"{p}.rotulo": f"{bfx} · {bt}", f"{p}.pv": a["pv"], f"{p}.defesa": a["defesa"], f"{p}.rd": a["rd"],
                    f"{p}.ten": a["ten"], f"{p}.vel": a["vel"], f"{p}.ataque": f"+{a['ataque']}", f"{p}.dt": a["dt"],
                    f"{p}.tr": f"+{a['tr']}", f"{p}.nfraq": a["nfraq"],
                    f"{p}.dano": f"{a['dano_e']} · média {a['dano_m']}",
                    f"{p}.nafila": f"VEL {a['vel']}, {a['firmeza']}",
                    f"{p}.regra": "NPC aliado também usa as âncoras de 28.3 (sem Esquiva, PH, Energia nem Ultimate, "
                                  "28.2); Execução: ser racional pode Executar (23.5).",
                    f"{p}.combate": f"Para levar ao combate: na aba Inimigos, crie uma linha com o nome {nome}, tipo "
                                    f"{bt} e faixa {bfx} (modo Faixa do livro); ele entra nas listas de Encontros e "
                                    f"Combate."})
    _elenco(E, out)


def _elenco(E, out):
    campos = ("nome", "raca", "caminho", "ocupacao", "aparencia", "maneirismo", "motivacao", "segredo", "atitude",
              "papel")
    el = [{c: tx(E.get(f"npcs.elenco.{i}.{c}")) for c in campos} for i in range(1, 31)]
    nomes = [x["nome"] for x in el]
    best = {f["nome"] for f in bestiario_md()}
    for i, x in enumerate(el, start=1):
        n = x["nome"]
        out[f"npcs.elenco.{i}.aviso"] = ("" if not n else "Nome repetido no elenco" if nomes.count(n) > 1 else
                                         "Nome igual a uma criatura do Bestiário" if n in best else "")
        for s in range(2, 6):
            out[f"npcs.elenco.{i}.rot{s}"] = n if n else f"NPC {i}"
        rel = E.get(f"npcs.elenco.{i}.relacao")
        out[f"npcs.elenco.{i}.aviso_rel"] = ("Relação fora de −3 a +3" if num(rel) and (rel < -3 or rel > 3 or
                                                                                    int(rel) != rel) else
                                             "Relação não é número (−3 a +3)" if rel is not None and tx(rel) != ""
                                             and not num(rel) else "")
    for c in range(1, 5):
        v = tx(E.get(f"npcs.cartao{c}.npc"))
        k = nomes.index(v) + 1 if v and v in nomes else None
        x = el[k - 1] if k else None
        out.update({
            f"npcs.cartao{c}.nome": x["nome"] if x else "",
            f"npcs.cartao{c}.quem": f'{x["raca"]} · {x["caminho"]} · {x["ocupacao"]}' if x else "",
            f"npcs.cartao{c}.aparencia": f'Aparência: {x["aparencia"]}' if x else "",
            f"npcs.cartao{c}.voz": f'Voz: {x["maneirismo"]}' if x else "",
            f"npcs.cartao{c}.quer": f'Quer: {x["motivacao"]}' if x else "",
            f"npcs.cartao{c}.esconde": f'Esconde: {x["segredo"]}' if x else "",
            f"npcs.cartao{c}.atitude": f'Atitude: {x["atitude"]} · Papel: {x["papel"]}' if x else ""})
    for b in (1, 2):
        msg = ""
        for c in (2 * b - 1, 2 * b):
            v = tx(E.get(f"npcs.cartao{c}.npc"))
            if v and v not in nomes and not msg:
                msg = f"Cartão {c}: NPC fora do Elenco"
        out[f"npcs.cartoes.aviso{b}"] = msg


# ---------------------------------------------------------------------------
# Aventuras (6.11)
# ---------------------------------------------------------------------------

def aventuras(E, out):
    sem, rol = out["inicio.semente_ef"], S.rolagem_efetiva(E.get("aventuras.rolagem"))
    x = lambda c: S.valor(sem, 300, rol, c)  # noqa: E731
    TP, FA, FC = tx(E.get("aventuras.tipo")), tx(E.get("aventuras.faixa")), tx(E.get("aventuras.faccao"))
    tipos, facs = lista("tipo_aventura"), lista("faccoes")
    tipo = TP if (TP and TP != "Sortear" and TP in ["Sortear"] + tipos) else sorteia("tipo_aventura", x(1))
    fx = efetiva_faixa(FA, out["campanha.faixa"])
    fi = FAIXAS.index(fx) + 1
    fc = FC if (FC and FC != "Sortear" and FC in facs) else sorteia("faccoes", x(20))
    fc_k = facs.index(fc) + 1 if fc in facs else 0
    n = len(facs)
    if n:
        k = sem_repetir(n, x(3), fc_k) if fc_k > 0 else S.escolha(x(3), n)
        c_fac = facs[k - 1]
    else:
        c_fac = ""
    c_raca = RACAS[(x(4) * 7) // M]
    c_nome = nome_cultura(c_raca, x(17), x(21))
    c_ocup = sorteia("ocupacao", x(16))
    contr = c_nome + (f", {c_ocup}" if c_ocup else "") + (f" ({c_fac})" if c_fac else "")
    ob = [v for k_, v in pares("objetivo") if tipo and k_ == tipo and v]
    objetivo = ob[S.escolha(x(5), len(ob)) - 1] if ob else ""
    ll = [v for v in lista("locais") if v not in LOCAIS_FAIXA or LOCAIS_FAIXA[v][0] <= fi <= LOCAIS_FAIXA[v][1]]
    local_l = ll[S.escolha(x(6), len(ll)) - 1] if ll else ""
    local_t = sorteia("local", x(6))
    fonte_l = S.inteiro(x(18), 1, 2)
    local = (local_l or local_t) if fonte_l == 1 else (local_t or local_l)
    # antagonista (H22)
    fichas = bestiario_md()
    da = FACCAO_FICHAS.get(fc, [])
    bf = [f for f in fichas if f["tipo"] == "Boss" and f["faixa"] == fx and f["nome"] in da]
    ef = [f for f in fichas if f["tipo"] == "Elite" and f["faixa"] == fx and f["nome"] in da]
    ba = [f for f in fichas if f["tipo"] == "Boss" and f["faixa"] == fx]
    deg, cands = (1, bf) if bf else (2, ef) if ef else (3, ba) if ba else (0, [])
    an = cands[S.escolha(x(7), len(cands)) - 1] if cands else None
    antag = f"{an['nome']} ({an['tipo']}, faixa {an['faixa']}, {an['faccao']})" if an else ""
    compl, revir, prazo = sorteia("complicacao", x(8)), sorteia("reviravolta", x(9)), sorteia("prazo", x(10))
    comp = COMPOS[S.escolha(x(11), 4) - 1][0]
    c4 = S.inteiro(x(12), 1, 2)
    sm = SEM_MATAR[S.escolha(x(19), 3) - 1]
    cons = consumiveis(fi)
    item, preco = cons[S.escolha(x(13), len(cons)) - 1]
    pista = sorteia("pista", x(14))
    fg = S.inteiro(x(15), 1, 2)
    gl, gt = sorteia(f"ganchos_{fx.replace('-', '_')}", x(2)), sorteia("gancho", x(2))
    gancho = (gl or gt) if fg == 1 else (gt or gl)
    npj = out["campanha.jogadores_ef"]
    orc4 = ORC[fi - 1][1]
    orc = orc4 if npj == 4 else orc4 * npj // 4
    dtm, dtd = DT["Média"][fi - 1], DT["Difícil"][fi - 1]
    verba = VERBA[fi - 1]
    p = "aventuras"
    out.update({
        f"{p}.res.tipo": tipo, f"{p}.res.gancho": gancho, f"{p}.res.contratante": contr, f"{p}.res.objetivo": objetivo,
        f"{p}.res.local": local, f"{p}.res.faccao": fc, f"{p}.res.antagonista": antag, f"{p}.res.complicacao": compl,
        f"{p}.res.reviravolta": revir, f"{p}.res.prazo": prazo,
        f"{p}.res.recompensa": f"Verba de marco: {verba} Cr para o grupo ao subir de nível (24.5) · equipamento de faixa "
                               f"se esta aventura fechar o arco (25.1) · 1 {item} ({preco} Cr, 24.3) · pista: {pista}",
        f"{p}.res.recompensa.origem": f"Faixa {fx}; o consumível é {SUG} (H9)",
        f"{p}.res.tipo.origem": "Escolhido" if TP and TP != "Sortear" else "Tabelas: Tipo de aventura",
        f"{p}.res.gancho.origem": (f"27.18, faixa {fx}" if fg == 1 else "Tabelas: Gancho de aventura") + f" · {SUG} (H20)",
        f"{p}.res.local.origem": ("27.17 (lugar que serve à faixa)" if fonte_l == 1 and local_l else "Tabelas: Local")
        + f" · {SUG} (H26)",
        f"{p}.res.faccao.origem": "Escolhida" if FC and FC != "Sortear" else "Sorteada (27.16)",
        f"{p}.res.antagonista.origem": ["Boss da facção na faixa", "Elite da facção na faixa", "Boss da faixa"][
            max(1, deg) - 1] + f" (28.11, 27.16) · {SUG} (H22)",
        f"{p}.cena1.tipo": "Social ou descoberta", f"{p}.cena1.texto": gancho, f"{p}.cena1.dt": dtm,
        f"{p}.cena1.orc": "—", f"{p}.cena2.dt": dtm, f"{p}.cena2.orc": "—",
        f"{p}.cena3.texto": f"Composição {comp} a 50% do orçamento: cena de passagem, cerca de 2 Ciclos (27.4).",
        f"{p}.cena3.dt": "—", f"{p}.cena3.orc": orc // 2,
        f"{p}.cena4.tipo": "Combate" if c4 == 1 else "Objetivo",
        f"{p}.cena4.texto": compl + " — " + ("e um encontro típico (27.4)." if c4 == 1 else
                                             f"e um objetivo que não seja matar todo mundo: {sm} (27.7)."),
        f"{p}.cena4.dt": dtd, f"{p}.cena4.orc": orc if c4 == 1 else "—",
        f"{p}.cena5.texto": f"Antagonista: {antag}. Cumpra o contrato da Fraqueza (27.5); três combates no dia é um dia "
                            f"de verdade, quatro é emergência (27.7).",
        f"{p}.cena5.dt": dtd, f"{p}.cena5.orc": orc,
        f"{p}.cenas.resumo": f"Orçamento da faixa {fx} para {npj} jogadores: {orc} PV (27.4" +
                             (")" if npj == 4 else f"; grupo diferente de 4: {SUG} (H5))") +
                             f". DT Média {dtm}, Difícil {dtd} (27.2).",
        f"{p}.saida.missao": f"{tipo} — {local}", f"{p}.saida.tipo": tipo, f"{p}.saida.contratante": contr,
        f"{p}.saida.objetivo": objetivo, f"{p}.saida.local": local, f"{p}.saida.prazo": prazo,
        f"{p}.saida.estado": "Oferecida",
        f"{p}.saida.recompensa": f"{verba} Cr · {item} · equipamento de faixa se fechar o arco",
        f"{p}.aviso.objetivo": "" if ob else f"Sem objetivo para o tipo {tipo} na aba Tabelas",
    })
    rv = E.get("aventuras.rolagem")
    av = ""
    if rv is not None and tx(rv) != "" and not (num(rv) and 1 <= rv <= 1000000 and int(rv) == rv):
        av = "Rolagem nº fora de 1 a 1.000.000: usando 1"
    elif TP and TP != "Sortear" and TP not in tipos:
        av = "Tipo fora da lista (aba Tabelas): sorteando"
    elif FA and FA not in FAIXAS:
        av = "Faixa fora da lista: usando a do grupo"
    elif FC and FC != "Sortear" and FC not in facs:
        av = "Facção fora da lista (aba Tabelas): sorteando"
    out[f"{p}.aviso.entradas"] = av
    out["aventuras.g.an_degrau"] = deg
    out["aventuras.g.an_nome"] = an["nome"] if an else ""
    out["aventuras.g.fonte_gancho"], out["aventuras.g.fonte_local"] = fg, fonte_l
    out["aventuras.g.c_fac"], out["aventuras.g.fc"] = c_fac, fc


def consumiveis(fi):
    """24.3 na ordem da tabela: as poções da faixa (H9) e os itens comuns."""
    c = [(f"Poção {nm}", pr) for nm, _, _, pr in POCOES if FAIXA_POCAO[nm][0] <= fi <= FAIXA_POCAO[nm][1]]
    return c + [(nm, pr) for nm, _, pr, _ in ITENS]


# ---------------------------------------------------------------------------
# Recompensas (6.12)
# ---------------------------------------------------------------------------

def recompensas(E, out):
    NV, MC = E.get("recompensas.nivel"), tx(E.get("recompensas.marco"))
    ef = clamp(int(NV), 1, 20) if num(NV) else out["campanha.nivel_ef"]
    fx = FAIXAS[(ef - 1) // 4]
    fi = FAIXAS.index(fx) + 1
    cone = fi
    tr = tier(ef)
    res = RESSONANCIAS.get(ef)
    nova = ef in (1, 5, 9, 13, 17) or MC == "Entrou em faixa nova"
    p = "recompensas.marco"
    out.update({
        f"{p}.nivel_ef": ef, f"{p}.faixa": fx, f"{p}.verba": VERBA[fi - 1],
        f"{p}.verba.regra": f"Para o grupo, por nível ganho (24.5). O que ela compra: {COMPRA[fi - 1]}.",
        f"{p}.cone": f"Nível {cone}",
        f"{p}.cone.regra": (f"Faixa nova: cada personagem pode trocar o Cone pelo Nível {cone} (25.1). Trocar é "
                            f"opcional: Nível maior não é automaticamente melhor (25.2)." if nova else
                            "Mesmo Cone máximo da faixa (25.1)."),
        f"{p}.bonus": CONE[cone - 1][0], f"{p}.bonus.regra": f"Efeito Condicional: {CONE[cone - 1][1]} (25.2).",
        f"{p}.tier": tr,
        f"{p}.tier.regra": (f"Tier novo: cada personagem recebe o {tr} nos slots que já possui (25.1)."
                            if ef in (1, 7, 13, 18) else "Mesmo Tier (viradas nos níveis 7, 13 e 18; 25.1 e 25.3)."),
        f"{p}.ressonancia": res[0] if res else "—",
        f"{p}.ressonancia.regra": (res[1] + ". Marco de história, escolhida na hora com o Mestre, e não muda (26.7)."
                                   if res else "Só nos níveis 5, 10, 15 e 20 (26.7)."),
        f"{p}.ritmo": "2 a 4 sessões por nível" if ef <= 8 else "4 a 6 sessões por nível",
    })
    lv = NV
    out["recompensas.aviso.nivel"] = ("" if lv is None or tx(lv) == "" else
                                      "Nível não é número: usando o da aba Campanha" if not num(lv) else
                                      "Nível fora de 1 a 20: usando o limite" if (lv < 1 or lv > 20 or int(lv) != lv)
                                      else "")
    marcos = ("Subiu de nível", "Entrou em faixa nova", "Fim de arco")
    out["recompensas.aviso.marco"] = ("Fora da lista: escolha o que aconteceu" if MC and MC not in marcos else
                                      f"O nível {ef} não é o primeiro de uma faixa (1, 5, 9, 13, 17)"
                                      if MC == "Entrou em faixa nova" and ef not in (1, 5, 9, 13, 17) else "")
    # Sobreposições (25.2) a partir da aba Grupo
    tot = 0
    for i in range(1, 7):
        nm = tx(E.get(f"grupo.pj{i}.nome"))
        c, so = E.get(f"grupo.pj{i}.cone"), E.get(f"grupo.pj{i}.sobrep")
        to = E.get(f"grupo.pj{i}.sobrep_total")
        teto = (3 - CONE[int(c) - 1][2]) if num(c) and 1 <= c <= 5 else ""
        q = f"recompensas.sob.{i}"
        # 25.2: o teto é do Cone e conta as de faixas anteriores; o total nunca é menor que as desta faixa
        total = max(to if num(to) else 0, so if num(so) else 0) if (nm or num(to) or num(so)) else ""
        out[f"{q}.total"] = total
        out[f"{q}.teto"] = teto
        out[f"{q}.situacao"] = ("" if not nm else "Preencha o Nível do Cone (1 a 5) na aba Grupo" if teto == "" else
                                "Cone no teto (+3): nenhuma Sobreposição (25.2)" if teto == 0 else
                                "Cone no teto (+3) com as que já tem: nenhuma a mais (25.2)" if total >= teto else
                                "Já recebeu a desta faixa (máximo 1 por faixa, 25.2)" if num(so) and so >= 1 else
                                "Pode receber 1 nesta faixa, se a história entregar (25.2)")
        out[f"{q}.aviso"] = (f"Acima do teto do Cone: o Nível {int(c)} aceita até {teto} (25.2)"
                             if teto != "" and total != "" and total > teto else
                             "Mais de 1 Sobreposição nesta faixa (25.2)" if num(so) and so > 1 else
                             "Total no Cone menor que as desta faixa: confira na aba Grupo"
                             if num(to) and num(so) and to < so else "")
        tot += so if num(so) else 0
    out["recompensas.sob.total"] = tot
    # preços (24.1–24.3)
    for i, (t, d, o, e, pr) in enumerate(ARMADURAS, start=1):
        out.update({f"recompensas.preco.armaduras.{i}.item": f"Armadura {t}",
                    f"recompensas.preco.armaduras.{i}.desc": f"Defesa +{d}" + ("" if o == "—" else f"; {o}"),
                    f"recompensas.preco.armaduras.{i}.espaco": e, f"recompensas.preco.armaduras.{i}.preco": pr})
    for i, (c, dd, al, at, e, pr) in enumerate(ARMAS, start=1):
        out.update({f"recompensas.preco.armas.{i}.item": f"Arma {c}", f"recompensas.preco.armas.{i}.desc": f"{dd}, {al}, {at}",
                    f"recompensas.preco.armas.{i}.espaco": e, f"recompensas.preco.armas.{i}.preco": pr})
    for i, (nm, cu, e, pr) in enumerate(POCOES, start=1):
        out.update({f"recompensas.preco.pocoes.{i}.item": f"Poção {nm}",
                    f"recompensas.preco.pocoes.{i}.desc": f"Cura {cu} (fixa)",
                    f"recompensas.preco.pocoes.{i}.espaco": e, f"recompensas.preco.pocoes.{i}.preco": pr})
    for i, (nm, e, pr, uso) in enumerate(ITENS, start=1):
        out.update({f"recompensas.preco.itens.{i}.item": nm, f"recompensas.preco.itens.{i}.desc": uso,
                    f"recompensas.preco.itens.{i}.espaco": e, f"recompensas.preco.itens.{i}.preco": pr})
    _achados(E, out)
    _cone(E, out)
    _conjunto(E, out)
    # Relíquias por Tier (25.3) e o Tier do grupo agora
    tg = tier(out["campanha.nivel_ef"])
    for i, (sl, da, *ts) in enumerate(RELIQUIAS, start=1):
        for j in range(4):
            out[f"recompensas.rel.{i}.t{j + 1}"] = f"+{ts[j]}"
        out[f"recompensas.rel.{i}.agora"] = f"+{ts[['Tier I', 'Tier II', 'Tier III', 'Tier IV'].index(tg)]}"
    out["recompensas.rel.tier_grupo"] = (f"Tier do grupo agora: {tg} (nível {out['campanha.nivel_ef']}). Bônus de slot "
                                         f"são permanentes e ficam fora do teto; Mãos e Esfera Planar nunca somam na "
                                         f"mesma rolagem (25.3).")
    saldo = 0
    for i in range(1, 26):
        v = E.get(f"recompensas.tes.{i}.cr")
        saldo += v if num(v) else 0
        out[f"recompensas.tes.{i}.aviso"] = ("Créditos não é número: a linha não entra no saldo"
                                             if v is not None and tx(v) != "" and not num(v) else "")
    out["recompensas.tes.saldo"] = saldo
    out["recompensas.tes.aviso.saldo"] = "Saldo de Créditos negativo" if saldo < 0 else ""


def _rolagem_av(v):
    return ("" if v is None or tx(v) == "" or (num(v) and 1 <= v <= 1000000 and int(v) == v)
            else "Rolagem nº fora de 1 a 1.000.000: usando 1")


def _achados(E, out):
    sem, rol = out["inicio.semente_ef"], S.rolagem_efetiva(E.get("recompensas.ach.rolagem"))
    x = lambda c: S.valor(sem, 400, rol, c)  # noqa: E731
    FA, LE = tx(E.get("recompensas.ach.faixa")), tx(E.get("recompensas.ach.leitura"))
    fx = efetiva_faixa(FA, out["campanha.faixa"])
    fi = FAIXAS.index(fx) + 1
    fator = 5 if LE == "Passagem" else 15 if LE == "Pesado" else 10
    verba = VERBA[fi - 1]
    cred = verba * fator // 100
    qtd = S.inteiro(x(1), 0, 2)
    cons = consumiveis(fi)
    j1 = S.escolha(x(2), len(cons))
    j2 = sem_repetir(len(cons), x(3), j1)
    it = [cons[j1 - 1] if qtd >= 1 else ("", ""), cons[j2 - 1] if qtd >= 2 else ("", "")]
    bug, pis = sorteia("bugiganga", x(4)), sorteia("pista", x(5))
    p = "recompensas.ach"
    out.update({
        f"{p}.creditos.valor": cred, f"{p}.qtd.qtd": qtd,
        f"{p}.creditos.origem": f"{fator}% da verba de marco da faixa {fx} (24.5) · {SUG} (H18)",
        f"{p}.item1": it[0][0], f"{p}.item1.valor": it[0][1], f"{p}.item1.qtd": 1 if it[0][0] else "",
        f"{p}.item2": it[1][0], f"{p}.item2.valor": it[1][1], f"{p}.item2.qtd": 1 if it[1][0] else "",
        f"{p}.bugiganga": bug, f"{p}.pista": pis,
        f"{p}.saida1.valor": cred, f"{p}.saida2.item": it[0][0], f"{p}.saida3.item": it[1][0],
        f"{p}.saida4.item": bug,
        f"{p}.saida2.notas": f"Preço de referência {it[0][1]} Cr (24.3)" if it[0][0] else "",
        f"{p}.saida3.notas": f"Preço de referência {it[1][1]} Cr (24.3)" if it[1][0] else "",
    })
    av = _rolagem_av(E.get("recompensas.ach.rolagem"))
    if not av and FA and FA not in FAIXAS:
        av = "Faixa fora da lista: usando a do grupo"
    if not av and LE and LE not in ("Passagem", "Típico", "Pesado"):
        av = "Leitura fora da lista: usando Típico"
    out[f"{p}.aviso"] = av
    ses = E.get("campanha.sessao")
    for k in range(1, 5):
        out[f"{p}.saida{k}.sessao"] = ses if num(ses) else ""


def _cone(E, out):
    sem, rol = out["inicio.semente_ef"], S.rolagem_efetiva(E.get("recompensas.cone.rolagem"))
    x = lambda c: S.valor(sem, 410, rol, c)  # noqa: E731
    FA = tx(E.get("recompensas.cone.faixa"))
    fx = efetiva_faixa(FA, out["campanha.faixa"])
    nv = FAIXAS.index(fx) + 1
    a, b = sorteia("cone_nome_a", x(1)), sorteia("cone_nome_b", x(2))
    nome = a if not b else f"{a} {b}"
    mem = sorteia("cone_memoria", x(3))
    bm = BONUS_MAIOR[S.escolha(x(4), 10) - 1]
    gat = lista("cone_gatilho")
    if gat:
        i1 = S.escolha(x(5), len(gat))
        g1, g2 = gat[i1 - 1], gat[sem_repetir(len(gat), x(6), i1) - 1]
    else:
        g1 = g2 = ""
    bonus, cond, numr, pv = CONE[nv - 1]
    teto = 3 - numr
    p = "recompensas.cone"
    out.update({f"{p}.nome": nome, f"{p}.nivel": f"Nível {nv}", f"{p}.memoria": mem, f"{p}.bonus": f"{bm}: {bonus}",
                f"{p}.efeito": (f"Dois efeitos: {g1}; {g2}" if nv >= 4 else g1) + f" — {cond}",
                f"{p}.sobreposicao": f"Até {teto} neste Nível (parte numérica até +3; PV até +{pv + 10 * teto}); no "
                                     f"máximo 1 por faixa",
                f"{p}.nivel.origem": f"Cone de Luz máximo da faixa {fx} (25.1)"})
    av = _rolagem_av(E.get("recompensas.cone.rolagem"))
    if not av and FA and FA not in FAIXAS:
        av = "Faixa fora da lista: usando a do grupo"
    out[f"{p}.aviso"] = av


def _conjunto(E, out):
    sem, rol = out["inicio.semente_ef"], S.rolagem_efetiva(E.get("recompensas.conj.rolagem"))
    x = lambda c: S.valor(sem, 420, rol, c)  # noqa: E731
    p = "recompensas.conj"
    out.update({f"{p}.nome": sorteia("conjunto_nome", x(1)), f"{p}.origem": sorteia("conjunto_origem", x(2)),
                f"{p}.p2": CONJUNTO_2[S.escolha(x(3), 4) - 1], f"{p}.p4": sorteia("conjunto_4", x(4)),
                f"{p}.aviso": _rolagem_av(E.get("recompensas.conj.rolagem"))})


def calcular(E, out):
    npcs(E, out)
    aventuras(E, out)
    recompensas(E, out)
