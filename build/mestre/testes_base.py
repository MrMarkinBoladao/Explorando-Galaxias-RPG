# -*- coding: utf-8 -*-
"""Suítes spike, protegidos, lint e texto da Planilha do Mestre (plano seção 8)."""

import json
import re
import shutil
import tempfile
from pathlib import Path

import testar_ficha as TF
import oraculo_mestre as O
from mestre import nucleo as N
from mestre import sorteio as S

Resultado = TF.Resultado
SPIKE_JSON = N.BUILD / "mestre_spike.json"


# ---------------------------------------------------------------------------
# origem (revisão da Fase 3, F2): as planilhas medidas são as que o gerador gravou
# ---------------------------------------------------------------------------

ORIGEM_JANELA_S = 900       # o gerador grava modelo, mapa e Exemplo em ≈ 3 min


def suite_origem(args):
    """Os dois .xlsx são os do gerador (openpyxl) e não uma regravação do Google Planilhas (modo Office, sincronizada
    pelo Drive): docProps/app.xml presente, sem xl/metadata.xml, área de impressão do Escudo gravada, o Exemplo com as
    mesmas validações do modelo e os dois gravados até ORIGEM_JANELA_S do mapa. A fase3 roda esta suíte no início e
    no fim, para que uma regravação no meio da bateria apareça como falha."""
    import zipfile
    r = Resultado("origem")
    t_mapa = N.SAIDA_MAPA.stat().st_mtime
    nval = {}
    for arq in (N.SAIDA_MODELO, N.SAIDA_EXEMPLO):
        if not arq.exists():
            r.falha(f"{arq.name}: não existe")
            continue
        with zipfile.ZipFile(arq) as z:
            nomes = set(z.namelist())
            r.ok("docProps/app.xml" in nomes, f"{arq.name}: sem docProps/app.xml (não é a gravação do gerador)")
            r.ok(not any(n.startswith("xl/metadata") for n in nomes),
                 f"{arq.name}: tem xl/metadata (formato de regravação do Google Planilhas)")
            wbx = z.read("xl/workbook.xml").decode("utf-8")
            r.ok(re.search(r"_xlnm\.Print_Area[^>]*>'?Escudo do Mestre'?!", wbx) is not None,
                 f"{arq.name}: Escudo do Mestre sem área de impressão")
            nval[arq] = sum(len(re.findall(r"<dataValidation[ >]", z.read(n).decode("utf-8")))
                            for n in nomes if n.startswith("xl/worksheets/sheet"))
        dt = arq.stat().st_mtime - t_mapa
        r.ok(abs(dt) <= ORIGEM_JANELA_S, f"{arq.name}: gravado {dt:+.0f} s do mapa (máximo ±{ORIGEM_JANELA_S} s): "
                                         f"o arquivo foi regravado fora do gerador?")
        r.info(f"{arq.name}: {nval[arq]} validações, {dt:+.0f} s do mapa")
    if len(nval) == 2:
        r.ok(nval[N.SAIDA_MODELO] == nval[N.SAIDA_EXEMPLO],
             f"validações: modelo {nval[N.SAIDA_MODELO]} × Exemplo {nval[N.SAIDA_EXEMPLO]}")
    return r


# ---------------------------------------------------------------------------
# spike (plano 1.1, design §12.1)
# ---------------------------------------------------------------------------

def suite_spike(args):
    import openpyxl
    from openpyxl.worksheet.datavalidation import DataValidation
    r = Resultado("spike")
    pasta = Path(tempfile.mkdtemp(prefix="mestre-spike-", dir=N.TEMP))
    casos = []          # (célula, esperado, descrição)
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Início"
        dd = wb.create_sheet("Dados")
        lin = [1]

        def caso(formula, esperado, desc):
            cel = f"A{lin[0]}"
            ws[cel] = formula
            casos.append((cel, esperado, desc))
            lin[0] += 1
            return cel

        # (1) MOD com 2^31·16807 e três MOD aninhados (L3)
        caso("=MOD(2147483646*16807,2147483647)", (2147483646 * 16807) % O.M, "MOD de 2^31·16807")
        for z in (1, 2, 65535, 65536, 123456789, 2147483646):
            caso("=" + S._l3(str(z)), O.l3(z), f"L3({z})")
            caso("=" + S._q(str(z)), O.q(z), f"Q({z})")
            caso("=" + S._x(str(z)), O.x_mix(z), f"X({z})")
        # (2) cadeia u, y, x nos extremos de S, G, R e C (semente e rolagem em células)
        ws["C1"], ws["C2"] = 1, 1
        k = 0
        for s in (1, 2147483646):
            for g in (100, 710):
                for rr in (1, 1000000):
                    for c in (1, 400):
                        row = 10 + k
                        ws[f"D{row}"], ws[f"E{row}"] = s, rr
                        ws[f"F{row}"] = S.formula_u(f"D{row}", f"E{row}", g, c)
                        ws[f"G{row}"] = S.formula_y(f"F{row}")
                        ws[f"H{row}"] = S.formula_x(f"G{row}")
                        u, y, x = O.uyx(s, g, rr, c)
                        casos += [(f"F{row}", u, f"u S={s} G={g} R={rr} C={c}"),
                                  (f"G{row}", y, f"y S={s} G={g} R={rr} C={c}"),
                                  (f"H{row}", x, f"x S={s} G={g} R={rr} C={c}")]
                        # (3) INT(x·n/M), dado e inteiro
                        ws[f"I{row}"] = "=" + S.escolha(f"H{row}", 100)
                        ws[f"J{row}"] = "=" + S.dado(f"H{row}", 6)
                        ws[f"K{row}"] = "=" + S.inteiro(f"H{row}", 0, 2)
                        casos += [(f"I{row}", O.escolha(x, 100), "INT(x·100/M)+1"),
                                  (f"J{row}", O.dado(x, 6), "d6"), (f"K{row}", O.inteiro(x, 0, 2), "inteiro 0–2")]
                        k += 1
        # (4) contador corrido + INDEX(MATCH(k)) com vaga vazia no meio e cabeçalho repetido
        valores = ["Alfa", None, "Beta", "#cab", "Gama", None, "Delta"]
        for i, v in enumerate(valores, start=1):
            if v and v != "#cab":
                dd[f"A{i + 1}"] = v
            elif v == "#cab":
                dd[f"A{i + 1}"] = "Lista (cabeçalho)"
            prev = "0" if i == 1 else f"B{i}"
            dd[f"B{i + 1}"] = f"={prev}" if v == "#cab" else f"={prev}+IF(LEN(A{i + 1})>0,1,0)"
        cheios = [v for v in valores if v and v != "#cab"]
        for kk in range(1, len(cheios) + 1):
            caso(f"=INDEX('Dados'!$A$2:$A$8,MATCH({kk},'Dados'!$B$2:$B$8,0))", cheios[kk - 1],
                 f"INDEX(MATCH({kk}, contador corrido))")
        # (5) SMALL + MATCH com chaves fracionárias e "" no meio
        chaves = [3.501, "", 1.002, 2.503, "", 2.004]
        for i, v in enumerate(chaves, start=1):
            dd[f"D{i}"] = f'=""' if v == "" else v
        nums = sorted(v for v in chaves if v != "")
        for kk in range(1, len(nums) + 1):
            pos = [v for v in chaves].index(nums[kk - 1]) + 1
            caso(f"=MATCH(SMALL('Dados'!$D$1:$D$6,{kk}),'Dados'!$D$1:$D$6,0)", pos, f"SMALL+MATCH k={kk}")
        caso("=SUMPRODUCT(ISNUMBER('Dados'!$D$1:$D$6)*1)", len(nums), "SUMPRODUCT(ISNUMBER)*1 (sem COUNT)")
        caso("=COUNTIF('Dados'!$D$1:$D$6,\"<\"&2.5)", sum(1 for v in nums if v < 2.5), "COUNTIF \"<\"& com \"\"")
        caso("=MIN('Dados'!$D$1:$D$6)", min(nums), "MIN ignora \"\"")
        # (6) CHOOSE(MATCH()) sobre 7 intervalos
        limites = [0, 10, 20, 30, 40, 50, 60]
        for i, v in enumerate(limites, start=1):
            dd[f"F{i}"] = v
        rotulos = ["a", "b", "c", "d", "e", "f", "g"]
        for v in (0, 9, 10, 35, 61, 59):
            idx = max(i for i, l in enumerate(limites) if l <= v)
            caso(f"=CHOOSE(MATCH({v},'Dados'!$F$1:$F$7,1)," + ",".join(f'"{x}"' for x in rotulos) + ")",
                 rotulos[idx], f"CHOOSE(MATCH({v}))")
        # (7) REPT com ● e ○; número concatenado; COUNTIFS
        caso('=REPT("●",3)&REPT("○",2)&" "&3&"/"&5', "●●●○○ 3/5", "REPT ●/○")
        caso('="média "&INT(3*(6+1)/2)', "média 10", "número concatenado")
        dd["H1"], dd["I1"], dd["H2"], dd["I2"] = "Fragmentum", "Ruína", "Legião", "Lua-forja"
        caso("=COUNTIFS('Dados'!$H$1:$H$2,\"Legião\",'Dados'!$I$1:$I$2,\"Lua-forja\")", 1, "COUNTIFS 2 critérios")
        # (7b) revisão 2 da ficha (Google real): ""+1 → #VALUE! (igual no Google); célula vazia: vazio no Google e 0
        # na formulas (diferente: a Mestre não depende disso, toda entrada passa pela leitura protegida, que devolve
        # "" explícito); ISBLANK("") = FALSE (igual). A camada e o contador conferidos aqui.
        caso('=""+1', "#VALUE!", 'revisão 2: ""+1 dá #VALUE!')
        caso("='Dados'!$Z$1", 0, "revisão 2: célula vazia = 0 na formulas (no Google é vazio)")
        caso('=ISBLANK("")', False, 'revisão 2: ISBLANK("") é FALSE')
        caso("=IF(ISERROR('Dados'!$Z$1),\"\",IF(ISBLANK('Dados'!$Z$1),\"\",'Dados'!$Z$1))", "",
             "revisão 2: leitura protegida de célula vazia = \"\" (igual ao Google)")
        dd["Z2"] = "=1/0"
        caso("=IF(ISERROR('Dados'!$Z$2),\"\",IF(ISBLANK('Dados'!$Z$2),\"\",'Dados'!$Z$2))", "",
             "revisão 2: leitura protegida de célula com erro = \"\"")
        caso("=SUMPRODUCT(ISERROR('Dados'!$Z$1:$Z$3)*1)", 1, "revisão 2: contador SUMPRODUCT(ISERROR)*1 = 1")
        # (8) validação de lista apontando para outra aba (com acento no nome)
        dv = DataValidation(type="list", formula1="'Dados'!$A$2:$A$8", allow_blank=True, errorStyle="warning")
        dv.add("B1")
        ws.add_data_validation(dv)
        arq = pasta / "mestre-spike.xlsx"
        wb.save(arq)
        modelo = TF.Modelo(arq)
        sol = modelo.calcular(saidas=[f"'Início'!{c}" for c, _, _ in casos])
        divergencias = []
        for cel, esperado, desc in casos:
            obtido = sol.get(f"'Início'!{cel}")
            ok = TF._igual(obtido, esperado) if not isinstance(esperado, float) else \
                isinstance(obtido, (int, float)) and abs(obtido - esperado) < 1e-9
            if not r.ok(ok, f"{desc}: formulas={obtido!r} Python={esperado!r}"):
                divergencias.append({"caso": desc, "formulas": obtido, "python": esperado})
        wb2 = openpyxl.load_workbook(arq)
        dvs = wb2["Início"].data_validations.dataValidation
        r.ok(len(dvs) == 1 and dvs[0].formula1 == "'Dados'!$A$2:$A$8",
             "validação de lista apontando para outra aba não voltou igual ao reabrir")
        r.info(f"{len(casos)} casos; carga do modelo {modelo.tempo_carga:.2f} s; maior intermediário do "
               f"oráculo {O.maior_intermediario():.3e} (< 2^53 = {2**53:.3e})")
        SPIKE_JSON.write_text(json.dumps({
            "gerado_por": "build/testar_mestre.py --suite spike", "biblioteca": "formulas 1.3.4",
            "casos": len(casos), "divergencias": divergencias,
            "maior_intermediario": O.maior_intermediario(),
            "padroes": ["MOD 2^31·16807", "L3", "Q", "X", "u/y/x nos extremos", "INT(x·n/M)",
                        "INDEX(MATCH(k, contador corrido))", "SMALL+MATCH chaves fracionárias",
                        "SUMPRODUCT(ISNUMBER()*1)", "COUNTIF \"<\"&", "MIN com \"\"", "CHOOSE(MATCH()) 7 intervalos",
                        "REPT ●○", "COUNTIFS", "validação de lista em outra aba"],
        }, ensure_ascii=False, indent=2), encoding="utf-8")
    finally:
        shutil.rmtree(pasta, ignore_errors=True)
    return r


# ---------------------------------------------------------------------------
# protegidos — só leitura (plano P4, P5)
# ---------------------------------------------------------------------------

BASELINE = N.RAIZ / ".agents" / "tasks" / "mestre" / "baseline-hashes.json"
# P5: ausente já antes da Fase 1; a Mestre nunca grava em ficha-automatizada\
AUSENTES_CONHECIDOS = {"ficha-automatizada/Cópia de Ficha Automatizada - Explorando Galáxias V1.1.xlsx":
                       "ausente já antes da Fase 1; a Mestre nunca grava em ficha-automatizada\\"}


def suite_protegidos(args):
    r = Resultado("protegidos")
    gravados = json.loads(TF.PROTEGIDOS_JSON.read_text(encoding="utf-8"))["sha256"]
    atuais = {p.relative_to(N.RAIZ).as_posix(): p for p in TF._arquivos_protegidos() if p.exists()}
    for rel, h in gravados.items():
        if rel.lower().endswith("desktop.ini"):
            # metadado do Google Drive (o Drive cria, troca e apaga sozinho; plano §1): não é conteúdo do livro
            r.info(f"ignorado (metadado do Drive): {rel} — {'presente' if rel in atuais else 'ausente'}")
            continue
        if r.ok(rel in atuais, f"livro/protegido sumiu: {rel}"):
            r.ok(TF._sha256(atuais[rel]) == h, f"livro/protegido ALTERADO: {rel}")
    for rel in atuais:
        r.ok(rel in gravados or rel.lower().endswith("desktop.ini"), f"arquivo novo em área protegida (livro): {rel}")
    base = json.loads(BASELINE.read_text(encoding="utf-8"))
    hashes = base.get("sha256", base) if isinstance(base, dict) else base
    n = 0
    for rel, h in hashes.items():
        if not isinstance(h, str) or len(h) != 64:
            continue
        p = Path(rel) if Path(rel).is_absolute() else N.RAIZ / rel
        try:
            chave = p.resolve().relative_to(N.RAIZ.resolve()).as_posix()
        except ValueError:
            chave = rel.replace("\\", "/")
        if not p.exists():
            if chave in AUSENTES_CONHECIDOS:
                r.info(f"ausente conhecido (P5): {chave} — {AUSENTES_CONHECIDOS[chave]}")
            else:
                r.falha(f"arquivo da ficha ausente: {rel}")
            continue
        n += 1
        r.ok(TF._sha256(p) == h, f"arquivo da ficha ALTERADO: {rel}")
    r.info(f"{len(gravados)} arquivos do livro e {n} da ficha conferidos (SHA-256)")
    return r


# ---------------------------------------------------------------------------
# lint — compatibilidade com o Google Planilhas (design §12.1 + P3)
# ---------------------------------------------------------------------------

PROIBIDAS_EXTRA = {"RAND", "RANDBETWEEN", "SUBSTITUTE", "CHAR", "TRIM", "COUNT", "INDIRECT", "OFFSET", "TEXT",
                   "FILTER", "SORT", "UNIQUE", "SEQUENCE", "XLOOKUP", "LET", "LAMBDA", "TEXTJOIN", "IFS", "SWITCH",
                   "ROUNDDOWN", "ROUND", "ROUNDUP"}     # D1: arredondamento só com INT (revisão Fase 1, F7)


def suite_lint(args):
    import openpyxl
    import renderizar_ficha as R
    from openpyxl.utils import get_column_letter
    r = Resultado("lint")
    do_json = set(TF.ler_funcoes_ok()["funcoes"])
    permitidas = do_json - PROIBIDAS_EXTRA
    r.ok("TRIM" not in permitidas and "COUNT" not in permitidas, "lista efetiva contém TRIM/COUNT")
    for ruim in ("=RAND()", "=COUNT(A1:A3)", "=INDIRECT(A1)", "=TRIM(A1)", "=SUBSTITUTE(A1,\"a\",\"b\")",
                 "=_xlfn.XLOOKUP(1,A1:A2,B1:B2)", "=SUM(A1;A2)", "=ROUND(A1,0)"):
        sonda = Resultado("sonda")
        TF._checar_formula(sonda, ruim, "sonda", permitidas)
        r.ok(bool(sonda.falhas), f"autoteste do lint não pegou {ruim}")
    mapa = json.loads(N.SAIDA_MAPA.read_text(encoding="utf-8"))
    for arq in (N.SAIDA_MODELO, N.SAIDA_EXEMPLO):
        r.ok(arq.exists(), f"entregável ausente: {arq.name}")
    wb = openpyxl.load_workbook(N.SAIDA_MODELO)
    r.ok(wb.sheetnames == N.ABAS, f"abas fora de D3: {wb.sheetnames}")
    r.ok(len(wb.defined_names) == 0, "nomes definidos no arquivo")
    r.ok(bool(wb.calculation.fullCalcOnLoad), "fullCalcOnLoad desligado")
    for f in wb._fonts:
        r.ok(f.name == "Arial", f"fonte {f.name}")
    nf = ndv = ncf = 0
    for ws in wb.worksheets:
        r.ok(not ws.protection.sheet, f"{ws.title}: proteção")
        r.ok(not ws.tables, f"{ws.title}: Tabela do Excel")
        r.ok(ws.freeze_panes is None and ws.sheet_view.pane is None, f"{ws.title}: painel congelado")
        for linha in ws.iter_rows():
            for c in linha:
                v = c.value
                if isinstance(v, bool):
                    r.falha(f"{ws.title}!{c.coordinate}: booleano (caixa de seleção?)")
                if isinstance(v, str) and v.startswith("="):
                    nf += 1
                    onde = f"'{ws.title}'!{c.coordinate}"
                    r.ok("_xlfn" not in v.lower() and "«" not in v, f"{onde}: _xlfn ou marcador não resolvido")
                    TF._checar_formula(r, v, onde, permitidas)
                if c.has_style and c.font is not None and c.font.name != "Arial":
                    r.falha(f"{ws.title}!{c.coordinate}: fonte {c.font.name}")
        for dv in ws.data_validations.dataValidation:
            ndv += 1
            for fx in (dv.formula1, dv.formula2):
                if fx and dv.type == "custom":
                    TF._checar_formula(r, fx, f"{ws.title} validação", permitidas, permite_exclamacao=False)
                elif fx and dv.type == "list":
                    r.ok("[" not in fx and "«" not in fx, f"{ws.title} validação {dv.sqref}: {fx}")
                    r.ok(dv.errorStyle == "warning", f"{ws.title} validação {dv.sqref}: errorStyle {dv.errorStyle}")
        for cf in ws.conditional_formatting:
            for regra in cf.rules:
                ncf += 1
                for fx in regra.formula or []:
                    TF._checar_formula(r, fx, f"{ws.title} CF {cf.sqref}", permitidas, permite_exclamacao=False)
        # D11: área ≤ 1360 px em A:L; nada visível à direita de L (exceto Dados e Tabelas); auxiliares ocultas
        if ws.title not in ("Dados", "Tabelas"):
            larg = sum(R.px_coluna(ws, k) for k in range(1, 13))
            r.ok(larg <= N.LARGURA_TELA, f"{ws.title}: A:L = {larg} px (> 1360)")
            fora = [c.coordinate for linha in ws.iter_rows(min_col=13) for c in linha
                    if c.value is not None and not ws.column_dimensions[get_column_letter(c.column)].hidden]
            r.ok(not fora, f"{ws.title}: conteúdo visível à direita de L: {fora[:6]}")
            for t in R.tabelas(ws):
                r.ok(len(t["linhas"]) <= 13, f"{ws.title}: tabela com cabeçalho na linha {t['cab']} tem "
                                             f"{len(t['linhas'])} linhas (máximo 13, D11)")
        r.ok(ws.title in N.ABAS_PRONTAS, f"{ws.title}: aba fora das prontas (Fase 3: as 18 estão prontas)")
    # toda entrada vazia no modelo, nunca em coluna oculta nem à direita de L
    for nome, info in mapa["entradas"].items():
        aba, cel = TF.separar_ref(mapa["celulas"][nome])
        ws = wb[aba]
        r.ok(ws[cel].value is None, f"entrada {nome} ({aba}!{cel}) não está vazia no modelo")
        col = ws[cel].column_letter
        r.ok(not ws.column_dimensions[col].hidden, f"entrada {nome} em coluna oculta")
    r.info(f"{nf} fórmulas, {ndv} validações, {ncf} formatações condicionais, {len(mapa['entradas'])} entradas; "
           f"funções permitidas: {len(permitidas)} das {len(do_json)} de ficha_funcoes_ok.json "
           f"(sem {', '.join(sorted(do_json & PROIBIDAS_EXTRA))})")
    lint_google(r, wb, mapa)
    from mestre import testes_fase3
    testes_fase3.lint_fase3(r, wb)     # Fase 3: nenhuma aba "Em construção"; Escudo em A4 paisagem
    return r


def lint_google(r, wb, mapa):
    """Requisito do Google (revisão 2 da ficha): testar_ficha._lint_google por import, com o mapa da Mestre (as vagas
    das listas editáveis da aba Tabelas contam como entradas). (1) nenhuma opção de lista que o Google leia como
    fórmula, data, hora, número ou booleano; (2) nenhuma fórmula lê uma entrada fora da camada de leitura protegida;
    toda entrada com sinal de erro e todo sinal num aviso; contador presente."""
    m = json.loads(json.dumps(mapa))
    for id_, t in m["tabelas"].items():
        for k, cel in enumerate(t["vagas"] + t.get("ambientes", []) + t.get("segunda", [])):
            nome = f"{id_}.vaga_lint{k + 1}"
            m["celulas"][nome] = f"'Tabelas'!{cel}"
            m["entradas"][nome] = {"tipo": "texto"}
    antes = TF.ler_mapa
    TF.ler_mapa = lambda: m
    try:
        n0 = len(r.falhas)
        TF._lint_google(r, wb)
        r.info(f"requisito do Google: {len(r.falhas) - n0} achado(s) (testar_ficha._lint_google com o mapa da Mestre)")
    finally:
        TF.ler_mapa = antes
    # fora da camada, ISBLANK sempre erra: a entrada chega como "" e ISBLANK("") = FALSE (spike, revisão 2)
    camada = set(mapa.get("leitura", {}))
    isb = [f"'{ws.title}'!{c.coordinate}" for ws in wb.worksheets for ln in ws.iter_rows() for c in ln
           if isinstance(c.value, str) and "ISBLANK(" in c.value and f"'{ws.title}'!{c.coordinate}" not in camada]
    r.ok(not isb, f"ISBLANK fora da camada de leitura protegida (use LEN(x)=0): {isb[:8]}")
    ini = wb["Início"]
    dica = ini[mapa["celulas"]["inicio.dica_formula"].split("!")[1]].value or ""
    r.ok(dica.startswith("Não comece um texto com +, - ou ="), f"Início sem a dica do Google: {dica!r}")
    r.ok(all(f"inicio.erros.{N.SLUG[a]}" in mapa["celulas"] for a in N.ABAS) and "inicio.erros.total" in mapa["celulas"],
         "Início sem o contador de erros por aba")


# ---------------------------------------------------------------------------
# texto — ortografia PT-BR, glossário 30.1, termos de 30.2, rótulos de Sugestão
# ---------------------------------------------------------------------------

LEXICO_MESTRE = N.BUILD / "mestre_lexico_extra.txt"


def suite_texto(args):
    r = Resultado("texto")
    itens = []
    for arq in (N.SAIDA_MODELO, N.SAIDA_EXEMPLO):
        itens += [(f"{arq.stem}: {o}", s) for o, s in TF._textos_visiveis(arq)]
    from mestre import testes_fase4
    guia = testes_fase4.textos_guia()          # Fase 4: o COMO-USAR (ortografia, 30.1, 30.2, léxico)
    itens += guia
    r.info(f"COMO-USAR: {len(guia)} linhas de texto conferidas")
    proib = TF._proibidas_ps1()
    pad = [(t, re.compile(r"\b" + re.escape(t) + r"\b", re.I)) for t in proib]
    pad += [(t, re.compile(r"(?<!\w)" + re.escape(t) + r"(?!\w)", re.I)) for t in TF._APOSENTADOS_30_2]
    for onde, s in itens:
        for t, rx in pad:
            r.ok(not rx.search(s), f"(1) termo proibido/aposentado '{t}' em {onde}: {s[:100]!r}")
    corpo = "\n".join(re.sub(r"\*\*|\*|`", "", p.read_text(encoding="utf-8")) for p in sorted(TF.LIVRO.glob("*.md")))
    for termo in TF._glossario_30_1():
        ch = TF._sem_acento_mesmo_tamanho(termo)
        rx = re.compile(r"(?<!\w)" + re.escape(ch) + r"(?!\w)")
        pal = termo.split()
        for onde, s in itens:
            for m in rx.finditer(TF._sem_acento_mesmo_tamanho(s)):
                orig = s[m.start():m.end()]
                errado = orig.lower() != termo.lower() or (len(pal) > 1 and orig.split()[1:] != pal[1:])
                r.ok(not errado or s.strip() in corpo, f"(2) '{termo}' escrito '{orig}' em {onde}: {s[:100]!r}")
    livro, _ = TF._lexico()
    extra = {l.strip().lower() for l in LEXICO_MESTRE.read_text(encoding="utf-8").splitlines()
             if l.strip() and not l.startswith("#")}
    r.ok(not (extra & livro), f"mestre_lexico_extra.txt repete palavras do livro: {sorted(extra & livro)[:10]}")
    fora, usadas = {}, set()
    # listas de nomes (sem_ortografia, design §6.16) e os nomes do Elenco do Exemplo ficam fora da ortografia (3)
    mp = json.loads(N.SAIDA_MAPA.read_text(encoding="utf-8"))
    sem_orto = {f"'Tabelas'!{c}" for t in mp["tabelas"].values() if t.get("sem_ortografia") for c in t["vagas"]}
    sem_orto |= {v for k, v in mp["celulas"].items() if re.fullmatch(r"npcs\.(elenco\.\d+\.nome|cartao\d\.npc)", k)}
    for onde, s in itens:
        if onde.split(": ", 1)[-1] in sem_orto:
            continue
        for w in TF._RE_PALAVRA.findall(s):
            r.checagens += 1
            if w.lower() in extra:
                usadas.add(w.lower())
            if w.lower() not in livro and w.lower() not in extra:
                fora.setdefault(w, onde)
    for w, onde in sorted(fora.items()):
        r.falha(f"(3) palavra fora do léxico: '{w}' (em {onde})")
    r.ok(not (extra - usadas), f"léxico extra com palavras sem uso: {sorted(extra - usadas)[:10]}")
    for forma in TF._SEM_ACENTO:
        rx = re.compile(r"(?<!\w)" + re.escape(forma) + r"(?!\w)", re.I)
        for onde, s in itens:
            r.ok(not rx.search(s), f"(4) forma sem acento '{forma}' em {onde}: {s[:80]!r}")
    # ● e ○ são a barra dos relógios (design H16, validada no spike); não são emoji
    for onde, s in itens:
        r.ok(not TF._tem_emoji(s.replace("●", "").replace("○", "")), f"emoji em {onde}")
    # toda célula de heurística tem o rótulo "Sugestão da planilha — não é regra do livro (H#)"
    import openpyxl
    wb = openpyxl.load_workbook(N.SAIDA_MODELO)
    mapa = json.loads(N.SAIDA_MAPA.read_text(encoding="utf-8"))
    hs = set()
    for s in mapa["sugestoes"]:
        aba, cel = TF.separar_ref(s["rotulo"])
        v = str(wb[aba][cel].value or "")
        r.ok(N.ROTULO_SUGESTAO in v and (f"({s['h']}" in v or f", {s['h']})" in v or f"{s['h']})" in v),
             f"rótulo de Sugestão ausente em {s['rotulo']} ({s['h']}): {v[:80]!r}")
        hs.add(s["h"])
    r.info(f"{len(itens)} trechos de texto visível; léxico do livro {len(livro)} + {len(extra)} extras; "
           f"heurísticas rotuladas: {', '.join(sorted(hs, key=lambda h: int(h[1:])))}")
    from mestre import testes_fase3
    testes_fase3.texto_fase3(r)        # Fase 3: H24 (seção citada em cada item da Sessão Zero), custos da falha
    return r
