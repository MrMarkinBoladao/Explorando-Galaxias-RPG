# -*- coding: utf-8 -*-
"""
nucleo.py — infraestrutura da Planilha do Mestre (design §3.1, D2, D4, D9, D11).

Caminhos e nomes de saída (constantes únicas), mapa lógico próprio (MapaMestre → build\\mestre_mapa.json),
registro de referências «nome.logico» (independente do REFS/MAPA de gerar_ficha), os construtores de
célula (ent/cal/av/aux/rot/cab/sugestao), a aba Dados e o ajuste de layout pela régua
renderizar_ficha.medir. Os ESTILOS vêm de build\\gerar_ficha.py por import (a ficha não é editada).
"""

import json
import math
import os
import re
import sys
import tempfile
import unicodedata
from copy import copy
from pathlib import Path

BUILD = Path(__file__).resolve().parent.parent
RAIZ = BUILD.parent
if str(BUILD) not in sys.path:
    sys.path.insert(0, str(BUILD))

import gerar_ficha as G  # noqa: E402  (estilos, cores, legenda, q, _abs, _px_largura, marcar_tabelas)
import renderizar_ficha as R  # noqa: E402  (régua de medida)
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter, column_index_from_string  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

TEMP = os.environ.get("TEMP") or tempfile.gettempdir()
LIVRO = RAIZ / "livro-v1.0"
PASTA_MESTRE = RAIZ / "Mestre"
SAIDA_MODELO = PASTA_MESTRE / "Planilha do Mestre - Explorando Galáxias V1.2.xlsx"
SAIDA_EXEMPLO = PASTA_MESTRE / "Planilha do Mestre - Exemplo.xlsx"
SAIDA_GUIA = PASTA_MESTRE / "COMO-USAR-PLANILHA-DO-MESTRE.md"
SAIDA_MAPA = BUILD / "mestre_mapa.json"
VERSAO = G.VERSAO

ABAS = ["Início", "Campanha", "Grupo", "Sessões", "Missões", "NPCs", "Inimigos", "Bestiário", "Encontros",
        "Combate", "Aventuras", "Recompensas", "Mundos", "Improviso", "Escudo do Mestre", "Minhas Tabelas",
        "Tabelas", "Dados"]
SLUG = {"Início": "inicio", "Campanha": "campanha", "Grupo": "grupo", "Sessões": "sessoes", "Missões": "missoes",
        "NPCs": "npcs", "Inimigos": "inimigos", "Bestiário": "bestiario", "Encontros": "encontros",
        "Combate": "combate", "Aventuras": "aventuras", "Recompensas": "recompensas", "Mundos": "mundos",
        "Improviso": "improviso", "Escudo do Mestre": "escudo", "Minhas Tabelas": "minhas", "Tabelas": "tabelas",
        "Dados": "dados"}
ABAS_FASE1 = ["Início", "Campanha", "Grupo", "Inimigos", "Bestiário", "Encontros", "Combate", "Tabelas", "Dados"]
ABAS_FASE2 = ["NPCs", "Aventuras", "Recompensas"]
ABAS_FASE3 = ["Sessões", "Missões", "Mundos", "Improviso", "Escudo do Mestre", "Minhas Tabelas"]
ABAS_PRONTAS = [a for a in ABAS if a in ABAS_FASE1 + ABAS_FASE2 + ABAS_FASE3]     # Fase 3: as 18 (ordem de D3)
EM_CONSTRUCAO = "Em construção: esta aba chega numa próxima fase da Planilha do Mestre."
ROTULO_SUGESTAO = "Sugestão da planilha — não é regra do livro"

SUBTITULOS = {
    "Início": "Comece por aqui: semente da campanha, índice das abas e como a planilha se organiza.",
    "Campanha": "A mesa, os marcos, as facções e a reputação, os relógios, a linha do tempo, a Ficha de Decisões e a "
                "Sessão Zero.",
    "Grupo": "Os personagens jogadores, copiados da ficha de cada jogador (uma linha por PJ).",
    "Sessões": "Preparar a próxima sessão (cenas, pistas, ganchos, checklist) e o diário de campanha.",
    "Missões": "As missões oferecidas, ativas e concluídas (cole a linha de saída da aba Aventuras).",
    "NPCs": "Gerador de NPCs (nome por cultura, Raça, Caminho, ocupação, voz, segredo, gancho), elenco e cartões.",
    "Inimigos": "Criador de Inimigos: uma linha por inimigo da campanha, com a ficha de 28.1 pronta.",
    "Bestiário": "As 32 fichas do capítulo 28, com filtro e ficha completa.",
    "Encontros": "Orçamento de 27.4, três encontros salvos, contrato da Fraqueza e encontro aleatório.",
    "Combate": "Fila de Ação, PV, Tenacidade e Quebra, condições, Morrendo e recursos do grupo.",
    "Aventuras": "Gerador de aventuras: tipo, gancho, contratante, objetivo, local, antagonista, complicação, "
                 "reviravolta, prazo, recompensa e cinco cenas.",
    "Recompensas": "Entrega de marco (24.5, 25.1, 26.7), achados de encontro, sabor de Cone e de Conjunto e o "
                   "tesouro do grupo.",
    "Mundos": "Geradores de planeta ou local, estação, nave, facção, organização e nomes avulsos; as facções e os "
              "locais do livro.",
    "Improviso": "Rumores e ganchos, eventos e complicações, loja com preços do livro, bugigangas, oráculo sim/não, "
                 "rolador de dados e DT rápida.",
    "Escudo do Mestre": "Referência de mesa para imprimir (A4 paisagem, 5 páginas): todo número vem da aba Dados.",
    "Minhas Tabelas": "Dez tabelas do próprio Mestre, com 100 vagas e sorteio de 1 a 5 resultados sem repetir.",
    "Tabelas": "Listas dos geradores, editáveis: 100 vagas por lista; apague para trocar, use as vagas vazias "
               "para ampliar (não insira linhas).",
    "Dados": "Tabelas do livro usadas pelas fórmulas. Não edite: o gerador reescreve esta aba.",
}

# Grade padrão de A:L (D11): A = 150 px, B…J = 100 px, K e L = 155 px → 1360 px
GRADE_PX = {"A": 150, **{get_column_letter(c): 100 for c in range(2, 11)}, "K": 155, "L": 155}
LARGURA_TELA = 1360
PRIMEIRA_AUX = 14        # N: colunas auxiliares (ocultas) começam aqui; M fica vazia e oculta


# ---------------------------------------------------------------------------
# Mapa lógico (contrato gerador × testes, design §4.3)
# ---------------------------------------------------------------------------

class MapaMestre:
    def __init__(self):
        self.celulas = {}
        self.blocos = {}
        self.listas = {}
        self.tabelas = {}
        self.geradores = {}
        self.subtabelas = []
        self.entradas = {}
        self.avisos = {}
        self.sugestoes = []
        self.constantes = {}
        self.numeros_de_regra = {}
        self.ocultas = {}
        self.pior = {}           # "'Aba'!A1" -> texto de pior caso (medida do layout)
        self.extra = {}          # Fase 3: "minhas" (vagas das Minhas Tabelas), "escudo" (célula -> bloco da aba Dados)

    @staticmethod
    def ref(aba, cel):
        return f"'{aba}'!{cel}"

    def salvar(self):
        dados = {
            "gerado_por": "build/gerar_mestre.py", "arquivo": SAIDA_MODELO.name, "abas": ABAS,
            "convencao": "celulas: '<aba>.<campo>' -> \"'Aba'!A1\" (aba em minúsculas sem acento); blocos: "
                         "'dados.<id>'; listas: 'lista.<id>'; tabelas: 'tab.<id>'; geradores: '<G>' -> campos "
                         "com as células u/y/x; numeros_de_regra: células calculadas com número de regra "
                         "(cobertura do oráculo, plano P10)",
            "celulas": self.celulas, "blocos": self.blocos, "listas": self.listas, "tabelas": self.tabelas,
            "geradores": self.geradores, "subtabelas": self.subtabelas, "entradas": self.entradas,
            "avisos": self.avisos, "sugestoes": self.sugestoes, "constantes": self.constantes,
            "numeros_de_regra": self.numeros_de_regra,
            "ocultas": {k: sorted(v, key=column_index_from_string) for k, v in self.ocultas.items()},
            **self.extra,
        }
        # requisito do Google (revisão 2 da ficha): camada de leitura protegida, sinais de erro e contador
        from mestre import protecao
        dados.update({"leitura": protecao.LEITURA, "sinais": protecao.SINAIS, "contador": protecao.CONTADOR})
        SAIDA_MAPA.write_text(json.dumps(dados, ensure_ascii=False, indent=1), encoding="utf-8")


MAPA = MapaMestre()
REFS = {}
_MARCADOR = re.compile(r"«([^»]+)»")


def iniciar():
    global MAPA, REFS
    MAPA = MapaMestre()
    REFS = {}


def T(nome):
    return f"«{nome}»"


def q(s):
    return G.q(s)


def slug(s):
    s = unicodedata.normalize("NFKD", str(s))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")


def _abs(cel):
    return G._abs(cel)


def reg(nome, ws, cel):
    if nome in REFS:
        raise ValueError(f"nome lógico repetido no mapa: {nome}")
    REFS[nome] = (ws.title, cel)
    MAPA.celulas[nome] = MapaMestre.ref(ws.title, cel)
    return cel


def ref(nome, aba):
    a, cel = REFS[nome]
    return _abs(cel) if a == aba else f"'{a}'!{_abs(cel)}"


def refx(nome):
    """Referência absoluta com o nome da aba, para usar em validação de lista."""
    a, cel = REFS[nome]
    return f"'{a}'!{_abs(cel)}"


def resolver_marcadores(wb):
    faltando = set()

    def troca(m, aba):
        if m.group(1) not in REFS:
            faltando.add(m.group(1))
            return "#REF!"
        return ref(m.group(1), aba)

    for ws in wb.worksheets:
        for linha in ws.iter_rows():
            for c in linha:
                if isinstance(c.value, str) and "«" in c.value:
                    c.value = _MARCADOR.sub(lambda m, a=ws.title: troca(m, a), c.value)
        for dv in ws.data_validations.dataValidation:
            if dv.formula1 and "«" in dv.formula1:
                dv.formula1 = _MARCADOR.sub(lambda m: refx(m.group(1)) if m.group(1) in REFS
                                            else faltando.add(m.group(1)) or "#REF!", dv.formula1)
    if faltando:
        raise ValueError(f"marcadores sem célula no mapa: {sorted(faltando)}")


def _col(letra):
    return column_index_from_string(letra)


def _mesclar(ws, cel, ate):
    if ate:
        linha = ws[cel].row
        ws.merge_cells(f"{cel}:{ate}{linha}")


# ---------------------------------------------------------------------------
# Validações (errorStyle="warning": deixa digitar e explica em PT-BR)
# ---------------------------------------------------------------------------

def _validacao(ws, cel, dv):
    dv.add(cel)
    ws.add_data_validation(dv)


def dv_lista(ws, cel, fonte, rotulo):
    _validacao(ws, cel, DataValidation(
        type="list", formula1=fonte, allow_blank=True, showErrorMessage=True, errorStyle="warning",
        errorTitle="Fora da lista",
        error=f"{rotulo}: escolha um valor da lista. A planilha pode não reconhecer outro valor."))


def dv_inteiro(ws, cel, minimo, maximo, rotulo):
    _validacao(ws, cel, DataValidation(
        type="whole", operator="between", formula1=str(minimo), formula2=str(maximo), allow_blank=True,
        showErrorMessage=True, errorStyle="warning", errorTitle="Valor fora da faixa",
        error=f"{rotulo}: use um número inteiro de {minimo} a {maximo}."))


def dv_texto(ws, cel, maximo, rotulo):
    _validacao(ws, cel, DataValidation(
        type="textLength", operator="lessThanOrEqual", formula1=str(maximo), allow_blank=True,
        showErrorMessage=True, errorStyle="warning", errorTitle="Texto longo",
        error=f"{rotulo}: até {maximo} caracteres."))


# ---------------------------------------------------------------------------
# Construtores de célula
# ---------------------------------------------------------------------------

def _alinhar(c, quebra=True, centro=False):
    c.alignment = Alignment(vertical="center", wrap_text=quebra, horizontal="center" if centro else None)


def ent(ws, cel, nome, tipo="texto", fonte=None, minimo=None, maximo=None, amostra=None, invalido=None,
        ate=None, rotulo="", centro=False, opcoes=None):
    """Entrada: sai VAZIA no arquivo (D4, D9). tipo: texto | inteiro | lista.
    fonte (lista): 'lista.<id>' (aba Dados), marcador «nome» de intervalo, ou literal '"a,b"'.
    `opcoes` = valores possíveis da lista (amostra e pior caso)."""
    if _col(re.match(r"[A-Z]+", cel).group(0)) > 12 and ws.title not in ("Tabelas",):
        raise ValueError(f"entrada {nome} em coluna auxiliar ({ws.title}!{cel})")
    c = G.entrada(ws, cel)
    _alinhar(c, quebra=(tipo in ("texto", "lista")), centro=centro)
    _mesclar(ws, cel, ate)
    reg(nome, ws, cel)
    rotulo = rotulo or nome
    info = {"tipo": tipo}
    if tipo == "lista":
        if fonte.startswith("lista."):
            formula = MAPA.listas[fonte]
        else:
            formula = fonte
        info["fonte"] = fonte
        dv_lista(ws, cel, formula, rotulo)
        if opcoes is None and fonte.startswith('"'):
            opcoes = [x for x in fonte.strip('"').split(",") if x]
        if opcoes is None and fonte.startswith("lista."):
            opcoes = list(MAPA_LISTAS_VALORES.get(fonte, []))
        info["opcoes"] = list(opcoes or [])
        if amostra is None and info["opcoes"]:
            amostra = info["opcoes"][0]
        if invalido is None:
            invalido = "Valor inventado"
    elif tipo == "inteiro":
        dv_inteiro(ws, cel, minimo, maximo, rotulo)
        info["minimo"], info["maximo"] = minimo, maximo
        if amostra is None:
            amostra = maximo
        if invalido is None:
            invalido = "abc"
    else:
        if maximo:
            dv_texto(ws, cel, maximo, rotulo)
            info["maximo"] = maximo
        if amostra is None:
            amostra = "Teste"
    info["amostra"] = amostra
    info["invalido"] = invalido
    MAPA.entradas[nome] = info
    return cel


MAPA_LISTAS_VALORES = {}     # 'lista.<id>' -> valores (preenchido por escrever_dados)


def vaga(ws, cel, nome, maximo=200, ate=None):
    """Entrada de texto livre SEM validação própria (as vagas das Minhas Tabelas: uma validação para a coluna toda,
    posta pelo chamador). Fica no mapa como entrada e passa pela camada de leitura protegida (requisito do Google)."""
    c = G.entrada(ws, cel)
    _alinhar(c, quebra=True)
    _mesclar(ws, cel, ate)
    reg(nome, ws, cel)
    MAPA.entradas[nome] = {"tipo": "texto", "maximo": maximo, "amostra": "Teste", "invalido": None}
    return cel


def cal(ws, cel, valor, nome=None, ate=None, quebra=True, negrito=False, regra=False, centro=False):
    c = G.calculada(ws, cel, valor, negrito=negrito)
    _alinhar(c, quebra=quebra, centro=centro)
    _mesclar(ws, cel, ate)
    if nome:
        reg(nome, ws, cel)
        if regra:
            MAPA.numeros_de_regra[nome] = MapaMestre.ref(ws.title, cel)
    elif regra:
        raise ValueError(f"número de regra sem nome lógico em {ws.title}!{cel}")
    return c


def av(ws, cel, nome, formula, ate=None):
    c = G.aviso(ws, cel, formula)
    c.alignment = Alignment(vertical="center", wrap_text=True)
    _mesclar(ws, cel, ate)
    reg(nome, ws, cel)
    MAPA.avisos.setdefault(ws.title, []).append(cel)
    return c


def aux(ws, cel, formula, nome=None):
    """Célula auxiliar (cinza) — sempre em coluna oculta à direita de L (D11)."""
    cel = cel.replace("$", "")
    letra = re.match(r"[A-Z]+", cel).group(0)
    if _col(letra) <= 12:
        raise ValueError(f"auxiliar em coluna visível: {ws.title}!{cel}")
    c = G.nao_se_aplica(ws, cel, formula)
    MAPA.ocultas.setdefault(ws.title, set()).add(letra)
    if nome:
        reg(nome, ws, cel)
    return c


def rot(ws, cel, s, ate=None, negrito=False, italico=False, quebra=True):
    c = G.texto(ws, cel, s, negrito=negrito, italico=italico)
    _alinhar(c, quebra=quebra)
    _mesclar(ws, cel, ate)
    return c


def cab(ws, cel, s, ate=None):
    c = G.cabecalho(ws, cel, s)
    _mesclar(ws, cel, ate)
    return c


def cabs(ws, linha, nomes):
    """nomes: {coluna inicial: texto} ou {coluna: (texto, coluna final)}."""
    for col, v in nomes.items():
        if isinstance(v, tuple):
            cab(ws, f"{col}{linha}", v[0], ate=v[1])
        else:
            cab(ws, f"{col}{linha}", v)


def titulo(ws, linha, s, ate="L"):
    G.titulo(ws, f"A{linha}", s, ate=ate)
    ws.merge_cells(f"A{linha}:{ate}{linha}")
    ws.row_dimensions[linha].height = 20


def sugestao(ws, cel, h, nome, celulas=(), ate=None, extra=""):
    """Rótulo fixo 'Sugestão da planilha — não é regra do livro (H#)' + registro em `sugestoes`."""
    s = f"{ROTULO_SUGESTAO} ({h})" + (f": {extra}" if extra else "")
    c = G.texto(ws, cel, s, italico=True)
    _alinhar(c)
    _mesclar(ws, cel, ate)
    reg(nome, ws, cel)
    MAPA.sugestoes.append({"h": h, "rotulo": MapaMestre.ref(ws.title, cel), "nome": nome,
                           "celulas": [MapaMestre.ref(ws.title, x) for x in celulas]})


def sugestao_formula(nome, h, celulas=()):
    """Célula calculada cujo próprio texto traz o rótulo de Sugestão (registro em `sugestoes`)."""
    aba, cel = REFS[nome]
    MAPA.sugestoes.append({"h": h, "rotulo": MapaMestre.ref(aba, cel), "nome": nome, "na_formula": True,
                           "celulas": [MapaMestre.ref(aba, x) for x in celulas]})


def constante(chave, valor, secao):
    MAPA.constantes[chave] = [valor, secao]
    return valor


def subtabela(ws, id_, cab_linhas, colunas, itens):
    MAPA.subtabelas.append({"aba": ws.title, "id": id_, "cabecalhos": list(cab_linhas), "colunas": colunas,
                            "itens": itens})


def pior(ws, cel, texto):
    MAPA.pior[MapaMestre.ref(ws.title, cel)] = texto


PALAVRA_MAX = 16        # o texto de pior caso pula trechos com "palavra" maior (nomes de arquivo do changelog)


def texto_pior(n, deslocamento=0):
    """renderizar_ficha.texto_pior_caso(n, d) no primeiro d a partir de `deslocamento` sem palavra de mais de
    PALAVRA_MAX letras: o livro tem nomes de arquivo ("00-changelog-v01-para-v10.md,") que nenhum Mestre digita, e
    o trecho escolhido mudava com o número de entradas. Usado pelo gerador e pelas suítes visual e preview."""
    for d in range(deslocamento, deslocamento + 400):
        s = R.texto_pior_caso(n, d)
        if max((len(w) for w in s.split()), default=0) <= PALAVRA_MAX:
            return s
    return R.texto_pior_caso(n, deslocamento)


# ---------------------------------------------------------------------------
# Workbook, cabeçalho das abas e aba "Em construção"
# ---------------------------------------------------------------------------

def novo_workbook():
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils.indexed_list import IndexedList
    from openpyxl.workbook.properties import CalcProperties
    wb = Workbook()
    arial = Font(name=G.FONTE, size=G.TAM_CORPO)
    wb._fonts = IndexedList([arial])
    wb._named_styles["Normal"].font = arial
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.active.title = ABAS[0]
    for nome in ABAS[1:]:
        wb.create_sheet(nome)
    for ws in wb.worksheets:
        larguras_grade(ws)
    return wb


def larguras_grade(ws, px=None):
    for letra, v in (px or GRADE_PX).items():
        ws.column_dimensions[letra].width = G._px_largura(v)


def cabecalho_aba(ws):
    G.titulo(ws, "A1", f"{ws.title} — Explorando Galáxias {VERSAO}", tamanho=G.TAM_TITULO_ABA, ate="L")
    ws.merge_cells("A1:L1")
    ws.row_dimensions[1].height = 24
    G.texto(ws, "A2", "Legenda:", negrito=True)
    rot(ws, "A3", SUBTITULOS[ws.title], ate="L", italico=True)
    reg(f"{SLUG[ws.title]}.titulo", ws, "A1")
    reg(f"{SLUG[ws.title]}.legenda", ws, "A2")
    reg(f"{SLUG[ws.title]}.subtitulo", ws, "A3")


def aba_em_construcao(ws):
    rot(ws, "A5", EM_CONSTRUCAO, ate="L")
    reg(f"{SLUG[ws.title]}.em_construcao", ws, "A5")


# ---------------------------------------------------------------------------
# Aba Dados (blocos do livro; listas de validação à direita)
# ---------------------------------------------------------------------------

LINHA_INICIAL_DADOS = 5


def escrever_dados(ws, blocos, listas):
    linha = LINHA_INICIAL_DADOS
    largura = max(len(b["cabecalho"]) for b in blocos)
    for b in blocos:
        ncol = len(b["cabecalho"])
        ultima = get_column_letter(ncol)
        G.titulo(ws, f"A{linha}", b["titulo"], ate=ultima if ncol > 1 else None)
        lin_titulo = linha
        linha += 1
        for j, nome in enumerate(b["cabecalho"], start=1):
            G.cabecalho(ws, ws.cell(linha, j).coordinate, nome)
        lin_cab = linha
        linha += 1
        primeira = linha
        for registro in b["linhas"]:
            for j, v in enumerate(registro, start=1):
                cel = G.calculada(ws, ws.cell(linha, j).coordinate, v if v is not None else "")
                cel.alignment = Alignment(vertical="center", wrap_text=True)
            linha += 1
        ultima_linha = linha - 1
        MAPA.blocos[f"dados.{b['id']}"] = {
            "titulo_texto": b["titulo"], "titulo": MapaMestre.ref("Dados", f"A{lin_titulo}"),
            "intervalo": MapaMestre.ref("Dados", f"$A${primeira}:${ultima}${ultima_linha}"),
            "primeira_linha": primeira, "ultima_linha": ultima_linha, "linhas": len(b["linhas"]),
            "colunas": {nome: get_column_letter(j) for j, nome in enumerate(b["cabecalho"], 1)},
            "fonte": b.get("fonte", ""),
        }
        linha += 1
    col0 = largura + 2
    lin_tit = LINHA_INICIAL_DADOS
    G.titulo(ws, ws.cell(lin_tit - 1, col0).coordinate, "Listas de validação (uma por coluna; a fonte está no "
             "cabeçalho)", ate=get_column_letter(col0 + len(listas) - 1))
    for k, l in enumerate(listas):
        letra = get_column_letter(col0 + k)
        G.cabecalho(ws, f"{letra}{lin_tit}", f"{l['titulo']} — {l['fonte']}")
        for i, v in enumerate(l["valores"], start=1):
            G.calculada(ws, f"{letra}{lin_tit + i}", v).alignment = Alignment(vertical="center", wrap_text=True)
        MAPA.listas[f"lista.{l['id']}"] = MapaMestre.ref(
            "Dados", f"${letra}${lin_tit + 1}:${letra}${lin_tit + len(l['valores'])}")
        MAPA_LISTAS_VALORES[f"lista.{l['id']}"] = list(l["valores"])
        ws.column_dimensions[letra].width = 26
    for j in range(1, largura + 1):
        ws.column_dimensions[get_column_letter(j)].width = 22
    ws.column_dimensions["A"].width = 34


def dcol(bloco, coluna):
    b = MAPA.blocos[f"dados.{bloco}"]
    letra = b["colunas"][coluna]
    return f"'Dados'!${letra}${b['primeira_linha']}:${letra}${b['ultima_linha']}"


def dcel(bloco, coluna, i):
    """Célula da linha i (1-based) do bloco."""
    b = MAPA.blocos[f"dados.{bloco}"]
    return f"'Dados'!${b['colunas'][coluna]}${b['primeira_linha'] + i - 1}"


_IDS_TEXTOS = []


def dtexto(id_):
    """Célula do trecho de regra `id_` do bloco 'textos' (aba Dados); a chave interna não vai para a fórmula."""
    if not _IDS_TEXTOS:
        import mestre_dados3 as D3
        _IDS_TEXTOS[:] = [t[0] for t in D3.textos()]
    return dcel("textos", "Texto", _IDS_TEXTOS.index(id_) + 1)


def dlista(id_):
    return MAPA.listas[f"lista.{id_}"]


def faixa_da_lista(x):
    """Faixa ("9-12") a partir da opção escolhida na lista ("Faixa 9-12"); "" se não for opção da lista.
    A lista não mostra "9-12" porque o Google leria como data (requisito do Google, revisão 2 da ficha)."""
    return f'IFERROR(INDEX({dcol("faixas", "Faixa")},MATCH({x},{dlista("faixas")},0)),"")'


def valor_digitado(nome, v):
    """Valor que o mestre escolhe na lista para o valor lógico `v` (testes e Exemplo): faixa "9-12" → "Faixa 9-12"."""
    info = MAPA.entradas.get(nome) if MAPA.entradas else None
    fonte = (info or {}).get("fonte") or ""
    if isinstance(v, str) and v in D_FAIXAS and fonte in LISTAS_DE_FAIXA:
        return f"Faixa {v}"
    return v


LISTAS_DE_FAIXA = ("lista.faixas", "lista.filtro_faixa", "lista.faixa_bloco")
D_FAIXAS = ("1-4", "5-8", "9-12", "13-16", "17-20")


def dbusca(bloco, coluna_chave, chave, coluna):
    """INDEX/MATCH exato num bloco da aba Dados, protegido por IFERROR."""
    return f'IFERROR(INDEX({dcol(bloco, coluna)},MATCH({chave},{dcol(bloco, coluna_chave)},0)),"")'


# ---------------------------------------------------------------------------
# Acabamento: avisos, colunas ocultas, layout pela régua
# ---------------------------------------------------------------------------

def formatar_avisos(wb):
    fundo = PatternFill(start_color=G.COR_AVISO_FUNDO, end_color=G.COR_AVISO_FUNDO, fill_type="solid")
    for aba, cels in MAPA.avisos.items():
        ws = wb[aba]
        for cel in cels:
            ws.conditional_formatting.add(cel, FormulaRule(formula=[f"LEN({cel})>0"], fill=fundo))


def ocultar_auxiliares(wb):
    for aba, letras in MAPA.ocultas.items():
        ws = wb[aba]
        for letra in letras:
            ws.column_dimensions[letra].hidden = True
        if aba not in ("Tabelas", "Dados"):
            ws.column_dimensions["M"].hidden = True


def _texto_formula(estima, s, aba):
    try:
        return estima.formula(s, aba)
    except Exception:  # noqa: BLE001
        return ""


def ajustar_layout(wb, entradas_pior):
    """Altura das linhas pela régua da ficha (renderizar_ficha.medir, com 10% de folga): texto fixo
    pelo próprio texto; fórmula pelo texto mais longo que ela pode mostrar (PiorTexto, ou o pior caso
    registrado em MAPA.pior); entradas de texto livre pelo pior caso. Não alarga coluna: se uma palavra
    não cabe na largura, é fatal (a área de jogo fica em 1360 px)."""
    estima = R.PiorTexto(wb, entradas_texto=entradas_pior)
    problemas = []
    listas = {MAPA.celulas[n] for n, i in MAPA.entradas.items() if i["tipo"] == "lista"}
    for ws in wb.worksheets:
        topo, coberta = R._mesclas(ws)
        alturas = {}
        for linha in ws.iter_rows():
            for cel in linha:
                r, c = cel.row, cel.column
                if (r, c) in coberta or ws.column_dimensions[get_column_letter(c)].hidden or c > 12 and \
                        ws.title not in ("Tabelas", "Dados"):
                    continue
                ref_ = MapaMestre.ref(ws.title, cel.coordinate)
                s = cel.value
                if ref_ in MAPA.pior:
                    s = MAPA.pior[ref_]
                    if isinstance(s, list):          # vários candidatos: o que pede mais linhas
                        f0 = R.fonte_da_celula(cel)
                        w0 = R.area_px(ws, r, c, topo)[0] - R.RESERVA - (R.SETA_LISTA if ref_ in listas else 0)
                        s = max(s, key=lambda x: len(R.quebrar(x, f0, w0 / R.FOLGA)))
                elif s is None:
                    s = entradas_pior.get(ref_)
                elif isinstance(s, str) and s.startswith("="):
                    s = _texto_formula(estima, s, ws.title)
                if s is None:
                    continue
                s = R.texto_exibido(s)
                if not s.strip():
                    continue
                f = R.fonte_da_celula(cel)
                asc, desc = f.getmetrics()
                lh = asc + desc
                w, _ = R.area_px(ws, r, c, topo)
                util = w - R.RESERVA - (R.SETA_LISTA if ref_ in listas else 0)   # a seta da lista suspensa
                quebra = bool(cel.alignment is not None and cel.alignment.wrap_text)
                if not quebra:
                    larga = max(f.getlength(x) for x in s.split("\n")) * R.FOLGA
                    if larga > util:
                        novo = copy(cel.alignment)
                        novo.wrap_text = True
                        cel.alignment = novo
                linhas = R.quebrar(s, f, util / R.FOLGA)
                for ln in linhas:
                    if f.getlength(ln) * R.FOLGA > util:
                        problemas.append(f"{ref_}: palavra {ln[:30]!r} não cabe em {util} px")
                r2 = topo.get((r, c), (r, c))[0]
                precisa = len(linhas) * lh + 4
                outras = sum(R.px_linha(ws, k) for k in range(r, r2))
                alturas[r2] = max(alturas.get(r2, 0), precisa - outras)
        for r2, px in alturas.items():
            dim = ws.row_dimensions[r2]
            atual = dim.height or ws.sheet_format.defaultRowHeight or 15
            if px > atual * 4 / 3:
                dim.height = math.ceil(px * 3 / 4)
    return problemas
