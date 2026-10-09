# -*- coding: utf-8 -*-
"""Requisito do Google (lição da revisão 2 da ficha) — leitura protegida das entradas e contador de erros.

O Google lê como FÓRMULA o que o mestre digita (ou escolhe numa lista) começando com =, +, - ou @: "+1 em dois"
vira #ERROR! na própria célula, e o erro se espalha para tudo que a lê. Para conter:
  1. nenhuma fórmula lê uma entrada diretamente: lê uma célula da camada de leitura protegida (colunas ocultas à
     direita da aba) com =IF(ISERROR(X),"",IF(ISBLANK(X),"",X)) — o vazio é explícito (no Google a referência a
     célula vazia é vazia; no Excel e na formulas é 0);
  2. cada linha com entrada tem um sinal de erro (coluna oculta) que acende o aviso da linha em PT-BR;
  3. a aba Início conta as células com erro de cada aba (SUMPRODUCT(ISERROR(intervalo)*1)), sem referência circular.
Entradas: as do mapa (ent) e as vagas das listas editáveis da aba Tabelas (que já saem preenchidas com o livro).
Reaproveita de build\\gerar_ficha.py, por import e sem alterar a ficha, o tokenizador (_tokens, _RE_REF, _limites)
e as mensagens (MSG_ERRO_*). O lint usa testar_ficha._lint_google com o mapa da Mestre.
"""

from openpyxl.utils import get_column_letter

from mestre import nucleo as N

G = N.G
LEITURA = {}       # "'Aba'!X5" (camada) -> "'Aba'!B5" (célula lida)
SINAIS = {}        # "'Aba'!Y5" (sinal da linha) -> ["'Aba'!B5", ...]
CONTADOR = {}      # "'Início'!B40" -> intervalos contados
LINHAS_CONTADOR = {}
PARES_POR_LINHA = 5     # aba/contagem em A:B, C:D, E:F, G:H, I:J (K:L fica para aviso, P7)


def iniciar():
    LEITURA.clear()
    SINAIS.clear()
    CONTADOR.clear()
    LINHAS_CONTADOR.clear()


def entradas_por_aba():
    """aba -> {(linha, coluna): (nome, é_lista)}: entradas do mapa + vagas das listas editáveis (Tabelas)."""
    saida = {}
    for nome, info in N.MAPA.entradas.items():
        aba, cel = N.REFS[nome]
        r1, c1, _, _ = G._limites(cel.split(":")[0], None)
        saida.setdefault(aba, {})[(r1, c1)] = (nome, info["tipo"] == "lista")
    for id_, t in N.MAPA.tabelas.items():
        for k, cel in enumerate(t["vagas"]):
            r1, c1, _, _ = G._limites(cel, None)
            saida.setdefault("Tabelas", {})[(r1, c1)] = (f"{id_}.vaga{k + 1}", False)
        for k, cel in enumerate(t.get("ambientes", [])):
            r1, c1, _, _ = G._limites(cel, None)
            saida.setdefault("Tabelas", {})[(r1, c1)] = (f"{id_}.ambiente{k + 1}", True)
        for k, cel in enumerate(t.get("segunda", [])):      # 2ª coluna das tabelas de duas colunas (Fase 2)
            r1, c1, _, _ = G._limites(cel, None)
            saida.setdefault("Tabelas", {})[(r1, c1)] = (f"{id_}.segunda{k + 1}", False)
    return saida


def montar_contador(ws, r):
    """Bloco do contador na Início (fórmulas preenchidas por proteger()). Devolve a próxima linha livre."""
    N.titulo(ws, r, "Células com erro na planilha (o Google leu como fórmula um texto digitado)")
    r0 = r + 1
    N.rot(ws, f"A{r0}", "Células com erro na planilha", negrito=True)
    N.cal(ws, f"B{r0}", "=0", nome="inicio.erros.total", negrito=True, centro=True)
    N.cal(ws, f"C{r0}", '=""', nome="inicio.erros.explicacao", ate="J")
    rr = r0 + 1
    for k, aba in enumerate(N.ABAS):
        if k % PARES_POR_LINHA == 0:
            rr += 1 if k else 0
        col = 1 + 2 * (k % PARES_POR_LINHA)
        N.rot(ws, f"{get_column_letter(col)}{rr}", aba)
        N.cal(ws, f"{get_column_letter(col + 1)}{rr}", "=0", nome=f"inicio.erros.{N.SLUG[aba]}", centro=True)
    LINHAS_CONTADOR.update(inicio=r0, fim=rr)
    return rr + 1


def proteger(wb):
    """Camada de leitura protegida + sinais de erro por linha + avisos + contador da Início. Chamar depois de
    resolver_marcadores e antes de ocultar_auxiliares/formatar_avisos/ajustar_layout."""
    import renderizar_ficha as R
    entradas = entradas_por_aba()
    # 1) referências que tocam uma entrada
    pedidos, achados = {}, []
    for ws in wb.worksheets:
        for linha in ws.iter_rows():
            for cel in linha:
                v = cel.value
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                for k, t in enumerate(G._tokens(v)):
                    if t.type != "OPERAND" or t.subtype != "RANGE":
                        continue
                    m = G._RE_REF.match(t.value)
                    if not m:
                        continue
                    alvo = m.group("aba").strip("'").replace("''", "'") if m.group("aba") else ws.title
                    if alvo not in entradas:
                        continue
                    r0, c0, r1, c1 = lim = G._limites(m.group("c1"), m.group("c2"))
                    if not any(r0 <= r <= r1 and c0 <= c <= c1 for r, c in entradas[alvo]):
                        continue
                    achados.append((ws.title, cel.coordinate, k, alvo, lim, m.group("aba")))
                    pedidos.setdefault(alvo, set()).update((r, c) for r in range(r0, r1 + 1)
                                                           for c in range(c0, c1 + 1))
    # 2) colunas da camada (à direita de tudo; colunas contíguas: intervalo vira intervalo) e a coluna de sinal
    base, cmin, sinal_col = {}, {}, {}
    for aba in entradas:
        ws = wb[aba]
        cols = [c for _, c in pedidos.get(aba, ())]
        base[aba] = max(ws.max_column, N.PRIMEIRA_AUX) + 2
        cmin[aba] = min(cols) if cols else 1
        sinal_col[aba] = base[aba] + ((max(cols) - cmin[aba] + 1) if cols else 0)
    for aba, cels in pedidos.items():
        ws = wb[aba]
        for r, c in sorted(cels):
            orig = f"${get_column_letter(c)}${r}"
            esp = f"{get_column_letter(base[aba] + c - cmin[aba])}{r}"
            N.aux(ws, esp, f'=IF(ISERROR({orig}),"",IF(ISBLANK({orig}),"",{orig}))')
            LEITURA[N.MapaMestre.ref(aba, esp)] = N.MapaMestre.ref(aba, orig.replace("$", ""))
    # 3) reescreve as referências para a camada
    por_celula = {}
    for aba_f, coord, k, alvo, lim, prefixo in achados:
        por_celula.setdefault((aba_f, coord), []).append((k, alvo, lim, prefixo))
    for (aba_f, coord), trocas in por_celula.items():
        cel = wb[aba_f][coord]
        toks = G._tokens(cel.value)
        for k, alvo, (r0, c0, r1, c1), prefixo in trocas:
            a = f"${get_column_letter(base[alvo] + c0 - cmin[alvo])}${r0}"
            b = f"${get_column_letter(base[alvo] + c1 - cmin[alvo])}${r1}"
            nova = a if (r0, c0) == (r1, c1) else f"{a}:{b}"
            toks[k].value = f"{prefixo}!{nova}" if prefixo else nova
        cel.value = "=" + "".join(t.value for t in toks)
    # 4) sinal de erro por linha e o aviso da linha (o mais à direita; sem aviso, um novo em K:L)
    avisos_linha = {}
    for aba, cels in N.MAPA.avisos.items():
        for x in cels:
            r1, c1, _, _ = G._limites(x, None)
            avisos_linha.setdefault((aba, r1), []).append((c1, x))
    por_aviso, sem_aviso, adiadas = {}, [], []
    for aba, ents in entradas.items():
        ws = wb[aba]
        topo, coberta = R._mesclas(ws)
        por_linha = {}
        for (r, c), (nome, lista) in ents.items():
            por_linha.setdefault(r, []).append((c, lista))
        for r, itens in sorted(por_linha.items()):
            itens.sort()
            sinal = f"{get_column_letter(sinal_col[aba])}{r}"
            refs = [f"${get_column_letter(c)}${r}" for c, _ in itens]
            N.aux(ws, sinal, "=" + "+".join(f"ISERROR({x})*1" for x in refs))
            SINAIS[N.MapaMestre.ref(aba, sinal)] = [N.MapaMestre.ref(aba, x.replace("$", "")) for x in refs]
            tipos = {lista for _, lista in itens}
            alvo = avisos_linha.get((aba, r))
            rotulo = ""
            if alvo:
                cel_av = max(alvo)[1]
            else:
                livre = lambda k: (r, k) not in coberta and (r, k) not in topo and ws.cell(r, k).value is None  # noqa: E731
                if livre(11):
                    cel_av = f"K{r}"
                    N.av(ws, cel_av, f"{N.SLUG[aba]}.aviso.erro.{r}", '=""', ate="L" if livre(12) else None)
                    avisos_linha[(aba, r)] = [(11, cel_av)]
                else:
                    adiadas.append((aba, r, sinal, tipos, itens, ents))
                    continue
            por_aviso.setdefault((aba, cel_av), []).append((sinal, tipos, rotulo))
    for aba, r, sinal, tipos, itens, ents in adiadas:
        # K ocupada (campo mesclado até L): o aviso mais próximo da mesma aba, com a linha na mensagem
        perto = sorted((abs(rr - r), rr, max(v)[1]) for (a_, rr), v in avisos_linha.items() if a_ == aba)
        if not perto or perto[0][0] > 6:
            sem_aviso.append((aba, r, [ents[(r, c)][0] for c, _ in itens]))
            continue
        por_aviso.setdefault((aba, perto[0][2]), []).append((sinal, tipos, f"Linha {r}: "))
    if sem_aviso:
        raise ValueError(f"linhas com entrada e sem célula de aviso: {sem_aviso[:10]}")
    for (aba, cel_av), lst in por_aviso.items():
        ws = wb[aba]
        expr = ws[cel_av].value[1:]
        for sinal, tipos, rotulo in reversed(lst):
            msg = G.MSG_ERRO_GERAL if len(tipos) > 1 else (G.MSG_ERRO_LISTA if tipos == {True} else G.MSG_ERRO_TEXTO)
            expr = f"IF({sinal}>0,{N.q(rotulo + msg)},{expr})"
        ws[cel_av].value = "=" + expr
    # 5) contador da Início: ISERROR sobre a área usada de cada aba (sem as linhas do próprio contador)
    ini = wb["Início"]
    for aba in N.ABAS:
        ws = wb[aba]
        fim = get_column_letter(ws.max_column)
        if aba == "Início":
            a, b = LINHAS_CONTADOR["inicio"], LINHAS_CONTADOR["fim"]
            faixas = [f"$A$1:${fim}${a - 1}"] + ([f"$A${b + 1}:${fim}${ws.max_row}"] if ws.max_row > b else [])
        else:
            faixas = [f"'{aba}'!$A$1:${fim}${ws.max_row}"]
        _, cel = N.REFS[f"inicio.erros.{N.SLUG[aba]}"]
        ini[cel].value = "=" + "+".join(f"SUMPRODUCT(ISERROR({x})*1)" for x in faixas)
        CONTADOR[N.MapaMestre.ref("Início", cel)] = faixas
    cels = [N.REFS[f"inicio.erros.{N.SLUG[a]}"][1] for a in N.ABAS]
    _, tot = N.REFS["inicio.erros.total"]
    ini[tot].value = "=" + "+".join(N._abs(c) for c in cels)
    CONTADOR[N.MapaMestre.ref("Início", tot)] = cels
    _, exp = N.REFS["inicio.erros.explicacao"]
    ini[exp].value = (f'=IF({N._abs(tot)}=0,"Nenhuma célula com erro.","O Google leu como fórmula um texto que começa '
                      f'com +, - ou =. Veja abaixo a aba com erro: o aviso da linha diz o que fazer. Apague e escreva '
                      f'de novo com um apóstrofo (\') na frente, ou escolha de novo na lista.")')
    CONTADOR[N.MapaMestre.ref("Início", exp)] = [tot]
    return {"leitura": len(LEITURA), "sinais": len(SINAIS)}
