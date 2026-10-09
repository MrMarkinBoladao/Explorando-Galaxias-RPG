# -*- coding: utf-8 -*-
"""Aba Início (design §6.1, §5): comece por aqui, semente da campanha, painel da campanha (Fase 3), índice das abas,
painel de avisos por aba (Fase 3) e o contador de células com erro (requisito do Google)."""

import re

from openpyxl.utils import column_index_from_string

from mestre import nucleo as N
from mestre import sorteio as S
from mestre.nucleo import T, ent, cal, av, rot, cabs, titulo

INDICE = [
    ("Início", "Semente da campanha, painel, índice e avisos.", "Na primeira vez e no começo de cada sessão."),
    ("Campanha", "Mesa, marcos, facções, relógios, linha do tempo, Ficha de Decisões e Sessão Zero.",
     "Ao começar e entre sessões."),
    ("Grupo", "Os PJs copiados da ficha de cada jogador, com os testes que o Mestre pede.",
     "Ao começar e quando a ficha muda."),
    ("Sessões", "Preparar a próxima sessão (cenas, pistas, ganchos, checklist) e o diário.", "Entre sessões."),
    ("Missões", "Missões oferecidas, ativas e concluídas.", "Entre sessões."),
    ("NPCs", "Gerador de NPCs, elenco e cartões.", "Preparando a história."),
    ("Inimigos", "Criador de Inimigos: a ficha de 28.1 pronta.", "Preparando o combate."),
    ("Bestiário", "As 32 fichas do capítulo 28 com filtro.", "Preparando o combate."),
    ("Encontros", "Orçamento, encontros salvos, contrato da Fraqueza, encontro aleatório.", "Preparando o combate."),
    ("Combate", "Fila de Ação, PV, Tenacidade e Quebra, condições, Morrendo.", "Na mesa, durante a luta."),
    ("Aventuras", "Gerador de aventuras com 5 cenas.", "Preparando a história."),
    ("Recompensas", "Entrega de marco, achados, Cone e Conjunto, tesouro.", "Depois de um encontro ou marco."),
    ("Mundos", "Planetas, estações, naves, facções, organizações e nomes.", "Preparando a história."),
    ("Improviso", "Rumores, eventos, loja, bugigangas, oráculo, rolador e DT rápida.", "Na mesa, quando precisar."),
    ("Escudo do Mestre", "Referência de mesa para imprimir (A4 paisagem).", "Antes da sessão."),
    ("Minhas Tabelas", "Dez tabelas do próprio Mestre, com sorteio.", "Quando quiser."),
    ("Tabelas", "Listas dos geradores, editáveis (100 vagas).", "Para mudar o que os geradores sorteiam."),
    ("Dados", "Tabelas do livro usadas pelas fórmulas. Não edite.", "Nunca: o gerador reescreve."),
]
LINHAS_ALERTA = {}      # aba -> célula da contagem e do primeiro aviso (preenchidas por preencher_alertas)


def montar(wb):
    ws = wb["Início"]
    titulo(ws, 5, "Comece por aqui")
    passos = [
        "1. Abra no Google Planilhas (Arquivo → Importar) e trabalhe numa cópia por campanha.",
        "2. Aba Campanha: nível do grupo e nº de jogadores (uma vez; todas as abas leem daqui). Faça a Sessão Zero.",
        "3. Aba Grupo: copie da ficha de cada jogador os números de combate e os testes.",
        "4. Prepare: Encontros (com Inimigos e Bestiário), NPCs, Aventuras, Mundos; a sessão na aba Sessões.",
        "5. Jogue com Combate e Improviso; o Escudo do Mestre impresso fica na mesa.",
        "6. Depois: diário (Sessões), Missões, Recompensas, relógios e linha do tempo (Campanha).",
        "Amarelo = você preenche; azul = automático; vermelho = aviso; 'Sugestão da planilha' = não é regra do livro.",
        # requisito do Google (revisão 2 da ficha): o texto que começa com +, - ou = vira fórmula no Google
        "Não comece um texto com +, - ou =: o Google entende como fórmula. Se precisar, digite um apóstrofo (') "
        "antes.",
    ]
    for k, p in enumerate(passos):
        rot(ws, f"A{6 + k}", p, ate="L", negrito=k == len(passos) - 1)
    N.reg("inicio.dica_formula", ws, f"A{6 + len(passos) - 1}")
    r = 6 + len(passos) + 1
    titulo(ws, r, "Semente da campanha (geradores determinísticos, sem sorteio escondido)")
    rot(ws, f"A{r + 1}", "Semente da campanha", negrito=True)
    ent(ws, f"B{r + 1}", "inicio.semente", tipo="inteiro", minimo=1, maximo=2147483646, ate="C",
        rotulo="Semente da campanha", centro=True, amostra=2026)
    rot(ws, f"D{r + 1}", "Semente efetiva", negrito=True)
    cal(ws, f"E{r + 1}", S.semente_efetiva(T("inicio.semente")), nome="inicio.semente_ef", ate="F", centro=True,
        regra=True)
    av(ws, f"K{r + 1}", "inicio.aviso.semente", S.aviso_semente(T("inicio.semente")), ate="L")
    rot(ws, f"G{r + 1}", "Vazia = 12345.", ate="J", italico=True)
    rot(ws, f"A{r + 2}", "Mesma semente e mesma 'Rolagem nº' dão sempre o mesmo resultado; editar outra célula não "
                         "muda o sorteio. Para rolar de novo, some 1 na 'Rolagem nº' do gerador; para outra "
                         "campanha, troque a semente.", ate="L", italico=True)
    r = _painel(ws, r + 4)
    r += 1
    titulo(ws, r, "Índice das abas")
    rr = r + 1
    cabecalhos = []
    for k, (aba, serve, quando) in enumerate(INDICE):
        if k % 9 == 0:
            cabs(ws, rr, {"A": "Aba", "B": ("Para que serve", "G"), "H": ("Quando usar", "L")})
            cabecalhos.append(rr)
            rr += 1
        rot(ws, f"A{rr}", aba, negrito=True)
        rot(ws, f"B{rr}", serve, ate="G")
        rot(ws, f"H{rr}", quando, ate="L")
        rr += 1
    N.subtabela(ws, "inicio.indice", cabecalhos, "A:L", len(INDICE))
    rr = _alertas(ws, rr + 1)
    from mestre import protecao
    protecao.montar_contador(ws, rr + 1)


def _painel(ws, r):
    """Painel da campanha (calculado)."""
    titulo(ws, r, "Painel da campanha (automático)")
    cabs(ws, r + 1, {"A": "Campo", "B": ("Agora", "F"), "G": ("De onde vem", "L")})
    NV = T("campanha.nivel_ef")
    itens = [
        ("campanha", "Campanha", f'=IF(LEN({T("campanha.nome")})>0,{T("campanha.nome")},"(sem nome: aba Campanha)")',
         "Aba Campanha, Mesa."),
        ("sessao", "Sessão atual", f'=IF(ISNUMBER({T("campanha.sessao")}),{T("campanha.sessao")},"(vazia)")&" · dia de '
                                   f'campanha "&{T("campanha.dia_ef")}', "Aba Campanha, Mesa."),
        ("nivel", "Nível e faixa", f'="Nível "&{NV}&", faixa "&{T("campanha.faixa")}&" · Eficiência +"&'
                                   f'{T("campanha.ef")}', "26.1, 26.2, 26.3."),
        ("ph", "PH do grupo", f'="Máximo "&{T("campanha.ph_max")}&" · início do combate "&{T("campanha.ph_ini")}&" ("&'
                              f'{T("campanha.jogadores_ef")}&" jogadores)"', "16.2."),
        ("missoes", "Missões ativas", f'={T("missoes.n.ativa")}&" ativa(s) · "&{T("missoes.n.oferecida")}&'
                                      f'" oferecida(s) · "&{T("missoes.n.concluida")}&" concluída(s)"', "Aba Missões."),
        ("relogios", "Relógios a 1 de encher", f'=COUNTIF({T("campanha.col.rel_sit")},"Falta 1")',
         "Aba Campanha, relógios (Sugestão da planilha, H16)."),
        ("ressonancia", "Próxima Ressonância", f'={T("campanha.prox_ress")}', "26.7; sempre marco de história."),
        ("ritmo", "Sessões no nível atual", f'=IF(ISNUMBER({T("campanha.sessoes_nivel")}),{T("campanha.sessoes_nivel")}'
                                            f'&IF({T("campanha.sessoes_nivel")}=1," sessão"," sessões")&" no nível · ",'
                                            f'"")&{T("campanha.ritmo")}', "26.1."),
        ("proxima", "Próxima sessão preparada",
         f'=IF(LEN({T("sessoes.prep.titulo")})>0,IF(ISNUMBER({T("sessoes.prep.n")}),"Sessão "&{T("sessoes.prep.n")}&": ",'
         f'"")&{T("sessoes.prep.titulo")},"(nenhuma: aba Sessões)")', "Aba Sessões."),
    ]
    for k, (chave, rotulo, f, fonte) in enumerate(itens):
        rr = r + 2 + k
        rot(ws, f"A{rr}", rotulo, negrito=True)
        cal(ws, f"B{rr}", f, nome=f"inicio.painel.{chave}", ate="F", centro=chave == "relogios")
        rot(ws, f"G{rr}", fonte, ate="L")
    N.subtabela(ws, "inicio.painel", [r + 1], "A:L", len(itens))
    return r + 2 + len(itens)


def _alertas(ws, r):
    """Painel de avisos por aba: a contagem e o primeiro aviso aceso (as fórmulas entram em preencher_alertas, depois
    da camada de leitura protegida, que também cria avisos)."""
    titulo(ws, r, "Avisos por aba (automático): quantos estão acesos e o primeiro deles")
    rr = r + 1
    cabecalhos = []
    for k, aba in enumerate(N.ABAS):
        if k % 9 == 0:
            cabs(ws, rr, {"A": "Aba", "B": "Avisos acesos", "C": ("Primeiro aviso aceso", "L")})
            cabecalhos.append(rr)
            rr += 1
        rot(ws, f"A{rr}", aba, negrito=True)
        cal(ws, f"B{rr}", "=0", nome=f"inicio.alertas.{N.SLUG[aba]}.n", centro=True)
        cal(ws, f"C{rr}", '=""', nome=f"inicio.alertas.{N.SLUG[aba]}.primeiro", ate="L")
        LINHAS_ALERTA[aba] = (f"B{rr}", f"C{rr}")
        rr += 1
    N.subtabela(ws, "inicio.alertas", cabecalhos, "A:L", len(N.ABAS))
    return rr


def _corridas(cels):
    """Avisos da aba em corridas de células seguidas na mesma coluna: [(coluna, linha inicial, linha final)], na ordem
    de leitura (linha, coluna)."""
    pos = sorted((int(re.search(r"\d+", c).group(0)), re.match(r"[A-Z]+", c).group(0)) for c in cels)
    por_col = {}
    for lin, col in pos:
        por_col.setdefault(col, []).append(lin)
    corr = []
    for col, ls in por_col.items():
        ini = ant = ls[0]
        for x in ls[1:]:
            if x != ant + 1:
                corr.append((col, ini, ant))
                ini = x
            ant = x
        corr.append((col, ini, ant))
    return sorted(corr, key=lambda c: (c[1], column_index_from_string(c[0])))


def preencher_alertas(wb):
    """Chamar depois de protecao.proteger: contagem SUMPRODUCT((LEN(corrida)>0)*1) por corrida e o primeiro aviso por
    INDEX/MATCH("?*", corrida, 0), em colunas ocultas da Início; a aba soma as corridas e encadeia os primeiros."""
    ws = wb["Início"]
    lin = 1
    col_n, col_p, col_c = "AA", "AB", "AC"      # longe da camada de leitura protegida (que fica logo depois de M)
    for aba in N.ABAS:
        cels = N.MAPA.avisos.get(aba, [])
        cel_n, cel_p = LINHAS_ALERTA[aba]
        corr = _corridas(cels) if cels else []
        if not corr:
            ws[cel_n].value = "=0"
            ws[cel_p].value = '=""'
            continue
        pref = "" if aba == "Início" else f"'{aba}'!"
        a = lin
        for col, i0, i1 in corr:
            faixa = f"{pref}${col}${i0}:${col}${i1}"
            N.aux(ws, f"{col_n}{lin}", f"=SUMPRODUCT((LEN({faixa})>0)*1)")
            N.aux(ws, f"{col_p}{lin}", f'=IFERROR(INDEX({faixa},MATCH("?*",{faixa},0))&"","")')
            lin += 1
        b = lin - 1
        for k in range(b, a - 1, -1):          # encadeia: o primeiro aceso na ordem de leitura
            prox = f"{col_c}{k + 1}" if k < b else '""'
            N.aux(ws, f"{col_c}{k}", f"=IF(LEN({col_p}{k})>0,{col_p}{k},{prox})")
        ws[cel_n].value = f"=SUM(${col_n}${a}:${col_n}${b})"
        ws[cel_p].value = f"={col_c}{a}"
        lin += 1
    for c in (col_n, col_p, col_c):
        ws.column_dimensions[c].hidden = True
