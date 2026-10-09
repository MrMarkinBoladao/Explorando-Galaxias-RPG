# -*- coding: utf-8 -*-
"""Suítes e casos da Fase 2 (NPCs, Aventuras, Recompensas — R5, R6, R7; design §6.6, §6.11, §6.12, §6.16).

  suite_sabor          tabelas de sabor (mínimos, vazio, repetido, 200 caracteres, nomes ≥ 30 por cultura, NENHUM nome
                       de personagem do jogo — lista do próprio teste —, as entradas do livro no início das listas)
  dados_fase2          parser PRÓPRIO do teste (lê o .md) × aba Dados: equipamentos 24.1–24.3, verba 24.5, Cones 25.2,
                       Relíquias e Conjuntos 25.3, Ressonâncias 26.7, 27.7, 27.8, 27.16, 27.17 — item a item; e as
                       tabelas transcritas do oráculo (build\\oraculo_mestre_hist.py) contra o mesmo .md
  ouro_fase2           números do livro calculados NA planilha (24.3, 24.5, 25.1, 25.2, 25.3, 26.7)
  oraculo_fase2        casos oráculo × planilha (G = 200, 300, 400, 410, 420 em 300 Rolagens nº × 3 sementes, H22 em
                       todas as facções × faixas, Elenco e cartões, Tesouro) + validações de H9, H14, H18, H20, H22,
                       H25 e H26 no oráculo
Os casos de planilha entram na lista TAREFAS de testes_regras (calculados em paralelo, mestre\\paralelo.py)."""

import random
import re
import unicodedata

import openpyxl

import testar_ficha as TF
import oraculo_mestre as O
import oraculo_mestre_hist as OH
import oraculo_mestre_livro as L
from mestre import nucleo as N
from mestre import testes_comum as TC

Resultado = TF.Resultado

# ---------------------------------------------------------------------------
# Nomes de personagens do jogo (Honkai: Star Rail) — a planilha não pode usar nenhum (design §6.16)
# ---------------------------------------------------------------------------
PROIBIDOS = [
    "March 7th", "Março", "Dan Heng", "Himeko", "Welt", "Welt Yang", "Kafka", "Blade", "Silver Wolf", "Jing Yuan",
    "Bronya", "Seele", "Clara", "Svarog", "Gepard", "Serval", "Natasha", "Pela", "Sampo", "Hook", "Luka", "Lynx",
    "Qingque", "Yanqing", "Bailu", "Tingyun", "Sushang", "Yukong", "Luocha", "Fu Xuan", "Jingliu", "Topaz", "Numby",
    "Guinaifen", "Hanya", "Huohuo", "Argenti", "Ruan Mei", "Ratio", "Veritas", "Xueyi", "Black Swan", "Sparkle",
    "Misha", "Acheron", "Aventurine", "Kakavasha", "Gallagher", "Robin", "Boothill", "Firefly", "Sam", "Jade",
    "Yunli", "Jiaoqiu", "Feixiao", "Lingsha", "Moze", "Rappa", "Sunday", "Fugue", "Herta", "Aglaea", "Tribbie",
    "Mydei", "Castorice", "Anaxa", "Hyacine", "Cipher", "Phainon", "Cerydra", "Hysilens", "Saber", "Archer",
    "Arlan", "Asta", "Caelus", "Stelle", "Pom-Pom", "Elio", "Cocolia", "Oswaldo", "Schneider", "Hoolay", "Phantylia",
    "Screwllum", "Stephen", "Constance", "Huaiyan", "Yingxing", "Baiheng", "Dan Feng", "Imbibitor", "Lunae",
    "Diamond", "Duke Inferno", "Tayzzyronth", "Nanook", "Qlipoth", "Akivili", "Yaoshi", "Mythus", "Kephale",
    "Nikador", "Oronyx", "Cyrene", "Sigonia", "Danheng", "Jingyuan", "Fuxuan", "Ruanmei", "Mortenax", "Xipe",
    "Idrila", "Lan", "Nous", "Aha", "IX", "Ena", "Terminus", "Kafuka", "Mozé", "Hanabi", "Dahlia", "Evernight",
    "Lingsha", "Sushang", "Tingyun", "Qingque", "Arlan", "Pela",
]
# palavras comuns do português (ou nomes do próprio livro, como "Frota de Jade" em 27.16) que só contam como nome
# proibido nas listas de NOMES; nas outras tabelas de sabor elas são texto
COMUNS = {"pela", "clara", "jade", "aha", "ix", "ena", "lan", "sam", "stephen", "nous", "terminus", "saber"}


def _norm(s):
    s = unicodedata.normalize("NFKD", str(s))
    return "".join(c for c in s if not unicodedata.combining(c)).lower()


def _tokens(s):
    return [t for t in re.split(r"[^a-z0-9]+", _norm(s)) if t]


def _proibido(s, nomes_lista=True):
    """O nome proibido que aparece em `s` como palavra (ou sequência de palavras), ou None. Nas listas de nomes,
    também como pedaço de um nome colado (Vidyadhara) com 5 letras ou mais."""
    tk = _tokens(s)
    junto = "".join(tk)
    for p in PROIBIDOS:
        pt = _tokens(p)
        if not nomes_lista and " ".join(pt) in COMUNS:
            continue
        n = len(pt)
        if any(tk[i:i + n] == pt for i in range(len(tk) - n + 1)):
            return p
        if nomes_lista and len("".join(pt)) >= 5 and "".join(pt) in junto:
            return p
    return None


def suite_sabor(args):
    import mestre_sabor as SB
    import mestre_sabor_nomes as NM
    r = Resultado("sabor")
    # autoteste do verificador
    for ruim in ("Kafka Vasquez", "Dan Heng", "Wenshu Jingliu", "Haibailuming", "Silver-Wolf"):
        r.ok(_proibido(ruim) is not None, f"autoteste: {ruim!r} deveria ser nome proibido")
    for bom in ("Teodora Okonkwo-Reis", "Rin Sete Caudas", "KV-12, que se chama Paciência"):
        r.ok(_proibido(bom) is None, f"autoteste: {bom!r} não é nome do jogo ({_proibido(bom)})")
    r.ok(SB.conferir(), "mestre_sabor.conferir()")
    mapa = TC.mapa()
    wb = openpyxl.load_workbook(N.SAIDA_MODELO)
    ws = wb["Tabelas"]
    for t in SB.tabelas():
        vs, id_ = t["valores"], t["id"]
        r.ok(len(vs) >= t["minimo"], f"tab.{id_}: {len(vs)} entradas (mínimo {t['minimo']})")
        for v in vs:
            r.ok(bool(v.strip()) and len(v) <= 200 and v[0] not in "=+-@", f"tab.{id_}: entrada inválida {v!r}")
            p = _proibido(v, nomes_lista=t["sem_ortografia"])
            r.ok(p is None, f"tab.{id_}: {v!r} usa o nome de personagem do jogo {p!r}")
        if not t["pesos"]:
            r.ok(len(set(vs)) == len(vs), f"tab.{id_}: entrada repetida")
        # a aba Tabelas tem exatamente o conteúdo de mestre_sabor, na ordem, e a marca sem_ortografia
        m = mapa["tabelas"][f"tab.{id_}"]
        lidos = [ws[c].value for c in m["vagas"] if ws[c].value not in (None, "")]
        r.ok(lidos == vs, f"tab.{id_}: aba Tabelas ≠ mestre_sabor ({lidos[:2]} × {vs[:2]})")
        r.ok(m["sem_ortografia"] == t["sem_ortografia"], f"tab.{id_}: marca sem_ortografia")
        if t["segunda"]:
            l2 = [ws[c].value for c in m["segunda"] if ws[c].value not in (None, "")]
            r.ok(l2 == t["segunda"][1], f"tab.{id_}: 2ª coluna ≠ mestre_sabor")
            for v in l2:
                p = _proibido(v, nomes_lista=False)
                r.ok(p is None, f"tab.{id_}: {v!r} usa o nome de personagem do jogo {p!r}")
    # nomes ≥ 30 por cultura, no estilo de 05; toda combinação nome + sobrenome livre de nome do jogo
    for c in NM.CULTURAS:
        r.ok(len(NM.NOMES[c]) >= 30 and len(NM.SOBRENOMES[c]) >= 30, f"{c}: menos de 30 nomes ou sobrenomes")
        ordem, sep = NM.MONTAGEM[c]
        ruins = [f"{a} {b}" for a in NM.NOMES[c] for b in NM.SOBRENOMES[c]
                 if _proibido(b + sep + a if ordem == 2 else a + sep + b)]
        r.ok(not ruins, f"{c}: combinações com nome do jogo: {ruins[:5]}")
        r.checagens += len(NM.NOMES[c]) * len(NM.SOBRENOMES[c])
    r.ok(all(re.fullmatch(r"[A-Z]{2}-\d{1,3}", x) for x in NM.NOMES["Intellitron"]),
         "Intellitron: designação alfanumérica (AA-00) fora do formato")
    r.ok(all(re.match(r"d(a|o|as|os) ", x) for x in NM.SOBRENOMES["Vidyadhara"]),
         "Vidyadhara: linhagem de mar fora do formato 'da/do …'")
    r.ok(all(len(x) <= 6 for x in NM.NOMES["Vulpes"] + NM.NOMES["Avginiano"]), "Vulpes/Avginiano: nome longo demais")
    r.ok(all(len(x) >= 9 for x in NM.NOMES["Vidyadhara"]), "Vidyadhara: nome curto demais (05: nomes longos)")
    # nomes do Exemplo (PJs e NPCs do Elenco)
    from mestre import exemplo
    d = exemplo.entradas_nomes()
    nomes_ex = [v for k, v in d.items() if re.fullmatch(r"grupo\.pj\d\.nome|npcs\.elenco\.\d+\.nome", k)]
    for v in nomes_ex:
        r.ok(_proibido(v) is None, f"Exemplo: {v!r} usa nome do jogo ({_proibido(v)})")
    # entradas do livro no início das listas (27.7, 25.2, 25.3), na ordem do livro
    for id_, (cap, sec, vals) in SB.LIVRO_NO_INICIO.items():
        texto = L.limpo(" ".join(_secao(cap, sec)))
        pos = [texto.find(v) for v in vals]
        r.ok(all(p >= 0 for p in pos) and pos == sorted(pos), f"tab.{id_}: entradas do livro fora de {sec}: {pos}")
        t = next(x for x in SB.tabelas() if x["id"] == id_)
        r.ok(t["valores"][:len(vals)] == vals, f"tab.{id_}: as primeiras entradas não são as do livro")
    # H11: pesos e chaves das tabelas de duas colunas
    cam = SB.CAMINHO_NPC
    r.ok(cam.count("Nenhum") == 3 and all(cam.count(c) == 1 for c in OH.CAMINHOS), "H11: Caminho do NPC (Nenhum ×3)")
    mc = [k for k in SB.MOTIVACAO_CAMINHO for _ in SB.MOTIVACAO_CAMINHO[k]]
    r.ok(sorted(set(mc)) == sorted(OH.CAMINHOS) and all(mc.count(c) == 5 for c in OH.CAMINHOS),
         "motivação por Caminho: 5 por Caminho, os 9 de 06.3")
    r.ok(sorted(SB.OBJETIVO) == sorted(SB.TIPO_AVENTURA) and all(len(v) == 3 for v in SB.OBJETIVO.values()),
         "objetivo: 3 por tipo de aventura")
    r.ok(SB.ATITUDE == ["Hostil", "Desconfiado", "Indiferente", "Cordial", "Prestativo"], "atitude: os 5 degraus de 6.6")
    r.info(f"{len(SB.tabelas())} tabelas de sabor; {len(PROIBIDOS)} nomes do jogo conferidos (palavra, sequência e "
           f"pedaço colado de 5+ letras), {sum(len(NM.NOMES[c]) * len(NM.SOBRENOMES[c]) for c in NM.CULTURAS)} "
           f"combinações nome + sobrenome")
    return r


def _secao(cap, titulo):
    linhas = L._md(cap)
    nivel = len(titulo) - len(titulo.lstrip("#"))
    i = next(k for k, l in enumerate(linhas) if l.startswith(titulo))
    fim = next((k for k in range(i + 1, len(linhas)) if re.match(r"^#{1,%d} " % nivel, linhas[k])), len(linhas))
    return linhas[i:fim]


def _tab(cap, titulo, primeira):
    """Tabela markdown da seção cujo cabeçalho começa por `primeira` (parser do teste)."""
    ls = _secao(cap, titulo)
    for k, l in enumerate(ls):
        if l.startswith("|") and k + 1 < len(ls) and re.match(r"^\|(\s*:?-{3,}:?\s*\|)+$", ls[k + 1].strip()):
            cab = L._cels(l)
            if cab[0].startswith(primeira):
                corpo = []
                for x in ls[k + 2:]:
                    if not x.startswith("|"):
                        break
                    corpo.append(L._cels(x))
                return cab, corpo
    raise ValueError(f"teste: tabela {primeira} não achada em {titulo}")


def _n(s):
    s = s.replace("Cr", "").replace(".", "").replace(",", ".").replace("+", "").strip()
    v = float(s)
    return int(v) if v.is_integer() else v


def _bloco(wb, id_):
    b = TC.mapa()["blocos"][f"dados.{id_}"]
    ws = wb["Dados"]
    return [{c: ws[f"{l}{r}"].value for c, l in b["colunas"].items()} for r in range(b["primeira_linha"],
                                                                                      b["ultima_linha"] + 1)]


def livro_fase2():
    """O que o .md diz (parser do teste) para as tabelas da Fase 2."""
    d = {}
    _, c = _tab("24", "## 24.1", "Tipo")
    d["armaduras"] = [(x[0], _n(x[1]), x[2], x[3], _n(x[4]), _n(x[5])) for x in c]
    _, c = _tab("24", "## 24.2", "Categoria")
    d["armas"] = [(re.sub(r"\s*\(2 mãos\)", "", x[0]), x[1], x[2], x[3], _n(x[5]), _n(x[6])) for x in c]
    _, c = _tab("24", "## 24.3", "Poção")
    d["pocoes"] = [(x[0], _n(x[1]), _n(x[2]), _n(x[3])) for x in c]
    _, c = _tab("24", "## 24.3", "Item")
    d["itens"] = [(x[0], _n(x[1]), _n(x[2]), x[3]) for x in c]
    _, c = _tab("24", "## 24.5", "Faixa de nível")
    d["verba"] = [(x[0], _n(x[1]), x[2]) for x in c]
    _, c = _tab("25", "## 25.1", "Faixa de nível")
    d["faixas_equipamento"] = [(x[0], x[1], x[2]) for x in c]
    _, c = _tab("25", "## 25.2", "Nível do Cone")
    d["cone"] = [(_n(x[0]), _n(x[1]), x[2], x[3]) for x in c]
    cab, c = _tab("25", "## 25.3", "Slot")
    d["tiers"] = [(m.group(1), int(m.group(2)), int(m.group(3))) for h in cab[2:]
                  for m in [re.match(r"(Tier [IV]+) \((\d+)-(\d+)\)", h)]]
    d["reliquias"] = [(x[0], x[1], *[_n(v) for v in x[2:6]]) for x in c]
    _, c = _tab("25", "## 25.3", "Peças do mesmo Conjunto")
    d["conjuntos"] = [(x[0], x[1]) for x in c]
    _, c = _tab("26", "## 26.7", "Ressonância")
    d["ressonancias"] = [(x[0], _n(x[1]), x[2]) for x in c]
    texto25 = L.limpo(" ".join(_secao("25", "## 25.2")))
    d["sobreposicao_texto"] = texto25
    d["bonus_maior"] = [L.limpo(l[2:]) for l in _secao("25", "### O Bônus Maior") if l.startswith("- ")]
    fac, atual = {}, None
    for l in _secao("27", "## 27.16"):
        if l.startswith("### "):
            atual = l[4:].strip()
        m = re.search(r"\*Fichas: (.+?)\.\*", l)
        if m:
            fac[atual] = m.group(1)
    d["fichas_texto"] = fac
    _, c = _tab("27", "## 27.17", "Lugar")
    d["locais"] = [(x[0], x[2]) for x in c]
    t7 = L.limpo(" ".join(_secao("27", "## 27.7")))
    m = re.search(r"objetivo que não seja matar todo mundo \((.+?)\)", t7)
    d["sem_matar"] = [x.strip() for x in m.group(1).split(",")]
    ls = _secao("27", "## 27.8")
    i = next(k for k, l in enumerate(ls) if "O que você não deve dar" in l)
    nd = []
    for l in ls[i + 1:]:
        if l.startswith("- "):
            nd.append(L.limpo(l[2:]))
        elif nd and l.strip():
            break
    d["nao_dar"] = nd
    return d


def dados_fase2(r, wb):
    lv = livro_fase2()
    n0 = r.checagens

    def conf(rot, a, b):
        r.ok(TF._igual(a, b) if not isinstance(a, (list, tuple)) else list(a) == list(b),
             f"{rot}: planilha/oráculo={a!r} livro={b!r}")

    for row, x in zip(_bloco(wb, "armaduras"), lv["armaduras"]):
        conf(f"24.1 armadura {x[0]}", (row["Tipo"], row["Defesa"], row["Outros efeitos"], row["Esquiva"], row["Espaço"],
                                       row["Preço (Cr)"]), x)
    for row, x in zip(_bloco(wb, "armas"), lv["armas"]):
        conf(f"24.2 arma {x[0]}", (row["Categoria"], row["Dados base"], row["Alcance"], row["Atributo de Ataque"],
                                   row["Espaço"], row["Preço (Cr)"]), x)
    for row, x in zip(_bloco(wb, "pocoes"), lv["pocoes"]):
        conf(f"24.3 poção {x[0]}", (row["Poção"], row["Cura"], row["Espaço"], row["Preço (Cr)"]), x)
    for row, x in zip(_bloco(wb, "itens"), lv["itens"]):
        conf(f"24.3 item {x[0]}", (row["Item"], row["Espaço"], row["Preço (Cr)"], row["Para quê"]), x)
    for nome, n in (("armaduras", 3), ("armas", 6), ("pocoes", 3), ("itens", 11)):
        r.ok(len(_bloco(wb, nome)) == len(lv[nome]) == n, f"{nome}: nº de linhas ≠ livro")
    # consumíveis dos achados: todo item de 24.3, com o preço do livro; a poção só na faixa de H9
    pre = {f"Poção {x[0]}": x[3] for x in lv["pocoes"]} | {x[0]: x[2] for x in lv["itens"]}
    cons = _bloco(wb, "consumiveis")
    r.ok(sorted(x["Item"] for x in cons) == sorted(pre), "consumíveis ≠ itens e poções de 24.3")
    for row in cons:
        conf(f"24.3 consumível {row['Item']}", row["Preço (Cr)"], pre.get(row["Item"]))
        if row["Item"].startswith("Poção"):
            a, b = OH.FAIXA_POCAO[row["Item"].split()[1]]
            for k, fx in enumerate(L.FAIXAS, start=1):
                r.ok((row[f"Ordem {fx}"] not in (None, "")) == (a <= k <= b), f"H9: {row['Item']} na faixa {fx}")
    for row, x in zip(_bloco(wb, "verba"), lv["verba"]):
        conf(f"24.5 verba {x[0]}", (row["Faixa"], row["Verba de marco (Cr)"], row["O que ela compra"]), x)
    for row, x in zip(_bloco(wb, "faixas_equipamento"), lv["faixas_equipamento"]):
        conf(f"25.1 {x[0]}", (row["Faixa"], row["Cone de Luz máximo"], row["Relíquias"]), x)
    for row, x in zip(_bloco(wb, "cone"), lv["cone"]):
        conf(f"25.2 Cone Nível {x[0]}", (row["Nível do Cone"], row["Nível do personagem"], row["Bônus Maior"],
                                         row["Efeito Condicional"]), x)
        m = re.match(r"\+(\d+) (ou|e) \+(\d+) PV", x[2])
        conf(f"25.2 Cone Nível {x[0]} partes", (row["Parte numérica"], row["Modo"], row["PV"]),
             (int(m.group(1)), m.group(2), int(m.group(3))))
    sob = {row["Nível do Cone"]: (row["Sobreposições até o teto"], row["PV no teto"]) for row in _bloco(wb, "cone")}
    t = lv["sobreposicao_texto"]
    r.ok("2 Sobreposições num Cone de Nível 1 ou 2 (a parte numérica chega a +3; a de PV, a +30), 1 num de Nível 3 ou 4 "
         "(+3; PV +35) e nenhuma num de Nível 5" in t, "25.2: frase do teto da Sobreposição não encontrada no livro")
    conf("25.2 teto da Sobreposição (Níveis 1–5)", [sob[k] for k in range(1, 6)],
         [(2, 30), (2, 30), (1, 35), (1, 35), (0, 50)])
    for row, x in zip(_bloco(wb, "reliquias"), lv["reliquias"]):
        conf(f"25.3 Relíquia {x[0]}", (row["Slot"], row["O que dá"], row["Tier I"], row["Tier II"], row["Tier III"],
                                       row["Tier IV"]), x)
    for row, x in zip(_bloco(wb, "conjuntos"), lv["conjuntos"]):
        conf(f"25.3 Conjunto {x[0]}", (row["Peças"], row["O que você ganha"]), x)
    for row, x in zip(_bloco(wb, "ressonancias"), lv["ressonancias"]):
        conf(f"26.7 Ressonância {x[0]}", (row["Ressonância"], row["Nível"], row["O que ela dá"]), x)
    # equipamento por nível (25.1 + faixas de Tier de 25.3 + 26.7 + 24.5)
    tiers = lv["tiers"]
    res = {x[1]: x[0] for x in lv["ressonancias"]}
    vb = {x[0]: x[1] for x in lv["verba"]}
    cone = {x[0]: int(re.search(r"\d", x[1]).group(0)) for x in lv["faixas_equipamento"]}
    for row in _bloco(wb, "nivel_equipamento"):
        n = row["Nível"]
        fx = L.FAIXAS[(n - 1) // 4]
        tr = next(t_ for t_, a, b in tiers if a <= n <= b)
        novo = any(a == n for _, a, _b in tiers)
        conf(f"nível {n}", (row["Faixa"], row["Cone de Luz máximo (Nível)"], row["Tier de Relíquia"],
                            row["Primeiro nível da faixa"], row["Tier novo neste nível"], row["Ressonância"],
                            row["Verba de marco (Cr)"]),
             (fx, cone[fx], tr, "Sim" if n % 4 == 1 else "Não", "Sim" if novo else "Não", res.get(n, "—"), vb[fx]))
    # 25.1 × 25.3: "virando Tier II no nível 7" e "virando Tier IV no nível 18" são as fronteiras de 25.3
    for fx, _, rel in lv["faixas_equipamento"]:
        m = re.search(r"virando (Tier [IV]+) no nível (\d+)", rel)
        if m:
            r.ok(any(t_ == m.group(1) and a == int(m.group(2)) for t_, a, _ in tiers), f"25.1 × 25.3: {rel}")
    # 27.16 fichas, 27.17 faixa dos lugares, 27.7, 27.8, 25.2 Bônus Maior, 25.3 2 peças
    ff = {}
    for row in _bloco(wb, "faccao_fichas"):
        ff.setdefault(row["Facção"], []).append(row["Criatura"])
    for fac, s in lv["fichas_texto"].items():
        conf(f"27.16 fichas de {fac}", ", ".join(ff.get(fac, [])), s)
        conf(f"27.16 oráculo {fac}", ", ".join(OH.FACCAO_FICHAS[fac]), s)
    for row, (lug, cena) in zip(_bloco(wb, "locais_faixa"), lv["locais"]):
        m = re.match(r"Faixa (\d+-\d+)\.", cena)
        esp = ((L.FAIXAS.index(m.group(1)) + 1,) * 2 if m else (4, 5) if cena.startswith("Faixa alta.") else (1, 5))
        conf(f"27.17 {lug}", (row["Lugar"], row["Faixa mínima"], row["Faixa máxima"]), (lug, *esp))
        conf(f"27.17 oráculo {lug}", OH.LOCAIS_FAIXA.get(lug, (1, 5)), esp)
    conf("27.7 objetivo que não seja matar", [x["Objetivo"] for x in _bloco(wb, "sem_matar")], lv["sem_matar"])
    conf("27.7 oráculo", OH.SEM_MATAR, lv["sem_matar"])
    conf("27.8 o que não dar", [x["O que não dar"] for x in _bloco(wb, "nao_dar")], lv["nao_dar"])
    bm = [x["Bônus Maior"] for x in _bloco(wb, "bonus_maior")]
    conf("25.2 Bônus Maior (oráculo)", OH.BONUS_MAIOR, bm)
    md = lv["bonus_maior"]
    r.ok(len(md) == 6 and md[:3] == bm[:3] and all(k in md[3] for k in ("Ataque Básico", "Habilidade", "Ultimate"))
         and md[4].startswith("RD") and all(k in md[5] for k in ("Teste de Ataque", "um Teste de Resistência",
                                                                 "uma Perícia")), f"25.2 Bônus Maior × livro: {md}")
    g2 = [x["Bônus de 2 peças"] for x in _bloco(wb, "conjunto_2")]
    r.ok(all(o in lv["conjuntos"][0][1] for o in g2) and len(g2) == 4 and g2 == OH.CONJUNTO_2, f"25.3 2 peças: {g2}")
    # tabelas transcritas do oráculo × livro (o oráculo não lê o .md em tempo de execução)
    conf("oráculo 24.1", [tuple(x) for x in OH.ARMADURAS], [(a, b, c, e, f) for a, b, c, _, e, f in lv["armaduras"]])
    conf("oráculo 24.2", [tuple(x) for x in OH.ARMAS], lv["armas"])
    conf("oráculo 24.3 poções", [tuple(x) for x in OH.POCOES], lv["pocoes"])
    conf("oráculo 24.3 itens", [tuple(x) for x in OH.ITENS], lv["itens"])
    conf("oráculo 24.5", OH.COMPRA, [x[2] for x in lv["verba"]])
    conf("oráculo 24.5 verba", list(L.VERBA), [x[1] for x in lv["verba"]])
    conf("oráculo 25.2", [(a, b) for a, b, _, _ in OH.CONE], [(x[2], x[3]) for x in lv["cone"]])
    conf("oráculo 25.3", [tuple(x) for x in OH.RELIQUIAS], lv["reliquias"])
    conf("oráculo 26.7", [(v[0], k, v[1]) for k, v in sorted(OH.RESSONANCIAS.items())], lv["ressonancias"])
    r.info(f"Fase 2: {r.checagens - n0} checagens item a item contra 24.1, 24.2, 24.3, 24.5, 25.1, 25.2, 25.3, 26.7, "
           f"27.7, 27.8, 27.16 e 27.17 ({len(lv['armaduras'])} armaduras, {len(lv['armas'])} armas, "
           f"{len(lv['pocoes'])} poções, {len(lv['itens'])} itens, {len(lv['cone'])} Níveis de Cone, "
           f"{len(lv['reliquias'])} slots × 4 Tiers, {len(lv['conjuntos'])} Conjuntos, {len(lv['ressonancias'])} "
           f"Ressonâncias)")


def ouro_fase2(r):
    from mestre.testes_regras import planilha
    lv = livro_fase2()

    def conf(rot, E, nome, esperado):
        v = planilha(E, [nome]).get(nome)
        r.ok(TF._igual(v, esperado), f"{rot}: planilha={v!r} livro={esperado!r}")
    # 24.5 verba por faixa e 25.1 Cone/Tier, 26.7 Ressonância, nível a nível (aba Recompensas, entrega de marco)
    tiers = lv["tiers"]
    res = {x[1]: x[0] for x in lv["ressonancias"]}
    for n in range(1, 21):
        fx = L.FAIXAS[(n - 1) // 4]
        got = planilha({"recompensas.nivel": n}, ["recompensas.marco.verba", "recompensas.marco.cone",
                                                  "recompensas.marco.tier", "recompensas.marco.ressonancia",
                                                  "recompensas.marco.bonus"])
        vb = next(x[1] for x in lv["verba"] if x[0] == fx)
        cone = next(x[1] for x in lv["faixas_equipamento"] if x[0] == fx)
        tr = next(t for t, a, b in tiers if a <= n <= b)
        bon = lv["cone"][L.FAIXAS.index(fx)][2]
        r.ok((got["recompensas.marco.verba"], got["recompensas.marco.cone"], got["recompensas.marco.tier"],
              got["recompensas.marco.ressonancia"], got["recompensas.marco.bonus"]) ==
             (vb, cone, tr, res.get(n, "—"), bon), f"24.5/25.1/25.3/26.7 nível {n}: {got}")
    # 25.2 Sobreposição: Nível 1 ou 2 → 2; 3 ou 4 → 1; 5 → nenhuma (aba Recompensas, a partir da aba Grupo)
    for c, teto in ((1, 2), (2, 2), (3, 1), (4, 1), (5, 0)):
        conf(f"25.2 teto da Sobreposição, Cone Nível {c}", {"grupo.pj1.nome": "Vesper", "grupo.pj1.cone": c},
             "recompensas.sob.1.teto", teto)
    # 25.2 exemplo: O Último Trem Para Casa, Nível 2, recebe 1 Sobreposição → "já recebeu a desta faixa"
    conf("25.2 exemplo do Vesper", {"grupo.pj1.nome": "Vesper", "grupo.pj1.cone": 2, "grupo.pj1.sobrep": 1},
         "recompensas.sob.1.situacao", "Já recebeu a desta faixa (máximo 1 por faixa, 25.2)")
    # 25.2: o teto é do Cone e conta as Sobreposições de faixas anteriores (trocar é opcional, E14). Cone 3 com 1
    # anterior e Cone 1 com 2 anteriores já estão em +3: "no teto", mesmo com 0 nesta faixa
    no_teto = "Cone no teto (+3) com as que já tem: nenhuma a mais (25.2)"
    pode = "Pode receber 1 nesta faixa, se a história entregar (25.2)"
    for c, ant, esperado in ((3, 1, no_teto), (4, 1, no_teto), (1, 2, no_teto), (2, 2, no_teto), (1, 1, pode),
                             (2, 0, pode), (3, 0, pode)):
        E = {"campanha.nivel": 17, "grupo.pj1.nome": "Vesper", "grupo.pj1.cone": c, "grupo.pj1.sobrep": 0, "grupo.pj1.sobrep_total": ant}
        got = planilha(E, ["recompensas.sob.1.situacao", "recompensas.sob.1.total", "recompensas.sob.1.aviso",
                           "grupo.pj1.aviso3"])
        r.ok((got["recompensas.sob.1.situacao"], got["recompensas.sob.1.total"], got["recompensas.sob.1.aviso"],
              got["grupo.pj1.aviso3"]) == (esperado, ant, "", ""),
             f"25.2 Cone Nível {c} com {ant} de faixas anteriores: {got}")
    # acima do teto do Cone: aviso nas abas Recompensas e Grupo
    for c, tot in ((3, 2), (1, 3), (5, 1)):
        teto = {1: 2, 3: 1, 5: 0}[c]
        E = {"campanha.nivel": 17, "grupo.pj1.nome": "Vesper", "grupo.pj1.cone": c, "grupo.pj1.sobrep_total": tot}
        got = planilha(E, ["recompensas.sob.1.aviso", "grupo.pj1.aviso3"])
        r.ok(got["recompensas.sob.1.aviso"] == f"Acima do teto do Cone: o Nível {c} aceita até {teto} (25.2)" and
             got["grupo.pj1.aviso3"] == f"Acima do teto do Cone (25.2): o Nível {c} aceita até {teto}",
             f"25.2 Cone Nível {c} com {tot} no total: {got}")
    # total menor que as desta faixa (o total conta as desta faixa): aviso; o total usado é o desta faixa
    E = {"grupo.pj1.nome": "Vesper", "grupo.pj1.cone": 1, "grupo.pj1.sobrep": 1, "grupo.pj1.sobrep_total": 0}
    got = planilha(E, ["recompensas.sob.1.aviso", "recompensas.sob.1.total", "grupo.pj1.aviso3"])
    r.ok((got["recompensas.sob.1.aviso"], got["recompensas.sob.1.total"], got["grupo.pj1.aviso3"]) ==
         ("Total no Cone menor que as desta faixa: confira na aba Grupo", 1, "Total no Cone menor que as desta faixa"),
         f"25.2 total menor que as desta faixa: {got}")
    # 24.1–24.3 preços na aba Recompensas (calculados) × livro
    nomes = [n for n in TC.mapa()["celulas"] if n.startswith("recompensas.preco.") and n.endswith(".preco")]
    got = planilha({}, nomes)
    esp = {**{f"recompensas.preco.armaduras.{i}.preco": x[5] for i, x in enumerate(lv["armaduras"], 1)},
           **{f"recompensas.preco.armas.{i}.preco": x[5] for i, x in enumerate(lv["armas"], 1)},
           **{f"recompensas.preco.pocoes.{i}.preco": x[3] for i, x in enumerate(lv["pocoes"], 1)},
           **{f"recompensas.preco.itens.{i}.preco": x[2] for i, x in enumerate(lv["itens"], 1)}}
    r.ok(set(esp) == set(nomes), "preços da aba Recompensas ≠ itens do livro")
    for n, v in esp.items():
        r.ok(got.get(n) == v, f"24.1–24.3 {n}: planilha={got.get(n)!r} livro={v!r}")
    # 25.3 Relíquias na aba Recompensas e o Tier do grupo (nível 7 → Tier II)
    rel = planilha({"campanha.nivel": 7}, [f"recompensas.rel.{i}.t{j}" for i in range(1, 7) for j in range(1, 5)] +
                   [f"recompensas.rel.{i}.agora" for i in range(1, 7)])
    for i, x in enumerate(lv["reliquias"], start=1):
        r.ok([rel[f"recompensas.rel.{i}.t{j}"] for j in range(1, 5)] == [f"+{v}" for v in x[2:]] and
             rel[f"recompensas.rel.{i}.agora"] == f"+{x[3]}", f"25.3 {x[0]}: {rel}")
    # H18 na faixa 1-4: Típico = 10% de 200 = 20 Cr
    conf("H18 achados 1-4 Típico", {"campanha.nivel": 1}, "recompensas.ach.creditos.valor", 20)
    r.info("27.9 (Ultimate por faixa) não aparece nas abas da Fase 2: o ouro da Fase 3 o calcula no Escudo do Mestre")


# ---------------------------------------------------------------------------
# oráculo — casos da Fase 2 e validações das heurísticas
# ---------------------------------------------------------------------------

PREF = ["npcs.", "aventuras.", "recompensas."]
PESO = 8      # custo medido de um caso da Fase 2 (≈ 10 s) em relação a um caso comum (≈ 1,3 s): reparte entre os 8 processos


def _grupo_cones(rng, E):
    for i in range(1, rng.randint(1, 6) + 1):
        E[f"grupo.pj{i}.nome"] = f"PJ{i}"
        if rng.random() < 0.8:
            E[f"grupo.pj{i}.cone"] = rng.choice([1, 2, 3, 4, 5, 7])
        if rng.random() < 0.6:
            E[f"grupo.pj{i}.sobrep"] = rng.choice([0, 1, 2])
        if rng.random() < 0.6:     # 25.2: total no Cone, com as de faixas anteriores (inclui acima do teto e texto)
            E[f"grupo.pj{i}.sobrep_total"] = rng.choice([0, 1, 2, 3, "duas"])


def casos_oraculo(r, rng, _t):
    import mestre_sabor as SB
    racas = L.RACAS
    for s in (1, 2026, 2147483646):
        for R in range(1, 301):
            E = {"inicio.semente": s, "npcs.rolagem": R, "aventuras.rolagem": R, "recompensas.ach.rolagem": R,
                 "recompensas.cone.rolagem": R, "recompensas.conj.rolagem": R}
            if rng.random() < 0.7:
                E["campanha.nivel"] = rng.randint(1, 20)
            if rng.random() < 0.4:
                E["campanha.jogadores"] = rng.randint(1, 6)
            par = {"npcs.raca": [None, "Sortear"] + racas, "npcs.caminho": [None, "Sortear", "Nenhum"] + OH.CAMINHOS,
                   "npcs.papel": [None, "Sortear"] + OH.PAPEIS,
                   "npcs.bloco_faixa": [None, "Nenhum", "Do grupo"] + L.FAIXAS, "npcs.bloco_tipo": [None, "Comum", "Elite"],
                   "aventuras.tipo": [None, "Sortear"] + SB.TIPO_AVENTURA, "aventuras.faixa": [None] + L.FAIXAS,
                   "aventuras.faccao": [None, "Sortear"] + list(OH.FACCAO_FICHAS),
                   "recompensas.ach.faixa": [None] + L.FAIXAS,
                   "recompensas.ach.leitura": [None, "Passagem", "Típico", "Pesado"],
                   "recompensas.cone.faixa": [None] + L.FAIXAS, "recompensas.nivel": [None] + list(range(1, 21)),
                   "recompensas.marco": [None, "Subiu de nível", "Entrou em faixa nova", "Fim de arco"]}
            for k, op in par.items():
                if rng.random() < 0.5:
                    v = rng.choice(op)
                    if v is not None:
                        E[k] = v
            if R % 10 == 0:
                _grupo_cones(rng, E)
            if R % 25 == 0:        # entradas inválidas (avisos nos dois sentidos)
                E.update({"npcs.rolagem": rng.choice([0, 2000000, "abc"]), "npcs.raca": "Elfo",
                          "aventuras.faixa": "21-24", "recompensas.nivel": 25, "recompensas.ach.leitura": "Fácil"})
            _t(E, f"Fase 2 S={s} R={R}", PREF, limite=4, peso=PESO)
    r.info("Fase 2: G = 200, 300, 400, 410 e 420 em 300 Rolagens nº × 3 sementes (900 casos), com os parâmetros de §5 "
           "sorteados e entradas inválidas a cada 25")
    # H22: toda facção × faixa (incluindo as sem Boss na faixa)
    for fac in OH.FACCAO_FICHAS:
        for fx in L.FAIXAS:
            _t({"aventuras.faccao": fac, "aventuras.faixa": fx, "aventuras.rolagem": rng.randint(1, 10 ** 6)},
               f"H22 {fac} {fx}", ["aventuras."], limite=4, peso=PESO)
    # Elenco cheio, cartões e Tesouro
    from mestre import exemplo
    npcs = exemplo.npcs_do_gerador({"inicio.semente": 77}, range(1, 31))
    for k in range(6):
        E = {}
        for i, npc in enumerate(npcs[: rng.randint(1, 30)], start=1):
            E.update({f"npcs.elenco.{i}.{c}": v for c, v in npc.items()})
            if rng.random() < 0.2:
                E[f"npcs.elenco.{i}.relacao"] = rng.choice([-4, -3, 0, 2, 3, 4, 1.5, "muito"])
        if k == 1:
            E["npcs.elenco.2.nome"] = E["npcs.elenco.1.nome"]
            E["npcs.elenco.3.nome"] = "Casco Oco"
        for c in range(1, 5):
            E[f"npcs.cartao{c}.npc"] = rng.choice([npcs[0]["nome"], npcs[1]["nome"], "Ninguém", None])
        for i in range(1, 26):
            if rng.random() < 0.6:
                E[f"recompensas.tes.{i}.cr"] = rng.choice([200, -150, 50, -500, "dez"])
        _t({k_: v for k_, v in E.items() if v is not None}, f"Elenco e Tesouro {k + 1}", ["npcs.", "recompensas."],
           peso=PESO)
    validar_heuristicas(r)


def validar_heuristicas(r):
    """H9, H14, H18, H20, H22, H25 e H26 no oráculo (o oráculo é conferido contra a planilha nos casos acima)."""
    import mestre_sabor as SB
    S0 = 12345
    # H20 — fonte do gancho 50/50 (NPC: C = 14; aventura: C = 15) em 2 000 rolagens; gancho "livro" é da faixa
    for g, c in ((200, 14), (300, 15)):
        n1 = sum(O.inteiro(O.valor(S0, g, R, c), 1, 2) == 1 for R in range(1, 2001))
        r.ok(900 <= n1 <= 1100, f"H20 G={g}: fonte 'livro' em {n1 / 20:.1f}% (45% a 55%)")
        r.info(f"H20 G={g}: fonte 'livro' {n1 / 20:.1f}% e 'tabela' {(2000 - n1) / 20:.1f}% em 2 000 rolagens")
    gl = {fx: OH.lista(f"ganchos_{fx.replace('-', '_')}") for fx in L.FAIXAS}
    for k in range(200):
        E = {"campanha.nivel": 1 + k % 20, "npcs.rolagem": k + 1, "aventuras.rolagem": k + 1}
        out = O.calcular(E)
        fx = L.FAIXAS[(k % 20) // 4]
        if out["npcs.g.fonte_gancho"] == 1:
            r.ok(out["npcs.res.gancho"] in gl[fx], f"H20: gancho do NPC fora dos 4 de 27.18 da faixa {fx}")
        if out["aventuras.g.fonte_gancho"] == 1:
            r.ok(out["aventuras.res.gancho"] in gl[fx], f"H20: gancho da aventura fora de 27.18 {fx}")
        # H14: orçamentos 50% e 100% da faixa; DT Média e Difícil (27.2)
        fi = L.FAIXAS.index(fx) + 1
        npj = out["campanha.jogadores_ef"]
        orc = L.ORC[fi - 1][1] * (1 if npj == 4 else npj) // (1 if npj == 4 else 4)
        r.ok(out["aventuras.cena3.orc"] == orc // 2 and out["aventuras.cena5.orc"] == orc, "H14: orçamento das cenas")
        r.ok((out["aventuras.cena1.dt"], out["aventuras.cena5.dt"]) == (L.DT["Média"][fi - 1], L.DT["Difícil"][fi - 1]),
             "H14: DT das cenas")
        # H25: os 2 traços de aparência são distintos; H26: o contratante não é da facção do conflito
        ap = out["npcs.res.aparencia"].split("; ")
        r.ok(len(ap) == 2 and ap[0] != ap[1], f"H25: aparência {ap}")
        r.ok(out["aventuras.g.c_fac"] != out["aventuras.g.fc"], "H26: contratante da mesma facção do conflito")
    # H22: para cada facção × faixa, o antagonista é do primeiro degrau que tem candidato; nunca vazio
    fichas = L.bestiario_md()
    lin = []
    for fac, nomes in OH.FACCAO_FICHAS.items():
        for fx in L.FAIXAS:
            out = O.calcular({"aventuras.faccao": fac, "aventuras.faixa": fx, "aventuras.rolagem": 3})
            bf = [f["nome"] for f in fichas if f["tipo"] == "Boss" and f["faixa"] == fx and f["nome"] in nomes]
            ef = [f["nome"] for f in fichas if f["tipo"] == "Elite" and f["faixa"] == fx and f["nome"] in nomes]
            ba = [f["nome"] for f in fichas if f["tipo"] == "Boss" and f["faixa"] == fx]
            deg = 1 if bf else 2 if ef else 3
            cand = (bf, ef, ba)[deg - 1]
            r.ok(out["aventuras.res.antagonista"] != "" and out["aventuras.g.an_degrau"] == deg and
                 out["aventuras.g.an_nome"] in cand, f"H22 {fac} {fx}: {out['aventuras.res.antagonista']!r} (degrau "
                                                      f"{deg}, candidatos {cand})")
            lin.append(deg)
    r.info(f"H22: {len(lin)} combinações facção × faixa, nenhum antagonista vazio; degrau 1 (Boss da facção) em "
           f"{lin.count(1)}, 2 (Elite da facção) em {lin.count(2)}, 3 (Boss da faixa) em {lin.count(3)}")
    # H9 — 0 a 2 consumíveis, cada valor entre 28% e 39% em 3 000 rolagens; item e preço de 24.3; poção da faixa
    cont = [0, 0, 0]
    for R in range(1, 3001):
        cont[O.inteiro(O.valor(S0, 400, R, 1), 0, 2)] += 1
    r.ok(all(840 <= c <= 1170 for c in cont), f"H9: quantidade 0/1/2 em {[c / 30 for c in cont]}% (28% a 39%)")
    r.info(f"H9: 0, 1 e 2 consumíveis em {cont[0] / 30:.1f}%, {cont[1] / 30:.1f}% e {cont[2] / 30:.1f}% de 3 000 "
           f"rolagens")
    precos = {f"Poção {a}": d for a, _, _, d in OH.POCOES} | {a: c for a, _, c, _ in OH.ITENS}
    for k in range(300):
        fx = L.FAIXAS[k % 5]
        out = O.calcular({"recompensas.ach.rolagem": k + 1, "recompensas.ach.faixa": fx,
                          "recompensas.ach.leitura": ["Passagem", "Típico", "Pesado"][k % 3],
                          "aventuras.rolagem": k + 1, "aventuras.faixa": fx})
        for j in (1, 2):
            it = out[f"recompensas.ach.item{j}"]
            if it:
                r.ok(precos.get(it) == out[f"recompensas.ach.item{j}.valor"], f"H9: {it} com preço fora de 24.3")
                if it.startswith("Poção"):
                    a, b = OH.FAIXA_POCAO[it.split()[1]]
                    r.ok(a <= k % 5 + 1 <= b, f"H9: {it} fora da faixa {fx}")
        # H18 — créditos ≤ 15% da verba, inteiro
        cr, vb = out["recompensas.ach.creditos.valor"], L.VERBA[k % 5]
        r.ok(isinstance(cr, int) and cr <= vb * 15 // 100 and cr == vb * (5, 10, 15)[k % 3] // 100, f"H18: {cr} de {vb}")
    erro = max(abs(v * f / 100 - v * f // 100) for v in L.VERBA for f in (5, 10, 15))
    r.info(f"H18: créditos de 5%/10%/15% da verba de marco (24.5); erro máximo de arredondamento (INT) {erro:g} Cr; "
           f"teto num nível com o ritmo de 26.1 (até 4 sessões × 3 combates de 27.7): 12 × 15% = 180% da verba de marco "
           f"— informativo, porque nada no balanceamento depende de Créditos (24.5)")
    # H26 — o lugar do livro só aparece na faixa que o livro diz (Poço Sete 1-4, Ferro-Vazio alta, Hesperin 17-20)
    for k in range(400):
        fx = L.FAIXAS[k % 5]
        out = O.calcular({"aventuras.rolagem": k + 1, "aventuras.faixa": fx})
        loc = out["aventuras.res.local"]
        if loc in OH.LOCAIS_FAIXA:
            a, b = OH.LOCAIS_FAIXA[loc]
            r.ok(a <= k % 5 + 1 <= b, f"H26: {loc} na faixa {fx}")
    r.ok(set(OH.lista("tipo_aventura")) == set(SB.TIPO_AVENTURA), "lista de tipos lida da planilha ≠ mestre_sabor")
