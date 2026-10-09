# -*- coding: utf-8 -*-
"""
renderizar_ficha.py — Prévia visual da Ficha de Personagem automatizada
("Explorando Galáxias" v1.0) em PNG, com detecção de texto cortado.

O QUE ELE FAZ
    Desenha cada aba do .xlsx gerado por `build/gerar_ficha.py` com pillow,
    respeitando larguras de coluna, alturas de linha, mesclas, preenchimento,
    bordas, fonte Arial (normal/negrito/itálico, tamanho e cor) e alinhamento.
    Os valores das fórmulas NÃO são lidos do arquivo (ele não tem cache): quem
    chama passa os valores calculados pela biblioteca `formulas` (a suíte
    `preview` de build/testar_ficha.py faz isso para 4 estados: ficha em
    branco, Nadir nível 1 de 29.7, um personagem nível 20 completo e o pior
    caso de estado_pior_caso()). A suíte `visual` usa a régua medir() daqui.

    Formatação condicional: só a dos avisos é simulada (fundo rosa #FFC7CE
    quando a célula de aviso tem texto); as outras regras não são desenhadas.

    Texto cortado (o que conta como achado):
      - com quebra de linha: alguma palavra mais larga que a célula, ou as
        linhas não cabem na altura da linha (ou da mescla);
      - sem quebra: o texto passa da largura útil e não pode transbordar —
        texto alinhado à esquerda transborda para as células vazias à direita,
        como no Google Planilhas; célula com valor ou fórmula (mesmo que dê "")
        bloqueia; número, texto centralizado ou à direita não transbordam;
      - a fonte é mais alta que a linha.
    A medida usa ImageFont.truetype("arial.ttf").getlength, com 3 px de
    margem de cada lado.

    Conversões: largura de coluna em caracteres -> px = int(largura × 7 + 5)
    (padrão 8,43); altura em pontos -> px = pontos × 4/3 (padrão da aba).

    Saída por aba e estado:
      - <estado>-<aba>.png: a aba inteira, para registro (reduzida só se passar
        de MAX_LADO px, e aí o texto fica pequeno demais para ler);
      - <estado>-<aba>-parte-NN.png: RECORTES em escala 1:1 (nunca reduzidos),
        de no máximo LADO_RECORTE × LADO_RECORTE px cada, cortados na divisa
        de coluna/linha e com as letras das colunas e os números das linhas
        repetidos em cada recorte. É neles que a inspeção visual é feita;
      - indice-recortes.txt: qual intervalo de células cada recorte mostra.

COMO USAR
    Normalmente pela bateria:
        $env:PYTHONUTF8="1"; python "build\\testar_ficha.py" --suite preview
    Direto (faz o mesmo que a suíte preview):
        $env:PYTHONUTF8="1"; python "build\\renderizar_ficha.py"

REQUISITOS
    Python 3, openpyxl 3.1.5, pillow 12.1.0 (já instalados; não instalar nada).
"""

from pathlib import Path

from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont

MARGEM = 3                     # px de cada lado dentro da célula
CAB_COL, CAB_LIN = 20, 36      # faixas com as letras das colunas e os números das linhas
MAX_LADO = 6000                # PNG inteiro maior que isso é gravado reduzido (só o de registro)
LADO_RECORTE = 1800            # lado máximo de cada recorte 1:1 (com as faixas de cabeçalho)
COR_GRADE = (226, 226, 226)
COR_CAB = (245, 245, 245)
COR_AVISO_FUNDO = (255, 199, 206)
ARQ_FONTE = {(False, False): "arial.ttf", (True, False): "arialbd.ttf",
             (False, True): "ariali.ttf", (True, True): "arialbi.ttf"}

_fontes = {}


def fonte(px, negrito=False, italico=False):
    chave = (px, bool(negrito), bool(italico))
    if chave not in _fontes:
        _fontes[chave] = ImageFont.truetype(ARQ_FONTE[(bool(negrito), bool(italico))], max(1, px))
    return _fontes[chave]


def px_coluna(ws, c):
    d = ws.column_dimensions[get_column_letter(c)]
    return 0 if d.hidden else int((d.width or 8.43) * 7 + 5)


def px_linha(ws, r):
    d = ws.row_dimensions[r]
    if d.hidden:
        return 0
    return round((d.height or ws.sheet_format.defaultRowHeight or 15) * 4 / 3)


def _rgb(cor, padrao=None):
    """'FF1F3864' / '1F3864' -> (31, 56, 100); cor de tema/índice -> padrão."""
    try:
        v = cor.rgb if hasattr(cor, "rgb") else cor
    except Exception:  # noqa: BLE001 — cor de tema sem rgb
        return padrao
    if not isinstance(v, str) or len(v) not in (6, 8):
        return padrao
    v = v[-6:]
    try:
        return tuple(int(v[i:i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return padrao


def texto_exibido(v):
    """Valor da célula como o Google Planilhas mostraria (pt-BR)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "VERDADEIRO" if v else "FALSO"
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return f"{v:.2f}".rstrip("0").rstrip(".").replace(".", ",")
    return str(v)


def quebrar(texto, f, largura):
    """Quebra por palavras (como o Google): devolve a lista de linhas."""
    linhas = []
    for par in texto.split("\n"):
        atual = ""
        for p in par.split(" "):
            teste = p if not atual else atual + " " + p
            if not atual or f.getlength(teste) <= largura:
                atual = teste
            else:
                linhas.append(atual)
                atual = p
        linhas.append(atual)
    return linhas


# ---------------------------------------------------------------------------
# Régua da suíte visual (e de build/gerar_ficha.py ao ajustar o layout): o texto tem de caber
# na PRÓPRIA célula ou mescla, sem transbordar para as vizinhas, com 10% de folga porque o
# Google Planilhas desenha a Arial um pouco mais larga que o pillow.
# ---------------------------------------------------------------------------

FOLGA = 1.10                   # 10% de folga na largura medida
RESERVA = 8                    # px da célula que o texto não usa (margens internas)
SETA_LISTA = 24                # px da seta da lista suspensa no Google
COR_CAB_COLUNA = "B4C6E7"      # cabeçalho de coluna (gerar_ficha.COR_CAB_COLUNA)
COR_TITULO = "1F3864"
COR_ENTRADA = "FFF2CC"
COR_ENTRADA_ZEBRA = "FFE9B0"


def cor_fundo(cel):
    f = cel.fill
    if f is None or f.patternType != "solid":
        return None
    return str(f.fgColor.rgb or "")[-6:].upper() or None


def eh_entrada(cel):
    return cor_fundo(cel) in (COR_ENTRADA, COR_ENTRADA_ZEBRA)


def fonte_da_celula(cel):
    ft = cel.font
    pt = (ft.sz if ft is not None and ft.sz else 10)
    return fonte(round(pt * 4 / 3), bool(ft and ft.b), bool(ft and ft.i))


def area_px(ws, r, c, topo):
    """(largura, altura) em px da célula ou da mescla que começa em (r, c)."""
    r2, c2 = topo.get((r, c), (r, c))
    return (sum(px_coluna(ws, k) for k in range(c, c2 + 1)), sum(px_linha(ws, k) for k in range(r, r2 + 1)))


def medir(ws, r, c, s, topo, quebra=None):
    """Confere se o texto `s` cabe na célula/mescla (r, c) pela régua da suíte visual.
    Devolve None se cabe ou o motivo se não cabe."""
    cel = ws.cell(r, c)
    f = fonte_da_celula(cel)
    asc, desc = f.getmetrics()
    lh = asc + desc
    w, h = area_px(ws, r, c, topo)
    util = w - RESERVA
    if quebra is None:
        quebra = bool(cel.alignment is not None and cel.alignment.wrap_text)
    if quebra:
        linhas = quebrar(s, f, util / FOLGA)
        larga = [ln for ln in linhas if f.getlength(ln) * FOLGA > util]
        if larga:
            return f"palavra de {int(f.getlength(larga[0]) * FOLGA)} px > {util} px: {larga[0][:30]!r}"
        if len(linhas) * lh + 2 > h:                 # +2 px: a mesma régua do texto cortado da prévia
            return f"{len(linhas)} linha(s) × {lh} px + 2 = {len(linhas) * lh + 2} px > {h} px de altura"
        return None
    linhas = s.split("\n")
    maior = max(f.getlength(ln) for ln in linhas) * FOLGA
    if maior > util:
        return f"{int(maior)} px de texto > {util} px úteis (sem quebra de linha)"
    if len(linhas) * lh > h:
        return f"{len(linhas)} linha(s) × {lh} px > {h} px de altura"
    return None


def tabelas(ws):
    """Tabelas da aba: linha de cabeçalho de coluna (fundo COR_CAB_COLUNA) e as linhas de
    dados logo abaixo, até a linha em branco, um título de bloco ou o próximo cabeçalho.
    Devolve [{"cab": linha, "c0": col, "c1": col, "linhas": [linhas de dados]}]."""
    topo, coberta = _mesclas(ws)

    def eh_cab(r):
        if str(ws.cell(r, 1).value or "").startswith("Legenda"):
            return []
        return [c for c in range(1, ws.max_column + 1)
                if (r, c) not in coberta and ws.cell(r, c).value is not None
                and cor_fundo(ws.cell(r, c)) == COR_CAB_COLUNA]

    saida = []
    r = 1
    while r <= ws.max_row:
        cols = eh_cab(r)
        if not cols:
            r += 1
            continue
        c0, c1 = min(cols), max(topo.get((r, c), (r, c))[1] for c in cols)
        linhas = []
        k = r + 1
        while k <= ws.max_row and not eh_cab(k):
            cels = [ws.cell(k, c) for c in range(c0, c1 + 1)]
            if any(cor_fundo(x) == COR_TITULO for x in cels):
                break
            if not any(x.value is not None or eh_entrada(x) for x in cels
                       if (x.row, x.column) not in coberta):
                break
            linhas.append(k)
            k += 1
        saida.append({"cab": r, "c0": c0, "c1": c1, "linhas": linhas})
        r = k
    return saida


def luminancia(rgb):
    def canal(v):
        v /= 255
        return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
    r, g, b = (canal(x) for x in rgb)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def contraste(a, b):
    la, lb = sorted((luminancia(a), luminancia(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


# ---------------------------------------------------------------------------
# Texto mais longo que uma fórmula pode mostrar (avisos e textos buscados na aba Dados)
# ---------------------------------------------------------------------------

def _arvore(formula):
    """Árvore da fórmula a partir dos tokens do openpyxl (Tokenizer)."""
    from openpyxl.formula.tokenizer import Tokenizer, Token
    toks = [t for t in Tokenizer(formula).items if t.type != Token.WSPACE]
    prec = {"=": 1, "<>": 1, "<": 1, ">": 1, "<=": 1, ">=": 1, "&": 2, "+": 3, "-": 3, "*": 4, "/": 4, "^": 5}
    pos = [0]

    def prox():
        return toks[pos[0]] if pos[0] < len(toks) else None

    def primario():
        t = toks[pos[0]]
        pos[0] += 1
        if t.type == Token.OPERAND:
            return ("lit", t.subtype, t.value)
        if t.type == Token.FUNC and t.subtype == Token.OPEN:
            nome, args = t.value[:-1].upper(), []
            if prox() is not None and prox().type == Token.FUNC and prox().subtype == Token.CLOSE:
                pos[0] += 1
                return ("fn", nome, args)
            while True:
                args.append(expr(0))
                t2 = toks[pos[0]]
                pos[0] += 1
                if t2.type == Token.FUNC and t2.subtype == Token.CLOSE:
                    return ("fn", nome, args)
        if t.type == Token.PAREN and t.subtype == Token.OPEN:
            e = expr(0)
            pos[0] += 1
            return e
        if t.type == Token.OP_PRE:
            return ("pre", t.value, expr(6))
        return ("lit", "TEXT", '""')

    def expr(minimo):
        esq = primario()
        while True:
            t = prox()
            if t is not None and t.type == Token.OP_POST:
                pos[0] += 1
                esq = ("pos", t.value, esq)
                continue
            if t is None or t.type != Token.OP_IN or prec.get(t.value, 0) < minimo or not prec.get(t.value):
                return esq
            pos[0] += 1
            dir_ = expr(prec[t.value] + 1)
            esq = ("op", t.value, esq, dir_)

    return expr(0)


class PiorTexto:
    """Estima o texto MAIS LONGO que cada célula pode mostrar: literais da fórmula (IF e
    IFERROR = o maior dos ramos; & = concatena), texto mais longo da coluna buscada por
    INDEX/VLOOKUP, maior opção da lista de uma entrada e o texto de pior caso das entradas
    livres. Conta e número viram "999"."""

    NUMERO = "999"

    def __init__(self, wb, entradas_texto=None, opcoes=None):
        self.wb = wb
        self.textos = entradas_texto or {}          # "'Aba'!A1" -> texto de pior caso
        self.opcoes = opcoes or {}                  # "'Aba'!A1" -> [opções da lista]
        self.memo = {}
        self.pilha = set()

    @staticmethod
    def _ref(aba, valor):
        if "!" in valor:
            a, cel = valor.rsplit("!", 1)
            aba = a.strip("'")
        else:
            cel = valor
        return aba, cel.replace("$", "")

    def celula(self, aba, cel):
        chave = f"'{aba}'!{cel}"
        if chave in self.memo:
            return self.memo[chave]
        if chave in self.textos:
            return self.textos[chave]
        if chave in self.opcoes:
            return max(self.opcoes[chave], key=len, default="")
        if chave in self.pilha or aba not in self.wb.sheetnames:
            return self.NUMERO
        v = self.wb[aba][cel].value
        if isinstance(v, str) and v.startswith("="):
            self.pilha.add(chave)
            s = self.formula(v, aba)
            self.pilha.discard(chave)
        else:
            s = texto_exibido(v)
        self.memo[chave] = s
        return s

    def intervalo(self, aba, faixa):
        from openpyxl.utils.cell import range_boundaries
        try:
            c0, r0, c1, r1 = range_boundaries(faixa)
        except (ValueError, TypeError):
            return self.NUMERO
        if aba not in self.wb.sheetnames or (r1 - r0 + 1) * (c1 - c0 + 1) > 5000:
            return self.NUMERO
        ws = self.wb[aba]
        melhor = ""
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                s = self.celula(aba, ws.cell(r, c).coordinate)
                if len(s) > len(melhor):
                    melhor = s
        return melhor

    def formula(self, f, aba):
        try:
            return self._pior(_arvore(f), aba)
        except Exception:  # noqa: BLE001 — fórmula que o parser não entende: sem estimativa
            return ""

    def _pior(self, n, aba):
        tipo = n[0]
        if tipo == "lit":
            sub, v = n[1], n[2]
            if sub == "TEXT":
                return v[1:-1].replace('""', '"')
            if sub == "NUMBER":
                return v
            if sub == "RANGE":
                a, cel = self._ref(aba, v)
                return self.celula(a, cel) if ":" not in cel else self.NUMERO
            return ""
        if tipo == "op":
            if n[1] == "&":
                return self._pior(n[2], aba) + self._pior(n[3], aba)
            return "" if n[1] in ("=", "<>", "<", ">", "<=", ">=") else self.NUMERO
        if tipo in ("pre", "pos"):
            return self.NUMERO
        nome, args = n[1], n[2]
        maior = lambda xs: max((self._pior(x, aba) for x in xs), key=len, default="")  # noqa: E731
        if nome == "IF":
            return maior(args[1:3])
        if nome == "IFERROR":
            return maior(args[:2])
        if nome == "CHOOSE":
            return maior(args[1:])
        if nome in ("INDEX", "VLOOKUP", "HLOOKUP"):
            alvo = args[0] if nome == "INDEX" else (args[1] if len(args) > 1 else None)
            if alvo is not None and alvo[0] == "lit" and alvo[1] == "RANGE":
                a, faixa = self._ref(aba, alvo[2])
                if ":" in faixa and nome != "INDEX" and len(args) > 2 and args[2][0] == "lit" and args[2][1] == "NUMBER":
                    # VLOOKUP/HLOOKUP com índice fixo: só a coluna (ou linha) devolvida conta
                    # (auditoria visual final; antes valia o texto mais longo da tabela inteira)
                    from openpyxl.utils import get_column_letter as _L
                    from openpyxl.utils.cell import range_boundaries
                    try:
                        c0, r0, c1, r1 = range_boundaries(faixa)
                        k = int(float(args[2][2]))
                        if nome == "VLOOKUP" and 1 <= k <= c1 - c0 + 1:
                            faixa = f"{_L(c0 + k - 1)}{r0}:{_L(c0 + k - 1)}{r1}"
                        elif nome == "HLOOKUP" and 1 <= k <= r1 - r0 + 1:
                            faixa = f"{_L(c0)}{r0 + k - 1}:{_L(c1)}{r0 + k - 1}"
                    except (ValueError, TypeError):
                        pass
                return self.intervalo(a, faixa) if ":" in faixa else self.celula(a, faixa)
            return self._pior(alvo, aba) if alvo is not None else ""
        if nome in ("TRIM",):
            return self._pior(args[0], aba)
        if nome == "CONCATENATE":
            return "".join(self._pior(x, aba) for x in args)
        return self.NUMERO


# ---------------------------------------------------------------------------
# 4º estado da prévia: "pior-caso" (listas na opção mais longa e textos longos e realistas)
# ---------------------------------------------------------------------------

# Tamanho (caracteres) do texto de pior caso de cada entrada livre, pelo nome lógico. As
# outras entradas livres usam o tamanho da regra que casar primeiro em PIOR_CASO_REGRAS.
PIOR_CASO_TAMANHO = {
    "criacao.nome": 40, "criacao.jogador": 30, "criacao.conceito": 80, "criacao.proposito": 200,
    "criacao.aparencia": 300, "criacao.anotacoes": 1000,
}
PIOR_CASO_REGRAS = [                 # (sufixo do nome lógico, caracteres)
    ("habilidades.ultimate.efeito", 300), (".efeito", 300), (".efeito4", 300), (".texto", 300),
    (".nome", 40), (".descricao", 40), (".item", 40), (".frase", 80), ("crenca_esforco", 80),
    ("coisa_inutil", 80),
]


def _texto_do_livro():
    """Texto corrido do capítulo 01 do livro (sem marcação), fonte dos textos de pior caso."""
    import re
    raiz = Path(__file__).resolve().parent.parent / "livro-v1.0"
    arq = sorted(raiz.glob("01*.md"))
    linhas = arq[0].read_text(encoding="utf-8").splitlines() if arq else []
    corpo = " ".join(ln for ln in linhas if ln.strip() and not ln.startswith(("#", "---", "|")))
    corpo = re.sub(r"[*_`]", "", corpo)
    corpo = re.sub(r"^\d+\.\s*", "", corpo)
    return re.sub(r"\s+", " ", corpo).strip() or "Explorando Galáxias é um RPG de mesa. " * 40


def texto_pior_caso(n, deslocamento=0):
    """Exatamente `n` caracteres de texto real do livro, começando na palavra `deslocamento`."""
    global _LIVRO
    if _LIVRO is None:
        _LIVRO = _texto_do_livro()
    palavras = _LIVRO.split(" ")
    inicio = deslocamento % max(1, len(palavras) - 1)
    corrido = " ".join(palavras[inicio:] + palavras)
    while len(corrido) < n + 60:           # texto curto demais: repete
        corrido = corrido + " " + corrido
    for k in range(60):                    # sem espaço na ponta (o Google não mostra)
        s = corrido[k:k + n]
        if not s.startswith(" ") and not s.endswith(" "):
            return s
    return corrido[:n]


_LIVRO = None


def tamanho_pior_caso(nome):
    if nome in PIOR_CASO_TAMANHO:
        return PIOR_CASO_TAMANHO[nome]
    for suf, n in PIOR_CASO_REGRAS:
        if nome.endswith(suf):
            return n
    return 40


def opcoes_da_lista(wb, info, ref, valores=None):
    """Opções da lista suspensa de uma entrada: lista da aba Dados, intervalo local (valores
    calculados em `valores`) ou literal "a,b"."""
    from openpyxl.utils.cell import range_boundaries
    fonte_ = info.get("fonte") or ""
    if fonte_.startswith('"'):
        return [x for x in fonte_.strip('"').split(",") if x]
    aba = ref.rsplit("!", 1)[0].strip("'")
    if fonte_.startswith("lista."):
        return None                        # quem chama resolve pelo mapa (aba Dados)
    if fonte_.startswith("$"):
        c0, r0, c1, r1 = range_boundaries(fonte_.replace("$", ""))
        ws = wb[aba]
        saida = []
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                k = f"'{aba}'!{ws.cell(r, c).coordinate}"
                v = (valores or {}).get(k, ws.cell(r, c).value)
                if isinstance(v, str) and v.startswith("="):
                    v = None
                s = texto_exibido(v)
                if s:
                    saida.append(s)
        return saida
    return []


def estado_pior_caso(wb, mapa, calcular):
    """Entradas do estado "pior-caso" ({"'Aba'!A1": valor}): cada lista suspensa na opção mais
    longa (medida em px na fonte da célula), números no máximo da validação e textos longos
    e realistas do livro nas entradas livres (tamanhos em PIOR_CASO_TAMANHO e
    PIOR_CASO_REGRAS). `calcular(entradas)` devolve os valores calculados: as listas
    dependentes (intervalos locais) são resolvidas numa 2ª passada, com as opções que a 1ª
    passada calculou."""
    C, listas = mapa["celulas"], mapa.get("listas", {})
    dados = wb["Dados"]
    from openpyxl.utils.cell import range_boundaries

    def lista_dados(id_):
        faixa = listas[id_].split("!", 1)[1].replace("$", "")
        c0, r0, c1, r1 = range_boundaries(faixa)
        return [texto_exibido(dados.cell(r, c0).value) for r in range(r0, r1 + 1)
                if dados.cell(r, c0).value not in (None, "")]

    def mais_longa(ref, opcoes):
        aba, cel = ref.rsplit("!", 1)
        f = fonte_da_celula(wb[aba.strip("'")][cel.replace("$", "")])
        return max(opcoes, key=lambda s: (f.getlength(s), s)) if opcoes else None

    entradas, dependentes = {}, []
    for k, (nome, info) in enumerate(sorted(mapa["entradas"].items())):
        ref = C[nome]
        tipo = info["tipo"]
        if tipo == "lista":
            fonte_ = info.get("fonte") or ""
            if nome.endswith(".item"):                       # inventário: item livre de 40 caracteres
                entradas[ref] = texto_pior_caso(40, 7 * k)
            elif fonte_.startswith("lista."):
                entradas[ref] = mais_longa(ref, lista_dados(fonte_))
            elif fonte_.startswith('"'):
                entradas[ref] = mais_longa(ref, opcoes_da_lista(wb, info, ref))
            else:
                dependentes.append((ref, info))
        elif tipo in ("inteiro", "decimal"):
            entradas[ref] = info.get("maximo")
        else:
            entradas[ref] = texto_pior_caso(tamanho_pior_caso(nome), 7 * k)
    for _ in range(2):
        if not dependentes:
            break
        valores = calcular(entradas)
        for ref, info in dependentes:
            op = opcoes_da_lista(wb, info, ref, valores)
            escolha = mais_longa(ref, op)
            if escolha is not None:
                entradas[ref] = escolha
    return entradas


def busca_dados(formula):
    """A fórmula busca texto na aba Dados (INDEX/VLOOKUP/HLOOKUP sobre 'Dados'!)?"""
    import re
    f = formula or ""
    return "'Dados'!" in f and bool(re.search(r"\b(INDEX|VLOOKUP|HLOOKUP)\(", f))


def _mesclas(ws):
    topo, coberta = {}, set()
    for m in ws.merged_cells.ranges:
        topo[(m.min_row, m.min_col)] = (m.max_row, m.max_col)
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                if (r, c) != (m.min_row, m.min_col):
                    coberta.add((r, c))
    return topo, coberta


def _cortes(bordas, limite):
    """Divide [0, bordas[-1]] em faixas de até `limite` px, cortando numa divisa de
    coluna/linha (`bordas`, crescente) sempre que possível. Devolve [(início, fim)]."""
    faixas, ini, fim_total = [], 0, bordas[-1]
    while ini < fim_total:
        cabe = [b for b in bordas if ini < b <= ini + limite]
        fim = max(cabe) if cabe else min(ini + limite, fim_total)   # coluna mais larga que o limite: corta no px
        faixas.append((ini, fim))
        ini = fim
    return faixas


def _faixa_nomes(bordas, ini, fim, nome):
    """Primeira e última coluna/linha (1-based) visíveis entre os px `ini` e `fim`."""
    idx = [k for k in range(1, len(bordas)) if bordas[k] > bordas[k - 1] and bordas[k - 1] < fim and bordas[k] > ini]
    return (nome(idx[0]), nome(idx[-1])) if idx else ("", "")


def gravar_recortes(img, xs, ys, base):
    """Grava os recortes 1:1 de `img` (desenhada em escala 1:1, com as faixas de cabeçalho
    de CAB_LIN × CAB_COL px) como `<base>-parte-NN.png`, de cima para baixo e da esquerda
    para a direita. Cada recorte repete as letras das colunas e os números das linhas.
    Devolve [(caminho, largura, altura, "A1:L45")]."""
    base = Path(base)
    partes = []
    for y0, y1 in _cortes(ys, LADO_RECORTE - CAB_COL):
        for x0, x1 in _cortes(xs, LADO_RECORTE - CAB_LIN):
            w, h = CAB_LIN + (x1 - x0), CAB_COL + (y1 - y0)
            rec = Image.new("RGB", (w, h), "white")
            rec.paste(img.crop((0, 0, CAB_LIN, CAB_COL)), (0, 0))
            rec.paste(img.crop((CAB_LIN + x0, 0, CAB_LIN + x1, CAB_COL)), (CAB_LIN, 0))
            rec.paste(img.crop((0, CAB_COL + y0, CAB_LIN, CAB_COL + y1)), (0, CAB_COL))
            rec.paste(img.crop((CAB_LIN + x0, CAB_COL + y0, CAB_LIN + x1, CAB_COL + y1)), (CAB_LIN, CAB_COL))
            c_ini, c_fim = _faixa_nomes(xs, x0, x1, get_column_letter)
            l_ini, l_fim = _faixa_nomes(ys, y0, y1, str)
            destino = base.with_name(f"{base.name}-parte-{len(partes) + 1:02d}.png")
            rec.save(destino)
            partes.append((destino, w, h, f"{c_ini}{l_ini}:{c_fim}{l_fim}"))
    return partes


def renderizar_aba(ws, valores, avisos, destino):
    """Desenha a aba `ws` (openpyxl, com fórmulas) usando `valores` {"'Aba'!A1": valor}
    para as células com fórmula e devolve (largura_px, altura_px, cortes, partes), em que
    cortes é uma lista de (ref, texto, motivo) e partes a lista de recortes 1:1 de
    gravar_recortes(). Salva o PNG inteiro em `destino` (reduzido se passar de MAX_LADO)
    e os recortes ao lado dele."""
    n_lin, n_col = ws.max_row, ws.max_column
    xs = [0]
    for c in range(1, n_col + 1):
        xs.append(xs[-1] + px_coluna(ws, c))
    ys = [0]
    for r in range(1, n_lin + 1):
        ys.append(ys[-1] + px_linha(ws, r))
    largura, altura = xs[-1], ys[-1]
    esc = 1.0                      # sempre 1:1; a redução (se houver) é só no PNG inteiro, no fim
    img = Image.new("RGB", (largura + CAB_LIN, altura + CAB_COL), "white")
    d = ImageDraw.Draw(img)
    topo, coberta = _mesclas(ws)

    def X(x):
        return (x + CAB_LIN) * esc

    def Y(y):
        return (y + CAB_COL) * esc

    def valor(r, c):
        cel = ws.cell(r, c)
        v = cel.value
        ref = f"'{ws.title}'!{cel.coordinate}"
        if isinstance(v, str) and v.startswith("="):
            return valores.get(ref), True
        if v is None and valores.get(ref) is not None:      # entrada preenchida no estado
            return valores[ref], True
        return v, v is not None

    # Cabeçalhos de coluna e linha (para localizar a célula na imagem)
    f_cab = fonte(max(8, int(11 * esc)))
    d.rectangle([0, 0, img.width, CAB_COL * esc], fill=COR_CAB)
    d.rectangle([0, 0, CAB_LIN * esc, img.height], fill=COR_CAB)
    for c in range(1, n_col + 1):
        if xs[c] > xs[c - 1]:
            d.text(((X(xs[c - 1]) + X(xs[c])) / 2, CAB_COL * esc / 2), get_column_letter(c), fill=(90, 90, 90),
                   font=f_cab, anchor="mm")
    for r in range(1, n_lin + 1):
        if ys[r] > ys[r - 1]:
            d.text((CAB_LIN * esc / 2, (Y(ys[r - 1]) + Y(ys[r])) / 2), str(r), fill=(90, 90, 90), font=f_cab,
                   anchor="mm")

    cortes = []
    textos = []
    for r in range(1, n_lin + 1):
        for c in range(1, n_col + 1):
            if (r, c) in coberta:
                continue
            r2, c2 = topo.get((r, c), (r, c))
            x0, y0, x1, y1 = xs[c - 1], ys[r - 1], xs[c2], ys[r2]
            if x1 <= x0 or y1 <= y0:
                continue
            cel = ws.cell(r, c)
            ref = f"'{ws.title}'!{cel.coordinate}"
            v, ocupada = valor(r, c)
            s = texto_exibido(v)
            fundo = None
            if cel.fill is not None and cel.fill.patternType == "solid":
                fundo = _rgb(cel.fill.fgColor)
            if ref in avisos and s:
                fundo = COR_AVISO_FUNDO
            caixa = [X(x0), Y(y0), X(x1) - 1, Y(y1) - 1]
            if fundo:
                d.rectangle(caixa, fill=fundo)     # preenchimento esconde a grade, como no Google Planilhas
            else:
                d.rectangle(caixa, outline=COR_GRADE)
            b = cel.border
            for lado, pts in (("top", (X(x0), Y(y0), X(x1), Y(y0))), ("bottom", (X(x0), Y(y1) - 1, X(x1), Y(y1) - 1)),
                              ("left", (X(x0), Y(y0), X(x0), Y(y1))), ("right", (X(x1) - 1, Y(y0), X(x1) - 1, Y(y1)))):
                sd = getattr(b, lado, None) if b is not None else None
                if sd is not None and sd.style:
                    w = 2 if sd.style in ("medium", "double") else 3 if sd.style == "thick" else 1
                    d.line(pts, fill=_rgb(sd.color, (0, 0, 0)), width=max(1, int(w * esc)))
            if not s:
                continue
            ft = cel.font
            pt = (ft.sz if ft is not None and ft.sz else 10)
            px = round(pt * 4 / 3)
            negrito, italico = bool(ft and ft.b), bool(ft and ft.i)
            f = fonte(px, negrito, italico)
            asc, desc = f.getmetrics()
            lh = asc + desc
            al = cel.alignment
            numero = isinstance(v, (int, float)) and not isinstance(v, bool)
            horiz = (al.horizontal if al is not None and al.horizontal not in (None, "general")
                     else ("right" if numero else "left"))
            vert = al.vertical if al is not None and al.vertical else "bottom"
            quebra = bool(al is not None and al.wrap_text)
            util = (x1 - x0) - 2 * MARGEM
            h = y1 - y0
            if quebra:
                linhas = quebrar(s, f, util)
                larga = [ln for ln in linhas if f.getlength(ln) > util]
                if larga:
                    cortes.append((ref, s, f"palavra mais larga que a célula ({int(f.getlength(larga[0]))} px > "
                                           f"{util} px): {larga[0]!r}"))
                if len(linhas) * lh + 2 > h:
                    cortes.append((ref, s, f"{len(linhas)} linha(s) de {lh} px não cabem em {h} px de altura"))
                fim_x = x1
            else:
                linhas = s.split("\n")
                maior = max(f.getlength(ln) for ln in linhas)
                fim_x = x1
                if maior > util:
                    disponivel = util
                    cor_txt = _rgb(ft.color if ft is not None else None, (0, 0, 0)) or (0, 0, 0)
                    claro = sum(cor_txt) / 3 > 200          # texto branco/claro (títulos de bloco)
                    sumiu = False
                    if horiz == "left" and not numero:
                        cc = c2 + 1
                        while disponivel < maior:
                            if cc > n_col:
                                disponivel = maior          # depois da última coluna: tudo vazio
                                sumiu = sumiu or claro      # ...e sem fundo: texto claro some
                                break
                            if (r, cc) in coberta or (r, cc) in topo or valor(r, cc)[1]:
                                break
                            viz = ws.cell(r, cc).fill
                            viz_rgb = _rgb(viz.fgColor) if viz is not None and viz.patternType == "solid" else None
                            if claro and (viz_rgb is None or sum(viz_rgb) / 3 > 200):
                                sumiu = True
                            disponivel += px_coluna(ws, cc)
                            fim_x = xs[cc] if cc <= n_col else fim_x
                            cc += 1
                    if disponivel < maior:
                        cortes.append((ref, s, f"texto de {int(maior)} px não cabe em {int(disponivel)} px "
                                               f"(sem quebra de linha e sem célula vazia ao lado)"))
                    elif sumiu:
                        cortes.append((ref, s, "texto claro transborda para célula sem fundo escuro (some no "
                                               "branco): estenda a faixa do título"))
                if len(linhas) * lh > h + 1:
                    cortes.append((ref, s, f"{len(linhas)} linha(s) de {lh} px não cabem em {h} px de altura"))
            textos.append((linhas, f, _rgb(ft.color if ft is not None else None, (0, 0, 0)) or (0, 0, 0),
                           horiz, vert, x0, y0, fim_x, y1, lh, px, negrito, italico))
    # Texto por cima de tudo (o transbordamento passa sobre a grade vizinha, como no Google)
    for linhas, f, cor, horiz, vert, x0, y0, x1, y1, lh, px, negrito, italico in textos:
        fe = f if esc == 1 else fonte(max(1, round(px * esc)), negrito, italico)
        total = len(linhas) * lh
        topo_y = (y0 + 2 if vert == "top" else y0 + (y1 - y0 - total) / 2 if vert == "center"
                  else y1 - 2 - total)
        for k, ln in enumerate(linhas):
            ty = topo_y + k * lh
            if horiz == "center" or horiz == "centerContinuous":
                tx, anc = (x0 + x1) / 2, "ma"
            elif horiz == "right":
                tx, anc = x1 - MARGEM, "ra"
            else:
                tx, anc = x0 + MARGEM, "la"
            d.text((X(tx), Y(ty)), ln, fill=cor, font=fe, anchor=anc)
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    partes = gravar_recortes(img, xs, ys, destino.with_suffix(""))
    red = min(1.0, MAX_LADO / max(img.width, img.height))
    inteiro = img if red == 1 else img.resize((max(1, int(img.width * red)), max(1, int(img.height * red))),
                                              Image.LANCZOS)
    inteiro.save(destino)
    return largura, altura, cortes, partes


def renderizar_estado(wb, valores, avisos, pasta, estado):
    """Renderiza todas as abas de `wb` para um estado.
    Devolve {aba: (png, largura, altura, cortes, partes)}."""
    saida = {}
    for ws in wb.worksheets:
        nome = ws.title.lower().replace(" ", "-")
        for a, b in (("á", "a"), ("ã", "a"), ("ç", "c"), ("í", "i"), ("é", "e"), ("ó", "o"), ("õ", "o")):
            nome = nome.replace(a, b)
        png = Path(pasta) / f"{estado}-{nome}.png"
        w, h, cortes, partes = renderizar_aba(ws, valores, avisos, png)
        saida[ws.title] = (png, w, h, cortes, partes)
    return saida


def gravar_indice(pasta, saidas):
    """Grava `indice-recortes.txt`: um recorte por linha com tamanho e intervalo de células.
    `saidas` = {estado: retorno de renderizar_estado}."""
    linhas = ["Recortes 1:1 da prévia (no máximo %d × %d px cada). Formato: arquivo · largura × altura · "
              "intervalo de células" % (LADO_RECORTE, LADO_RECORTE), ""]
    for estado, saida in saidas.items():
        for aba, (png, w, h, cortes, partes) in saida.items():
            linhas.append(f"[{estado}] {aba}: aba inteira {w + CAB_LIN} × {h + CAB_COL} px -> {png.name}")
            for p, pw, ph, intervalo in partes:
                linhas.append(f"    {p.name} · {pw} × {ph} · {intervalo}")
    destino = Path(pasta) / "indice-recortes.txt"
    destino.write_text("\n".join(linhas) + "\n", encoding="utf-8")
    return destino


if __name__ == "__main__":
    import sys
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import testar_ficha
    res = testar_ficha.suite_preview(None)
    res.imprimir()
    sys.exit(1 if res.falhas else 0)
